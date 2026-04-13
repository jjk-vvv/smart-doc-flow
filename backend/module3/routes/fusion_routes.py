from __future__ import annotations

import threading
from pathlib import Path
from typing import Any

from fastapi import APIRouter, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, PlainTextResponse

from ..core.bootstrap import ensure_module3_storage
from ..core.config import (
    ALLOWED_SOURCE_SUFFIXES,
    ALLOWED_TEMPLATE_SUFFIXES,
    STORAGE_DIR,
)
from ..core.database import database_enabled
from ..core.security import resolve_preview_path, safe_filename
from ..schemas.run_schemas import (
    ExcelParseRequest,
    UpdateExcelRequest,
    UpdateWordRequest,
    WebRunRequest,
    WordApplyDiffRequest,
    WordParseRequest,
)
from ..services.docset_service import create_docset_record
from ..services.excel_json_service import parse_excel_workbook_to_json
from ..services.preview_service import build_preview_text
from ..services.run_service import cancel_run, execute_web_run, register_run_control
from ..services.storage_service import (
    create_run_record,
    get_owner_mode,
    load_docset_record,
    load_run_record,
    load_template_record,
    new_id,
    run_directory,
    save_docset_record,
    save_upload_file,
    save_template_record,
    set_run_status,
    unique_path,
    utcnow_iso,
)
from ..services.template_service import create_template_record
from ..services.word_json_service import apply_word_section_diffs, parse_word_document_to_json

router = APIRouter()
DEFAULT_OWNER_ID = "default"


def _normalize_legacy_text_lines(value: Any) -> list[str]:
    text = "" if value is None else str(value)
    return text.replace("\r\n", "\n").replace("\r", "\n").split("\n")


def _legacy_sections_to_diffs(
    current_document: dict[str, Any],
    sections: list[Any],
) -> list[dict[str, Any]]:
    current_section_map = {
        str(section.get("id", "")).strip(): section
        for section in current_document.get("sections") or []
    }
    section_diffs: list[dict[str, Any]] = []

    for section in sections:
        payload = section.model_dump() if hasattr(section, "model_dump") else section.dict()
        section_id = str(payload.get("id", "")).strip()
        section_type = str(payload.get("type", "")).strip().lower()
        current_section = current_section_map.get(section_id)
        if not section_id or current_section is None:
            continue

        if section_type == "text":
            current_paragraphs = current_section.get("paragraphs") or []
            if not current_paragraphs:
                continue

            raw_lines = _normalize_legacy_text_lines(payload.get("content"))
            merged_lines = list(raw_lines[: len(current_paragraphs)])
            if len(raw_lines) > len(current_paragraphs):
                merged_lines[-1] = "\n".join(raw_lines[len(current_paragraphs) - 1 :])

            section_diffs.append(
                {
                    "section_id": section_id,
                    "type": "text",
                    "paragraphs": [
                        {
                            "paragraph_index": index,
                            "text": merged_lines[index] if index < len(merged_lines) else "",
                        }
                        for index in range(len(current_paragraphs))
                    ],
                }
            )
            continue

        if section_type == "table":
            requested_rows = payload.get("rows") or []
            current_rows = current_section.get("rows") or []
            cell_diffs: list[dict[str, Any]] = []
            for row_index, row in enumerate(requested_rows):
                if row_index >= len(current_rows):
                    break
                current_row = current_rows[row_index] or []
                for col_index, value in enumerate(row or []):
                    if col_index >= len(current_row):
                        break
                    cell_diffs.append(
                        {
                            "row": row_index,
                            "col": col_index,
                            "value": value,
                        }
                    )

            section_diffs.append(
                {
                    "section_id": section_id,
                    "type": "table",
                    "cells": cell_diffs,
                }
            )

    return [
        section_diff
        for section_diff in section_diffs
        if section_diff.get("paragraphs") or section_diff.get("cells")
    ]


def _normalize_owner_id(value: Any) -> str:
    normalized = str(value or "").strip()
    return normalized or DEFAULT_OWNER_ID


def _resolve_request_owner_id(request: Request, *, explicit_owner_id: str | None = None) -> str:
    headers = request.headers
    authenticated_owner_id = (
        str(headers.get("X-Authenticated-User-Id") or "").strip()
        or str(headers.get("X-Authenticated-Tenant-Id") or "").strip()
    )
    explicit_candidate = (
        str(explicit_owner_id or "").strip()
        or str(request.query_params.get("owner_id") or "").strip()
        or str(request.query_params.get("tenant_id") or "").strip()
    )
    if authenticated_owner_id and explicit_candidate and authenticated_owner_id != explicit_candidate:
        return _normalize_owner_id(authenticated_owner_id)
    header_tenant_id = str(headers.get("X-Tenant-Id") or "").strip()
    if get_owner_mode() == "strict_required" and not (
        authenticated_owner_id or explicit_candidate or header_tenant_id
    ):
        raise HTTPException(status_code=401, detail="owner identity is required in strict mode")
    return _normalize_owner_id(authenticated_owner_id or explicit_candidate or header_tenant_id)


def _require_run_output_path(run_id: str, owner_id: str) -> tuple[dict[str, Any], Path]:
    try:
        record = load_run_record(run_id, owner_id=owner_id)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc))

    output_filename = record.get("output_filename")
    if not isinstance(output_filename, str) or not output_filename.strip():
        raise HTTPException(status_code=404, detail="该任务暂无输出文件")

    output_path = Path(output_filename).resolve()
    if not output_path.exists():
        raise HTTPException(status_code=404, detail="输出文件不存在")
    return record, output_path


def _require_docx_run_output_path(run_id: str, owner_id: str) -> tuple[dict[str, Any], Path]:
    record, output_path = _require_run_output_path(run_id, owner_id)
    if output_path.suffix.lower() != ".docx":
        raise HTTPException(status_code=400, detail="当前仅支持 DOCX 输出文件")
    return record, output_path


def _require_xlsx_run_output_path(run_id: str, owner_id: str) -> tuple[dict[str, Any], Path]:
    record, output_path = _require_run_output_path(run_id, owner_id)
    if output_path.suffix.lower() != ".xlsx":
        raise HTTPException(status_code=400, detail="当前仅支持 XLSX 输出文件")
    return record, output_path


@router.get("/module3/health")
def module3_health() -> dict[str, Any]:
    ensure_module3_storage()
    return {
        "ok": True,
        "service": "module3",
        "template_extensions": sorted(ALLOWED_TEMPLATE_SUFFIXES),
        "source_extensions": sorted(ALLOWED_SOURCE_SUFFIXES),
        "storage_dir": str(STORAGE_DIR),
        "database_enabled": database_enabled(),
        "owner_mode": get_owner_mode(),
    }


@router.post("/module3/template/upload")
async def upload_template(request: Request, file: UploadFile = File(...)) -> dict[str, Any]:
    ensure_module3_storage()
    owner_id = _resolve_request_owner_id(request)

    filename = safe_filename(file.filename or "template.xlsx")
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_TEMPLATE_SUFFIXES:
        raise HTTPException(status_code=400, detail="模板文件目前仅支持 .xlsx 或 .docx")

    template_id = new_id("tpl")
    stored_path = unique_path((STORAGE_DIR / "templates") / f"{template_id}__{filename}")

    try:
        await save_upload_file(file, stored_path)
        record = create_template_record(template_id, filename, stored_path, owner_id=owner_id)
        save_template_record(record)
        return {
            "template_id": record["template_id"],
            "filename": record["filename"],
            "fields": record["fields"],
        }
    except Exception as exc:
        if stored_path.exists():
            stored_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=str(exc))


@router.post("/module3/docset/upload")
async def upload_docset(request: Request, files: list[UploadFile] = File(...)) -> dict[str, Any]:
    ensure_module3_storage()
    owner_id = _resolve_request_owner_id(request)
    if not files:
        raise HTTPException(status_code=400, detail="请至少上传一个数据源文件")

    doc_set_id = new_id("docset")
    doc_dir = (STORAGE_DIR / "docsets") / doc_set_id
    doc_dir.mkdir(parents=True, exist_ok=True)
    saved_files: list[dict[str, str]] = []

    try:
        for upload in files:
            filename = safe_filename(upload.filename or "data")
            suffix = Path(filename).suffix.lower()
            if suffix not in ALLOWED_SOURCE_SUFFIXES:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"不支持的数据源格式 {suffix or '(无扩展名)'}，仅支持 "
                        ".docx/.xlsx/.txt/.md"
                    ),
                )

            stored_path = unique_path(doc_dir / filename)
            await save_upload_file(upload, stored_path)
            saved_files.append({"filename": filename, "path": str(stored_path)})

        record = create_docset_record(doc_set_id, saved_files, owner_id=owner_id)
        save_docset_record(record)
        return {
            "doc_set_id": record["doc_set_id"],
            "count": record["count"],
            "files": record["files"],
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@router.post("/module3/run")
def create_run(request: Request, payload: WebRunRequest) -> dict[str, Any]:
    ensure_module3_storage()
    owner_id = _resolve_request_owner_id(request, explicit_owner_id=payload.owner_id)

    try:
        load_template_record(payload.template_id, owner_id=owner_id)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc))

    try:
        load_docset_record(payload.doc_set_id, owner_id=owner_id)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc))

    run_id = new_id("run")
    run_dir = run_directory(run_id)
    run_dir.mkdir(parents=True, exist_ok=True)
    create_run_record(
        run_id,
        {
            "status": "running",
            "filled": 0,
            "total": 0,
            "fill_rate": 0,
            "output_filename": None,
            "error": None,
            "created_at": utcnow_iso(),
            "template_id": payload.template_id,
            "doc_set_id": payload.doc_set_id,
            "owner_id": owner_id,
        },
    )

    register_run_control(run_id)
    run_request = payload.model_copy(update={"owner_id": owner_id})
    worker = threading.Thread(target=execute_web_run, args=(run_id, run_request), daemon=True)
    worker.start()

    return {"run_id": run_id, "status": "running"}


@router.get("/module3/run/{run_id}")
def get_run_status(run_id: str, request: Request) -> dict[str, Any]:
    owner_id = _resolve_request_owner_id(request)
    try:
        return load_run_record(run_id, owner_id=owner_id)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc))


@router.post("/module3/run/{run_id}/cancel")
def cancel_run_request(run_id: str, request: Request) -> dict[str, Any]:
    owner_id = _resolve_request_owner_id(request)
    try:
        record = load_run_record(run_id, owner_id=owner_id)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc))

    requested = cancel_run(run_id)
    if requested:
        record = set_run_status(
            run_id,
            status="cancel_requested",
            error="Cancellation requested by user.",
        )
    return {
        "run_id": run_id,
        "status": record.get("status"),
        "cancel_requested": requested,
    }


@router.get("/module3/run/{run_id}/download")
def download_run_output(run_id: str, request: Request) -> FileResponse:
    owner_id = _resolve_request_owner_id(request)
    _, output_path = _require_run_output_path(run_id, owner_id)
    return FileResponse(path=str(output_path), filename=output_path.name)


@router.get("/module3/documents/preview")
def preview_document(path: str = Query(...)) -> PlainTextResponse:
    try:
        resolved = resolve_preview_path(path)
        content = build_preview_text(resolved)
        return PlainTextResponse(content, media_type="text/plain; charset=utf-8")
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@router.post("/module3/api/v1/word/parse-json")
def parse_word_document_json(request: Request, payload: WordParseRequest) -> dict[str, Any]:
    owner_id = _resolve_request_owner_id(request)
    _, output_path = _require_docx_run_output_path(payload.run_id, owner_id)

    try:
        parsed_payload = parse_word_document_to_json(output_path)
        return {
            "run_id": payload.run_id,
            "document_name": output_path.name,
            **parsed_payload,
        }
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@router.post("/module3/api/v1/excel/parse-json")
def parse_excel_document_json(request: Request, payload: ExcelParseRequest) -> dict[str, Any]:
    owner_id = _resolve_request_owner_id(request)
    _, output_path = _require_xlsx_run_output_path(payload.run_id, owner_id)

    try:
        parsed_payload = parse_excel_workbook_to_json(output_path)
        return {
            "run_id": payload.run_id,
            "document_name": output_path.name,
            **parsed_payload,
        }
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@router.post("/module3/api/v1/word/apply-diff")
def apply_word_document_diff(request: Request, payload: WordApplyDiffRequest) -> dict[str, Any]:
    owner_id = _resolve_request_owner_id(request)
    _, output_path = _require_docx_run_output_path(payload.run_id, owner_id)

    try:
        section_payloads = [
            section.model_dump() if hasattr(section, "model_dump") else section.dict()
            for section in payload.sections
        ]
        action_payloads = [
            action.model_dump() if hasattr(action, "model_dump") else action.dict()
            for action in payload.actions
        ]
        applied = apply_word_section_diffs(
            output_path,
            section_payloads,
            structure_actions=action_payloads,
        )
        return {
            "status": "ok",
            "message": "Word diff applied successfully",
            "applied": applied,
        }
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@router.post("/module3/run/{run_id}/save")
def save_run_output(run_id: str, request: Request, payload: UpdateExcelRequest) -> dict[str, Any]:
    owner_id = _resolve_request_owner_id(request)
    _, output_path = _require_xlsx_run_output_path(run_id, owner_id)

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(filename=str(output_path))
        known_sheet_names = set(workbook.sheetnames)

        for sheet_data in payload.sheets:
            worksheet = (
                workbook[sheet_data.name]
                if sheet_data.name in known_sheet_names
                else workbook.create_sheet(title=sheet_data.name)
            )
            rows = sheet_data.rows or [[]]
            target_row_count = max(1, len(rows))
            target_column_count = max(1, max((len(row) for row in rows), default=0))
            current_row_count = max(worksheet.max_row or 0, target_row_count)
            current_column_count = max(worksheet.max_column or 0, target_column_count)

            for row_idx in range(1, current_row_count + 1):
                source_row = rows[row_idx - 1] if row_idx - 1 < len(rows) else []
                for col_idx in range(1, current_column_count + 1):
                    cell_value = (
                        source_row[col_idx - 1]
                        if col_idx - 1 < len(source_row)
                        else None
                    )
                    worksheet.cell(
                        row=row_idx,
                        column=col_idx,
                        value=None if cell_value in (None, "") else cell_value,
                    )

        if not workbook.sheetnames:
            workbook.create_sheet(title="Sheet1")

        workbook.save(str(output_path))
        return {"status": "ok", "message": "Excel file saved successfully"}
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@router.post("/module3/run/{run_id}/save-file")
async def save_run_output_file(
    run_id: str,
    request: Request,
    file: UploadFile = File(...),
) -> dict[str, Any]:
    owner_id = _resolve_request_owner_id(request)
    _, output_path = _require_run_output_path(run_id, owner_id)

    incoming_suffix = Path(file.filename or "").suffix.lower()
    current_suffix = output_path.suffix.lower()
    if incoming_suffix and incoming_suffix != current_suffix:
        raise HTTPException(
            status_code=400,
            detail=f"上传文件后缀 {incoming_suffix} 与当前输出文件 {current_suffix} 不一致",
        )

    try:
        payload = await file.read()
        if not payload:
            raise HTTPException(status_code=400, detail="上传文件内容为空")
        output_path.write_bytes(payload)
        return {"status": "ok", "message": "Run output file saved successfully"}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@router.post("/module3/run/{run_id}/save-word")
def save_run_output_word(run_id: str, request: Request, payload: UpdateWordRequest) -> dict[str, Any]:
    owner_id = _resolve_request_owner_id(request)
    _, output_path = _require_docx_run_output_path(run_id, owner_id)

    try:
        ordered_sections = sorted(payload.sections, key=lambda item: item.order)
        if not ordered_sections:
            raise HTTPException(status_code=400, detail="未提供可保存的 Word 内容")

        current_document = parse_word_document_to_json(output_path)
        section_diffs = _legacy_sections_to_diffs(current_document, ordered_sections)
        if not section_diffs:
            raise HTTPException(status_code=400, detail="未检测到可写回的 Word 修改")

        applied = apply_word_section_diffs(output_path, section_diffs)
        return {
            "status": "ok",
            "message": "Word file saved successfully",
            "applied": applied,
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


__all__ = ["router"]
