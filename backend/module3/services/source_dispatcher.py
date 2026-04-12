from __future__ import annotations

import json
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Callable

import pandas as pd
from openpyxl import load_workbook

from ..core.config import DEFAULT_PROMPT
from .llm_service import (
    build_messages,
    build_repair_messages,
    normalize_rows_output,
    parse_model_json,
    request_chat_completions,
    request_responses,
    request_text_responses,
)
from .markdown_processor import extract_rows_from_markdown_source
from .text_source_common import log_event
from .text_processor import extract_rows_from_txt_source
from .word_processor import extract_docx_images, read_docx_content, squish


def prefixed_logger(logger: Callable[[str], None] | None, prefix: str) -> Callable[[str], None] | None:
    if logger is None:
        return None

    def _wrapped(message: str) -> None:
        logger(f"{prefix}{message}")

    return _wrapped


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    text = squish(str(value))
    return "" if text.lower() in {"nan", "none", "null", "nat"} else text


def _excel_engine(path: Path) -> str:
    return "xlrd" if path.suffix.lower() == ".xls" else "openpyxl"


def _normalize_docx_excel_header(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", str(text or ""))
    normalized = normalized.replace("（", "(").replace("）", ")")
    return " ".join(normalized.split()).strip()


def _normalize_header(text: str) -> str:
    return re.sub(r"[\s_\-—–\\()（）\[\]【】:：,.，、/]+", "", str(text or "").strip()).lower()


def _tokenize_header(text: str) -> list[str]:
    raw = str(text or "").strip().lower()
    ascii_tokens = [token for token in re.split(r"[^a-z0-9]+", raw) if token]
    if ascii_tokens:
        return ascii_tokens
    chinese_tokens = re.findall(r"[\u4e00-\u9fff]+", raw)
    if chinese_tokens:
        return chinese_tokens
    normalized = _normalize_header(raw)
    return [normalized] if normalized else []


def _non_empty_columns(values: list[str]) -> list[tuple[int, str]]:
    return [(idx, value) for idx, value in enumerate(values, start=1) if value]


def read_excel_source_sheets(
    source_path: Path,
    *,
    scan_rows: int = 12,
    max_cols: int = 80,
) -> list[dict[str, Any]]:
    wb = load_workbook(filename=str(source_path), read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []

    try:
        for ws in wb.worksheets:
            header_row_index = 0
            header_columns: list[tuple[int, str]] = []

            for row_index, row in enumerate(ws.iter_rows(values_only=True, max_row=scan_rows, max_col=max_cols), start=1):
                values = [_cell_text(value) for value in row]
                columns = _non_empty_columns(values)
                if not columns:
                    continue

                normalized = [_normalize_header(value) for _, value in columns]
                distinct = len({item for item in normalized if item})
                if len(columns) >= 2 and distinct >= 2:
                    header_row_index = row_index
                    header_columns = columns
                    break

                if not header_columns:
                    header_row_index = row_index
                    header_columns = columns

            if not header_columns:
                continue

            header_by_col = {col_idx: header for col_idx, header in header_columns}
            rows: list[dict[str, str]] = []
            for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if row_index <= header_row_index:
                    continue
                row_values = {col_idx: _cell_text(value) for col_idx, value in enumerate(row, start=1)}
                record = {header: row_values.get(col_idx, "") for col_idx, header in header_by_col.items()}
                if any(value for value in record.values()):
                    rows.append(record)

            sheets.append(
                {
                    "sheet_name": ws.title,
                    "header_row": header_row_index,
                    "header_columns": header_columns,
                    "headers": [header for _, header in header_columns],
                    "rows": rows,
                }
            )
    finally:
        wb.close()

    if not sheets:
        raise ValueError("Excel source does not contain any usable worksheets")

    return sheets


def _score_header_match(target: str, source: str) -> tuple[int, int, int, int, int]:
    target_norm = _normalize_header(target)
    source_norm = _normalize_header(source)
    if not target_norm or not source_norm:
        return (0, 0, 0, 0, 0)

    if target_norm == source_norm:
        return (4, 0, 0, 0, -len(source))

    if target_norm in source_norm or source_norm in target_norm:
        return (3, min(len(target_norm), len(source_norm)), 0, 0, -len(source))

    target_tokens = set(_tokenize_header(target))
    source_tokens = set(_tokenize_header(source))
    overlap = len(target_tokens & source_tokens)
    if overlap:
        coverage = int((overlap / max(1, len(target_tokens))) * 100)
        precision = int((overlap / max(1, len(source_tokens))) * 100)
        return (2, coverage, precision, -abs(len(target_tokens) - len(source_tokens)), -len(source))

    return (0, 0, 0, 0, 0)


def simple_match_excel_headers(source_headers: list[str], target_headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used_sources: set[str] = set()

    for target in target_headers:
        best_source = ""
        best_score = (0, 0, 0, 0, 0)

        for source in source_headers:
            if source in used_sources:
                continue
            score = _score_header_match(target, source)
            if score > best_score:
                best_source = source
                best_score = score

        if best_source and best_score[0] > 0:
            mapping[target] = best_source
            used_sources.add(best_source)

    return mapping


def _build_excel_header_mapping_messages(
    source_headers: list[str],
    target_headers: list[str],
    custom_prompt: str,
) -> tuple[str, str]:
    system_prompt = (
        "You align target fields to source Excel headers. "
        "Return strict JSON only. "
        "Do not extract rows. "
        "Only map each target field to one source header or an empty string."
    )
    user_prompt = (
        f"[Source Headers]\n{json.dumps(source_headers, ensure_ascii=False)}\n\n"
        f"[Target Headers]\n{json.dumps(target_headers, ensure_ascii=False)}\n\n"
        f"[User Rules]\n{custom_prompt}\n\n"
        "Return JSON in this format only:\n"
        '{\n  "mapping": {\n    "target_header_1": "source_header_or_empty"\n  }\n}\n\n'
        "Rules:\n"
        "1) Prefer exact and semantic header alignment.\n"
        "2) Only use headers from [Source Headers].\n"
        "3) If there is no confident match, return an empty string.\n"
        "4) Do not invent new headers.\n"
    )
    return system_prompt, user_prompt


def ai_match_excel_headers(
    source_headers: list[str],
    target_headers: list[str],
    *,
    prompt: str,
    settings: dict[str, Any],
    logger: Callable[[str], None] | None = None,
) -> dict[str, str]:
    if not target_headers:
        return {}

    system_prompt, user_prompt = _build_excel_header_mapping_messages(source_headers, target_headers, prompt)
    if settings["api_mode"] == "responses":
        try:
            log_event(logger, "[model] header mapping via /v1/responses")
            model_text = request_text_responses(
                base_url=settings["base_url"],
                api_key=settings["api_key"],
                model=settings["model"],
                system_prompt=system_prompt,
                user_prompt=user_prompt,
            )
        except Exception as exc:
            log_event(logger, f"[model fallback] header mapping responses failed, switching to chat. reason: {exc}")
            model_text = request_chat_completions(
                base_url=settings["base_url"],
                api_key=settings["api_key"],
                model=settings["model"],
                system_prompt=system_prompt,
                user_prompt=user_prompt,
                image_inputs=None,
            )
    else:
        log_event(logger, "[model] header mapping via /v1/chat/completions")
        model_text = request_chat_completions(
            base_url=settings["base_url"],
            api_key=settings["api_key"],
            model=settings["model"],
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            image_inputs=None,
        )

    parsed = parse_model_json(model_text)
    if isinstance(parsed, dict) and isinstance(parsed.get("mapping"), dict):
        parsed_mapping = parsed["mapping"]
    elif isinstance(parsed, dict):
        parsed_mapping = parsed
    else:
        raise ValueError("Header mapping output is not a JSON object")

    allowed = set(source_headers)
    mapping: dict[str, str] = {}
    for target in target_headers:
        value = parsed_mapping.get(target, "")
        if value is None:
            continue
        selected = str(value).strip()
        if selected in allowed:
            mapping[target] = selected

    return mapping


def _ordered_source_headers(sheets: list[dict[str, Any]]) -> list[str]:
    ordered: list[str] = []
    seen: set[str] = set()
    for sheet in sheets:
        for header in sheet.get("headers", []):
            if header and header not in seen:
                seen.add(header)
                ordered.append(header)
    return ordered


def transform_excel_rows(
    sheets: list[dict[str, Any]],
    mapping: dict[str, str],
    target_headers: list[str],
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []

    for sheet in sheets:
        for source_row in sheet.get("rows", []):
            target_row: dict[str, str] = {}
            for target_header in target_headers:
                source_header = mapping.get(target_header, "")
                value = source_row.get(source_header, "") if source_header else ""
                target_row[target_header] = _cell_text(value)
            for source_header, value in source_row.items():
                source_key = _cell_text(source_header)
                if not source_key or source_key in target_row:
                    continue
                target_row[source_key] = _cell_text(value)
            if any(value for value in target_row.values()):
                rows.append(target_row)

    return rows


def _match_docx_headers_to_excel(word_headers: list[str], excel_columns: pd.Index) -> tuple[dict[str, Any], list[str]]:
    labels = list(excel_columns)
    if not labels:
        raise ValueError("Excel 数据源没有可用表头")

    normalized_to_label: dict[str, Any] = {}
    stripped_to_label: dict[str, Any] = {}
    for label in labels:
        stripped = str(label).strip()
        normalized = _normalize_docx_excel_header(label)
        stripped_to_label.setdefault(stripped, label)
        normalized_to_label.setdefault(normalized, label)

    matched: dict[str, Any] = {}
    unresolved: list[str] = []
    used_labels: set[Any] = set()
    for index, header in enumerate(word_headers, start=1):
        stripped = str(header or "").strip()
        if not stripped:
            raise ValueError(
                f"DOCX 模板表格第 {index} 列表头为空，请保留非空表头行。"
            )
        if stripped in stripped_to_label and stripped_to_label[stripped] not in used_labels:
            label = stripped_to_label[stripped]
            matched[stripped] = label
            used_labels.add(label)
            continue
        normalized = _normalize_docx_excel_header(stripped)
        if normalized in normalized_to_label and normalized_to_label[normalized] not in used_labels:
            label = normalized_to_label[normalized]
            matched[stripped] = label
            used_labels.add(label)
            continue

        best_label: Any | None = None
        best_score = (0, 0, 0, 0, 0)
        for label in labels:
            if label in used_labels:
                continue
            score = _score_header_match(stripped, str(label))
            if score > best_score:
                best_score = score
                best_label = label

        if best_label is not None and best_score[0] >= 2:
            matched[stripped] = best_label
            used_labels.add(best_label)
            continue
        unresolved.append(stripped)
    return matched, unresolved


def _is_safe_docx_ai_header_match(target_header: str, source_header: str) -> bool:
    target_text = re.sub(r"[_\-\s]+", "", _normalize_docx_excel_header(target_header))
    source_text = re.sub(r"[_\-\s]+", "", _normalize_docx_excel_header(source_header))
    if not target_text or not source_text:
        return False
    if target_text == source_text:
        return True
    if target_text in source_text or source_text in target_text:
        return True
    return SequenceMatcher(None, target_text, source_text).ratio() >= 0.84


def _extract_rows_from_excel_for_docx_template(
    source_path: Path,
    headers: list[str],
    prompt: str,
    *,
    settings: dict[str, Any],
    table_context: dict[str, Any] | None,
    logger: Callable[[str], None] | None = None,
) -> list[dict[str, str]]:
    read_kwargs: dict[str, Any] = {
        "sheet_name": 0,
        "engine": _excel_engine(source_path),
        "header": 0,
    }

    log_event(logger, f"[step1] loaded Excel source: {source_path.name}")
    dataframe = pd.read_excel(source_path, **read_kwargs)
    matched_mapping, unresolved = _match_docx_headers_to_excel(headers, dataframe.columns)
    if unresolved:
        ai_prompt = prompt
        context_text = _cell_text((table_context or {}).get("context_text", ""))
        if context_text:
            ai_prompt = f"{prompt}\n\n[Current DOCX Table Context]\n{context_text}" if prompt else f"[Current DOCX Table Context]\n{context_text}"
        ai_mapping = ai_match_excel_headers(
            [str(label) for label in dataframe.columns],
            unresolved,
            prompt=ai_prompt or DEFAULT_PROMPT,
            settings=settings,
            logger=logger,
        )
        direct_ai_mapping = {
            header: label
            for header, label in ai_mapping.items()
            if _is_safe_docx_ai_header_match(header, str(label))
        }
        inferred_ai_mapping = {
            header: label
            for header, label in ai_mapping.items()
            if header not in direct_ai_mapping
        }
        matched_mapping.update(ai_mapping)
        if inferred_ai_mapping:
            log_event(
                logger,
                "[header ai detail] 以下 DOCX 表头由 AI 联想匹配："
                + json.dumps(inferred_ai_mapping, ensure_ascii=False),
            )
        unresolved = [header for header in headers if str(header).strip() not in matched_mapping]

    if not matched_mapping:
        log_event(
            logger,
            "[table skipped] 当前源文件与当前 DOCX 表格没有可匹配表头，已跳过。"
            + json.dumps(
                {
                    "target_headers": headers,
                    "excel_headers": [str(label) for label in dataframe.columns],
                },
                ensure_ascii=False,
            ),
        )
        return []

    matched_headers = [header for header in headers if str(header).strip() in matched_mapping]
    rows: list[dict[str, str]] = []
    invalid_row_count = 0
    for _, row in dataframe.iterrows():
        mapped_row = {header: "" for header in headers}
        for header in matched_headers:
            label = matched_mapping[str(header).strip()]
            mapped_row[header] = _cell_text(row[label])
        for source_header in dataframe.columns:
            source_key = _cell_text(source_header)
            if not source_key or source_key in mapped_row:
                continue
            mapped_row[source_key] = _cell_text(row[source_header])
        filled_matched_headers = [header for header in matched_headers if mapped_row.get(header)]
        if not filled_matched_headers:
            continue
        if len(matched_headers) >= 3 and len(filled_matched_headers) == 1:
            invalid_row_count += 1
            continue
        if any(mapped_row.values()):
            rows.append(mapped_row)

    if not rows:
        log_event(logger, "[table skipped] 当前源文件在已匹配列下没有可写入当前 DOCX 表格的有效行，已跳过。")
        return []

    log_event(logger, "[step2] DOCX table headers matched to Excel columns")
    for header in matched_headers:
        label = matched_mapping[str(header).strip()]
        log_event(logger, f"[header match] {header} -> {label}")
    non_exact_mapping = {
        header: matched_mapping[str(header).strip()]
        for header in matched_headers
        if str(header).strip() != str(matched_mapping[str(header).strip()]).strip()
    }
    if non_exact_mapping:
        log_event(
            logger,
            "[header inferred] 以下 DOCX 表头不是完全同名匹配，请用户确认："
            + json.dumps(non_exact_mapping, ensure_ascii=False),
        )
    if unresolved:
        log_event(
            logger,
            "[header unresolved] 以下 DOCX 表头未匹配到 Excel 列，将暂时按空值处理："
            + json.dumps(unresolved, ensure_ascii=False),
        )
    if invalid_row_count:
        log_event(logger, f"[step2.1] filtered invalid Excel rows={invalid_row_count}")
    log_event(logger, f"[step3] DOCX-oriented Excel transform complete, rows={len(rows)}")
    return rows


def extract_rows_from_excel_with_prompt(
    source_path: Path,
    headers: list[str],
    prompt: str,
    *,
    settings: dict[str, Any],
    table_context: dict[str, Any] | None = None,
    logger: Callable[[str], None] | None = None,
) -> list[dict[str, str]]:
    if table_context and str(table_context.get("template_type") or "").lower() == "docx":
        return _extract_rows_from_excel_for_docx_template(
            source_path=source_path,
            headers=headers,
            prompt=prompt,
            settings=settings,
            table_context=table_context,
            logger=logger,
        )

    log_event(logger, f"[step1] loaded Excel source: {source_path.name}")
    log_event(logger, f"[step1.1] extraction guidance received: {prompt}")

    sheets = read_excel_source_sheets(source_path)
    total_rows = sum(len(sheet.get("rows", [])) for sheet in sheets)
    source_headers = _ordered_source_headers(sheets)
    log_event(
        logger,
        f"[step1.2] worksheets={len(sheets)} source_headers={len(source_headers)} source_rows={total_rows}",
    )

    mapping = simple_match_excel_headers(source_headers, headers)
    unresolved = [header for header in headers if header not in mapping]
    log_event(logger, f"[step2] direct header match resolved {len(mapping)}/{len(headers)} fields")

    if unresolved:
        ai_mapping = ai_match_excel_headers(
            source_headers,
            unresolved,
            prompt=prompt,
            settings=settings,
            logger=logger,
        )
        mapping.update(ai_mapping)
        unresolved = [header for header in headers if header not in mapping]
        log_event(logger, f"[step2.1] AI header match resolved {len(ai_mapping)} additional fields")

    if unresolved:
        log_event(logger, f"[step2.2] unresolved target fields: {json.dumps(unresolved, ensure_ascii=False)}")

    rows = transform_excel_rows(sheets, mapping, headers)
    if not rows:
        raise ValueError("No writable row data extracted from Excel source")

    log_event(logger, f"[step3] structured Excel transform complete, rows={len(rows)}")
    log_event(logger, "[mapped rows]\n" + json.dumps({"rows": rows[:10]}, ensure_ascii=False, indent=2))
    return rows


def extract_rows_from_word_source(
    word_path: Path,
    headers: list[str],
    prompt: str,
    *,
    base_url: str,
    api_key: str,
    model: str,
    api_mode: str,
    max_images: int,
    max_image_bytes: int,
    logger: Callable[[str], None] | None = None,
) -> list[dict[str, str]]:
    if word_path.suffix.lower() != ".docx":
        raise ValueError("Word source files must use the .docx extension")

    image_inputs = extract_docx_images(word_path, max_images=max_images, max_image_bytes=max_image_bytes)
    try:
        doc_text = read_docx_content(word_path)
    except ValueError as exc:
        if image_inputs:
            doc_text = ""
            log_event(logger, f"[step1] Word 正文为空，改为仅基于图片继续抽取。原因: {exc}")
        else:
            raise
    system_prompt, user_prompt = build_messages(headers, prompt, doc_text, image_count=len(image_inputs))

    log_event(logger, f"[step1] loaded Word source: {word_path.name}")
    log_event(logger, f"[step1.1] target headers={len(headers)}")
    log_event(
        logger,
        f"[step1.2] extracted {len(image_inputs)} images (max_images={max_images}, max_image_bytes={max_image_bytes})",
    )
    log_event(logger, "[step1.3] sending Word content to model")

    if api_mode == "responses":
        try:
            log_event(logger, "[model] /v1/responses with Word file and images")
            model_text = request_responses(
                base_url=base_url,
                api_key=api_key,
                model=model,
                system_prompt=system_prompt,
                user_prompt=user_prompt,
                word_path=word_path,
                include_file=True,
                image_inputs=image_inputs,
            )
        except Exception as exc:
            log_event(logger, f"[model fallback] file upload failed, retrying without file. reason: {exc}")
            try:
                log_event(logger, "[model] /v1/responses with images only")
                model_text = request_responses(
                    base_url=base_url,
                    api_key=api_key,
                    model=model,
                    system_prompt=system_prompt,
                    user_prompt=user_prompt,
                    word_path=word_path,
                    include_file=False,
                    image_inputs=image_inputs,
                )
            except Exception as exc2:
                try:
                    if not image_inputs:
                        raise exc2
                    log_event(logger, f"[model fallback] image upload failed, retrying as text-only. reason: {exc2}")
                    log_event(logger, "[model] /v1/responses text-only fallback")
                    model_text = request_responses(
                        base_url=base_url,
                        api_key=api_key,
                        model=model,
                        system_prompt=system_prompt,
                        user_prompt=user_prompt,
                        word_path=word_path,
                        include_file=False,
                        image_inputs=None,
                    )
                except Exception as exc3:
                    log_event(logger, f"[model fallback] responses text-only failed, switching to chat. reason: {exc3}")
                    try:
                        if image_inputs:
                            log_event(logger, "[model] /v1/chat/completions with images fallback")
                            model_text = request_chat_completions(
                                base_url=base_url,
                                api_key=api_key,
                                model=model,
                                system_prompt=system_prompt,
                                user_prompt=user_prompt,
                                image_inputs=image_inputs,
                            )
                        else:
                            raise RuntimeError("no image inputs available for chat image fallback")
                    except Exception as exc4:
                        if image_inputs:
                            log_event(logger, f"[model fallback] chat image input failed, retrying text-only. reason: {exc4}")
                        log_event(logger, "[model] /v1/chat/completions text-only fallback")
                        model_text = request_chat_completions(
                            base_url=base_url,
                            api_key=api_key,
                            model=model,
                            system_prompt=system_prompt,
                            user_prompt=user_prompt,
                            image_inputs=None,
                        )
    else:
        try:
            log_event(logger, "[model] /v1/chat/completions with images")
            model_text = request_chat_completions(
                base_url=base_url,
                api_key=api_key,
                model=model,
                system_prompt=system_prompt,
                user_prompt=user_prompt,
                image_inputs=image_inputs,
            )
        except Exception as exc:
            if not image_inputs:
                raise
            log_event(logger, f"[model fallback] chat image input failed, retrying text-only. reason: {exc}")
            log_event(logger, "[model] /v1/chat/completions text-only fallback")
            model_text = request_chat_completions(
                base_url=base_url,
                api_key=api_key,
                model=model,
                system_prompt=system_prompt,
                user_prompt=user_prompt,
                image_inputs=None,
            )

    log_event(logger, "[model output]\n" + model_text)

    try:
        parsed_data = parse_model_json(model_text)
    except Exception as exc:
        log_event(logger, f"[json repair] initial parse failed, retrying with repair prompt. reason: {exc}")
        repair_system_prompt, repair_user_prompt = build_repair_messages(headers, model_text)
        if api_mode == "responses":
            log_event(logger, "[model] JSON repair via /v1/responses")
            repair_text = request_responses(
                base_url=base_url,
                api_key=api_key,
                model=model,
                system_prompt=repair_system_prompt,
                user_prompt=repair_user_prompt,
                word_path=word_path,
                include_file=False,
                image_inputs=None,
            )
        else:
            log_event(logger, "[model] JSON repair via /v1/chat/completions")
            repair_text = request_chat_completions(
                base_url=base_url,
                api_key=api_key,
                model=model,
                system_prompt=repair_system_prompt,
                user_prompt=repair_user_prompt,
                image_inputs=None,
            )
        log_event(logger, "[repair output]\n" + repair_text)
        parsed_data = parse_model_json(repair_text)

    rows_values = normalize_rows_output(headers, parsed_data)
    log_event(logger, f"[step2] normalized Word rows={len(rows_values)}")
    log_event(logger, "[normalized rows]\n" + json.dumps({"rows": rows_values}, ensure_ascii=False, indent=2))
    return rows_values


def extract_rows_from_source(
    source_path: Path,
    headers: list[str],
    prompt: str,
    *,
    settings: dict[str, Any],
    table_context: dict[str, Any] | None = None,
    logger: Callable[[str], None] | None = None,
) -> list[dict[str, str]]:
    suffix = source_path.suffix.lower()
    if suffix == ".docx":
        return extract_rows_from_word_source(
            word_path=source_path,
            headers=headers,
            prompt=prompt,
            base_url=settings["base_url"],
            api_key=settings["api_key"],
            model=settings["model"],
            api_mode=settings["api_mode"],
            max_images=settings["max_images"],
            max_image_bytes=settings["max_image_bytes"],
            logger=logger,
        )
    if suffix == ".xlsx":
        normalized_prompt = prompt.strip() or DEFAULT_PROMPT
        return extract_rows_from_excel_with_prompt(
            source_path=source_path,
            headers=headers,
            prompt=normalized_prompt,
            settings=settings,
            table_context=table_context,
            logger=logger,
        )
    if suffix == ".md":
        return extract_rows_from_markdown_source(
            markdown_path=source_path,
            headers=headers,
            prompt=prompt,
            base_url=settings["base_url"],
            api_key=settings["api_key"],
            model=settings["model"],
            api_mode=settings["api_mode"],
            logger=logger,
        )
    if suffix == ".txt":
        return extract_rows_from_txt_source(
            txt_path=source_path,
            headers=headers,
            prompt=prompt,
            base_url=settings["base_url"],
            api_key=settings["api_key"],
            model=settings["model"],
            api_mode=settings["api_mode"],
            logger=logger,
        )
    raise ValueError(f"Unsupported source file type: {source_path.suffix}")


__all__ = [
    "ai_match_excel_headers",
    "extract_rows_from_excel_with_prompt",
    "extract_rows_from_source",
    "extract_rows_from_word_source",
    "log_event",
    "prefixed_logger",
    "read_excel_source_sheets",
    "simple_match_excel_headers",
    "transform_excel_rows",
]
