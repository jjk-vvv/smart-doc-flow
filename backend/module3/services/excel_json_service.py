from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _is_meaningful_excel_value(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True


def _sheet_bounds(worksheet: Any) -> tuple[int, int]:
    last_row = 0
    last_col = 0

    for row_index, row in enumerate(
        worksheet.iter_rows(values_only=True),
        start=1,
    ):
        row_has_value = False
        row_last_col = 0
        for col_index, value in enumerate(row, start=1):
            if _is_meaningful_excel_value(value):
                row_has_value = True
                row_last_col = col_index
        if row_has_value:
            last_row = row_index
            last_col = max(last_col, row_last_col)

    return last_row, last_col


def _normalize_cell_value(value: Any) -> str | int | float | bool | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    return text if text != "" else None


def parse_excel_workbook_to_json(path: Path) -> dict[str, Any]:
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheets: list[dict[str, Any]] = []

        for index, worksheet in enumerate(workbook.worksheets):
            row_count, column_count = _sheet_bounds(worksheet)
            if row_count <= 0 or column_count <= 0:
                rows: list[list[str | int | float | bool | None]] = [[]]
                row_count = 1
                column_count = 1
            else:
                rows = []
                for row in worksheet.iter_rows(
                    min_row=1,
                    max_row=row_count,
                    min_col=1,
                    max_col=column_count,
                    values_only=True,
                ):
                    rows.append([_normalize_cell_value(cell) for cell in row])

            sheets.append(
                {
                    "id": f"sheet-{index + 1}",
                    "name": worksheet.title,
                    "order": index,
                    "row_count": row_count,
                    "column_count": column_count,
                    "rows": rows,
                }
            )

        return {
            "sheet_count": len(sheets),
            "sheets": sheets,
        }
    finally:
        workbook.close()


__all__ = ["parse_excel_workbook_to_json"]
