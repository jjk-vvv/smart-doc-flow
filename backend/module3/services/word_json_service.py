from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


def _squish(text: str) -> str:
    return " ".join(str(text or "").split()).strip()


def _iter_docx_blocks(doc: DocxDocument):
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield "paragraph", Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield "table", Table(child, doc)


def _font_size_to_half_points(value: Any) -> int | None:
    if value is None:
        return None
    try:
        pt_value = float(value.pt)
    except Exception:
        return None
    return max(1, int(round(pt_value * 2)))


def _rgb_to_hex(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text.lower() if len(text) == 6 else None


def _paragraph_alignment_name(paragraph: Paragraph) -> str | None:
    alignment = paragraph.alignment
    mapping = {
        0: "left",
        1: "center",
        2: "right",
        3: "justify",
        4: "distribute",
    }
    if alignment is None:
        return None
    return mapping.get(int(alignment))


def _run_style_payload(run: Any) -> dict[str, Any]:
    font = run.font
    payload = {
        "w:b": bool(run.bold) if run.bold is not None else None,
        "w:i": bool(run.italic) if run.italic is not None else None,
        "w:u": bool(run.underline) if run.underline is not None else None,
        "w:strike": bool(run.font.strike) if run.font.strike is not None else None,
        "w:sz": _font_size_to_half_points(font.size),
        "color": _rgb_to_hex(getattr(font.color, "rgb", None)),
        "highlight": str(font.highlight_color).lower() if font.highlight_color is not None else None,
        "font": font.name or None,
    }
    return {key: value for key, value in payload.items() if value not in {None, ""}}


def _paragraph_style_payload(paragraph: Paragraph) -> dict[str, Any]:
    style_id = getattr(getattr(paragraph, "style", None), "style_id", None)
    payload = {
        "style_id": str(style_id).strip() if style_id else None,
        "alignment": _paragraph_alignment_name(paragraph),
    }
    return {key: value for key, value in payload.items() if value not in {None, ""}}


def _run_payload(run: Any) -> dict[str, Any] | None:
    text = run.text or ""
    if not text:
        return None
    return {
        "text": text,
        "style": _run_style_payload(run),
    }


def _paragraph_payload(paragraph: Paragraph, paragraph_index: int) -> dict[str, Any] | None:
    paragraph_text = paragraph.text or ""
    if not paragraph_text.strip():
        return None

    runs = [_run_payload(run) for run in paragraph.runs]
    filtered_runs = [item for item in runs if item is not None]
    if not filtered_runs:
        filtered_runs = [{"text": paragraph_text, "style": {}}]

    return {
        "index": paragraph_index,
        "text": paragraph_text,
        "style": _paragraph_style_payload(paragraph),
        "runs": filtered_runs,
    }


def _cell_payload(cell: Any) -> dict[str, Any]:
    paragraphs = [paragraph for paragraph in cell.paragraphs if (paragraph.text or "").strip()]
    first_paragraph = paragraphs[0] if paragraphs else None
    first_run = first_paragraph.runs[0] if first_paragraph and first_paragraph.runs else None
    payload = {
        "text": "\n".join((paragraph.text or "") for paragraph in paragraphs).strip(),
        "style": {
            **(_paragraph_style_payload(first_paragraph) if first_paragraph is not None else {}),
            **(_run_style_payload(first_run) if first_run is not None else {}),
        },
    }
    if not payload["style"]:
        payload["style"] = {}
    return payload


def _word_length_to_pixels(value: Any) -> int | None:
    if value is None:
        return None

    try:
        inches = getattr(value, "inches", None)
        if inches is not None:
            return max(60, int(round(float(inches) * 96)))
    except Exception:
        pass

    try:
        twips = getattr(value, "twips", None)
        if twips is not None:
            return max(60, int(round((int(twips) / 20) * (96 / 72))))
    except Exception:
        pass

    try:
        raw_value = int(value)
    except Exception:
        return None

    # Raw XML widths are typically twips-sized numbers, but python-docx Length
    # objects stringify to EMU-like integers. Treat only moderate values as
    # plain twips and ignore obviously inflated raw lengths.
    if 0 < raw_value <= 10000:
        return max(60, int(round((raw_value / 20) * (96 / 72))))
    return None


def _extract_column_widths(table: Table) -> list[int]:
    widths: list[int] = []
    try:
        grid = table._tbl.tblGrid
        if grid is not None:
            for col in grid.gridCol_lst:
                width = getattr(col, "w", None)
                widths.append(_word_length_to_pixels(width) or 160)
    except Exception:
        widths = []

    if widths:
        return widths

    max_columns = max((len(row.cells) for row in table.rows), default=1)
    return [160 for _ in range(max_columns or 1)]


def _table_payload(table: Table, table_index: int, order: int) -> dict[str, Any] | None:
    rows: list[list[dict[str, Any]]] = []
    meaningful_cell_count = 0

    for row in table.rows:
        cells = [_cell_payload(cell) for cell in row.cells]
        if any((cell.get("text") or "").strip() for cell in cells):
            meaningful_cell_count += sum(1 for cell in cells if (cell.get("text") or "").strip())
        rows.append(cells)

    if meaningful_cell_count == 0:
        return None

    column_count = max((len(row) for row in rows), default=0)
    return {
        "type": "table",
        "id": f"word-table-{table_index}",
        "name": f"表格 {table_index}",
        "order": order,
        "row_count": len(rows),
        "column_count": column_count,
        "column_widths": _extract_column_widths(table),
        "rows": rows,
        "virtualized": len(rows) > 500,
    }


def parse_word_document_to_json(path: Path) -> dict[str, Any]:
    document = Document(str(path))
    sections: list[dict[str, Any]] = []
    paragraph_buffer: list[dict[str, Any]] = []
    text_section_index = 1
    table_section_index = 1
    order = 0

    def flush_text_section() -> None:
        nonlocal paragraph_buffer, text_section_index, order
        if not paragraph_buffer:
            return
        sections.append(
            {
                "type": "text",
                "id": f"word-text-{text_section_index}",
                "name": f"文本 {text_section_index}",
                "order": order,
                "paragraphs": paragraph_buffer,
            }
        )
        text_section_index += 1
        order += 1
        paragraph_buffer = []

    paragraph_index = 0
    for block_type, block in _iter_docx_blocks(document):
        if block_type == "paragraph":
            payload = _paragraph_payload(block, paragraph_index)
            paragraph_index += 1
            if payload is not None:
                paragraph_buffer.append(payload)
            continue

        flush_text_section()
        table_payload = _table_payload(block, table_section_index, order)
        table_section_index += 1
        if table_payload is None:
            continue
        sections.append(table_payload)
        order += 1

    flush_text_section()
    return {
        "sections": sections,
        "section_count": len(sections),
    }


class _TextSectionRef:
    def __init__(self, kind: str, paragraphs: list[Paragraph]) -> None:
        self.kind = kind
        self.paragraphs = paragraphs


class _TableSectionRef:
    def __init__(self, kind: str, table: Table) -> None:
        self.kind = kind
        self.table = table


def _build_section_refs(document: DocxDocument) -> dict[str, _TextSectionRef | _TableSectionRef]:
    refs: dict[str, _TextSectionRef | _TableSectionRef] = {}
    paragraph_buffer: list[Paragraph] = []
    text_section_index = 1
    table_section_index = 1

    def flush_text_section() -> None:
        nonlocal paragraph_buffer, text_section_index
        visible = [paragraph for paragraph in paragraph_buffer if (paragraph.text or "").strip()]
        if not visible:
            paragraph_buffer = []
            return
        refs[f"word-text-{text_section_index}"] = _TextSectionRef(kind="text", paragraphs=visible)
        text_section_index += 1
        paragraph_buffer = []

    for block_type, block in _iter_docx_blocks(document):
        if block_type == "paragraph":
            paragraph_buffer.append(block)
            continue

        flush_text_section()
        refs[f"word-table-{table_section_index}"] = _TableSectionRef(kind="table", table=block)
        table_section_index += 1

    flush_text_section()
    return refs


def _normalize_diff_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _iter_run_elements(container: Any) -> list[Any]:
    return [element for element in container.iter() if element.tag == qn("w:r")]


def _run_has_text(run_element: Any) -> bool:
    return any(element.tag == qn("w:t") for element in run_element.iter())


def _first_run_element_with_text(paragraph: Paragraph) -> Any | None:
    for run_element in _iter_run_elements(paragraph._p):
        if _run_has_text(run_element):
            return run_element
    return None


def _first_run_element(paragraph: Paragraph) -> Any | None:
    runs = _iter_run_elements(paragraph._p)
    return runs[0] if runs else None


def _ensure_carrier_run(paragraph: Paragraph) -> Any:
    carrier = _first_run_element_with_text(paragraph)
    if carrier is not None:
        return carrier

    carrier = _first_run_element(paragraph)
    if carrier is not None:
        return carrier

    return paragraph.add_run()._r


def _reset_run_content(run_element: Any) -> None:
    for child in list(run_element):
        if child.tag != qn("w:rPr"):
            run_element.remove(child)


def _append_text_node(run_element: Any, text: str) -> None:
    text_node = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        text_node.set(_XML_SPACE, "preserve")
    text_node.text = text
    run_element.append(text_node)


def _set_run_text_xml(run_element: Any, value: Any) -> None:
    normalized = _normalize_diff_text(value)
    parts = normalized.split("\n")
    _reset_run_content(run_element)

    if not parts:
        parts = [""]

    for index, part in enumerate(parts):
        _append_text_node(run_element, part)
        if index < len(parts) - 1:
            run_element.append(OxmlElement("w:br"))


def _remove_other_runs(container: Any, keep_run: Any) -> None:
    for run_element in list(_iter_run_elements(container)):
        if run_element is keep_run:
            continue
        parent = run_element.getparent()
        if parent is not None:
            parent.remove(run_element)


def _replace_paragraph_text_in_place(paragraph: Paragraph, value: Any) -> None:
    carrier_run = _ensure_carrier_run(paragraph)
    _set_run_text_xml(carrier_run, value)
    _remove_other_runs(paragraph._p, carrier_run)


def _replace_cell_text_in_place(cell: Any, value: Any) -> None:
    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        paragraphs = [cell.add_paragraph()]

    carrier_paragraph = next(
        (paragraph for paragraph in paragraphs if _first_run_element_with_text(paragraph) is not None),
        paragraphs[0],
    )
    carrier_run = _ensure_carrier_run(carrier_paragraph)
    _set_run_text_xml(carrier_run, value)

    for paragraph in paragraphs:
        _remove_other_runs(paragraph._p, carrier_run)


def _get_child_by_tag(parent: Any, tag: str) -> Any | None:
    if parent is None:
        return None
    qualified = qn(tag)
    for child in parent:
        if child.tag == qualified:
            return child
    return None


def _get_grid_span(tc_element: Any) -> int:
    tc_pr = getattr(tc_element, "tcPr", None)
    grid_span = _get_child_by_tag(tc_pr, "w:gridSpan")
    if grid_span is None:
        return 1
    try:
        return max(1, int(grid_span.get(qn("w:val")) or 1))
    except Exception:
        return 1


def _set_grid_span(tc_element: Any, span: int) -> None:
    tc_pr = getattr(tc_element, "tcPr", None)
    if tc_pr is None:
        tc_pr = OxmlElement("w:tcPr")
        tc_element.insert(0, tc_pr)
    grid_span = _get_child_by_tag(tc_pr, "w:gridSpan")
    if span <= 1:
        if grid_span is not None:
            tc_pr.remove(grid_span)
        return
    if grid_span is None:
        grid_span = OxmlElement("w:gridSpan")
        tc_pr.append(grid_span)
    grid_span.set(qn("w:val"), str(span))


def _get_vmerge_state(tc_element: Any) -> str | None:
    tc_pr = getattr(tc_element, "tcPr", None)
    vmerge = _get_child_by_tag(tc_pr, "w:vMerge")
    if vmerge is None:
        return None
    return str(vmerge.get(qn("w:val")) or "continue").strip().lower()


def _set_vmerge_state(tc_element: Any, state: str | None) -> None:
    tc_pr = getattr(tc_element, "tcPr", None)
    if tc_pr is None and state is None:
        return
    if tc_pr is None:
        tc_pr = OxmlElement("w:tcPr")
        tc_element.insert(0, tc_pr)

    vmerge = _get_child_by_tag(tc_pr, "w:vMerge")
    if state is None:
        if vmerge is not None:
            tc_pr.remove(vmerge)
        return
    if vmerge is None:
        vmerge = OxmlElement("w:vMerge")
        tc_pr.append(vmerge)
    vmerge.set(qn("w:val"), state)


def _clear_table_cell_xml(tc_element: Any) -> None:
    for child in list(tc_element):
        if child.tag != qn("w:tcPr"):
            tc_element.remove(child)
    tc_element.append(OxmlElement("w:p"))


def _clone_row_with_empty_content(row_element: Any) -> Any:
    cloned = deepcopy(row_element)
    for tc_element in getattr(cloned, "tc_lst", []):
        _clear_table_cell_xml(tc_element)
    return cloned


def _row_has_merge_markup(row_element: Any) -> bool:
    for tc_element in getattr(row_element, "tc_lst", []):
        if _get_vmerge_state(tc_element) is not None or _get_grid_span(tc_element) > 1:
            return True
    return False


def _fix_vertical_merge_after_row_delete(table_element: Any, row_index: int) -> None:
    rows = list(getattr(table_element, "tr_lst", []))
    if row_index < 0 or row_index >= len(rows):
        return

    deleted_row = rows[row_index]
    if not _row_has_merge_markup(deleted_row):
        return

    deleted_cells = list(getattr(deleted_row, "tc_lst", []))
    next_row = rows[row_index + 1] if row_index + 1 < len(rows) else None
    if next_row is None:
        return
    next_cells = list(getattr(next_row, "tc_lst", []))

    for col_index, deleted_cell in enumerate(deleted_cells):
        deleted_vmerge = _get_vmerge_state(deleted_cell)
        if deleted_vmerge != "restart":
            continue
        if col_index >= len(next_cells):
            continue
        next_state = _get_vmerge_state(next_cells[col_index])
        if next_state is not None:
            _set_vmerge_state(next_cells[col_index], "restart")


def _iter_visual_cell_entries(row_element: Any) -> list[tuple[int, int, Any]]:
    entries: list[tuple[int, int, Any]] = []
    cursor = 0
    for tc_element in getattr(row_element, "tc_lst", []):
        span = _get_grid_span(tc_element)
        entries.append((cursor, cursor + span - 1, tc_element))
        cursor += span
    return entries


def _delete_row(table_ref: _TableSectionRef, index: int) -> bool:
    table_element = table_ref.table._tbl
    rows = list(getattr(table_element, "tr_lst", []))
    if index < 0 or index >= len(rows):
        return False
    _fix_vertical_merge_after_row_delete(table_element, index)
    row_element = rows[index]
    parent = row_element.getparent()
    if parent is None:
        return False
    parent.remove(row_element)
    return True


def _insert_row(table_ref: _TableSectionRef, index: int) -> bool:
    table_element = table_ref.table._tbl
    rows = list(getattr(table_element, "tr_lst", []))
    if not rows:
        return False
    target_index = min(max(index, 0), len(rows))
    template_index = min(target_index, len(rows) - 1)
    template_row = rows[template_index]
    new_row = _clone_row_with_empty_content(template_row)

    if target_index >= len(rows):
        table_element.append(new_row)
    else:
        rows[target_index].addprevious(new_row)
    return True


def _delete_col(table_ref: _TableSectionRef, index: int) -> bool:
    if index < 0:
        return False
    changed = False
    for row_element in getattr(table_ref.table._tbl, "tr_lst", []):
        entries = _iter_visual_cell_entries(row_element)
        for start, end, tc_element in entries:
            if not (start <= index <= end):
                continue
            span = _get_grid_span(tc_element)
            if span > 1:
                _set_grid_span(tc_element, span - 1)
            else:
                parent = tc_element.getparent()
                if parent is not None:
                    parent.remove(tc_element)
            changed = True
            break
    return changed


def _insert_col(table_ref: _TableSectionRef, index: int) -> bool:
    if index < 0:
        return False
    changed = False
    for row_element in getattr(table_ref.table._tbl, "tr_lst", []):
        entries = _iter_visual_cell_entries(row_element)
        if not entries:
            continue

        inserted = False
        for start, end, tc_element in entries:
            if start < index <= end:
                _set_grid_span(tc_element, _get_grid_span(tc_element) + 1)
                inserted = True
                break
            if index == start:
                new_cell = deepcopy(tc_element)
                _clear_table_cell_xml(new_cell)
                tc_element.addprevious(new_cell)
                inserted = True
                break

        if not inserted:
            _, _, template_cell = entries[-1]
            new_cell = deepcopy(template_cell)
            _clear_table_cell_xml(new_cell)
            template_cell.addnext(new_cell)

        changed = True
    return changed


def _apply_structure_action(table_ref: _TableSectionRef, action: dict[str, Any]) -> int:
    action_type = str(action.get("action_type", "")).strip().lower()
    index = int(action.get("index", -1))
    count = max(0, int(action.get("count", 1) or 1))
    if index < 0 or count <= 0:
        return 0

    applied = 0
    for _ in range(count):
        if action_type == "delete_row":
            changed = _delete_row(table_ref, index)
        elif action_type == "insert_row":
            changed = _insert_row(table_ref, index)
        elif action_type == "delete_col":
            changed = _delete_col(table_ref, index)
        elif action_type == "insert_col":
            changed = _insert_col(table_ref, index)
        else:
            changed = False

        if not changed:
            break
        applied += 1

    return applied


def apply_word_section_diffs(
    path: Path,
    section_diffs: list[dict[str, Any]],
    structure_actions: list[dict[str, Any]] | None = None,
) -> int:
    document = Document(str(path))
    refs = _build_section_refs(document)
    applied = 0

    for action in structure_actions or []:
        section_id = str(action.get("section_id", "")).strip()
        if not section_id or section_id not in refs:
            continue
        section_ref = refs[section_id]
        if not isinstance(section_ref, _TableSectionRef):
            continue
        applied += _apply_structure_action(section_ref, action)

    if structure_actions:
        refs = _build_section_refs(document)

    for section_diff in section_diffs:
        section_id = str(section_diff.get("section_id", "")).strip()
        section_type = str(section_diff.get("type", "")).strip().lower()
        if not section_id or section_id not in refs:
            continue

        section_ref = refs[section_id]
        if section_type == "text" and isinstance(section_ref, _TextSectionRef):
            for paragraph_diff in section_diff.get("paragraphs") or []:
                paragraph_index = int(paragraph_diff.get("paragraph_index", -1))
                if paragraph_index < 0 or paragraph_index >= len(section_ref.paragraphs):
                    continue
                _replace_paragraph_text_in_place(
                    section_ref.paragraphs[paragraph_index],
                    paragraph_diff.get("text"),
                )
                applied += 1
            continue

        if section_type == "table" and isinstance(section_ref, _TableSectionRef):
            for cell_diff in section_diff.get("cells") or []:
                row_index = int(cell_diff.get("row", -1))
                col_index = int(cell_diff.get("col", -1))
                if row_index < 0 or col_index < 0:
                    continue
                if row_index >= len(section_ref.table.rows):
                    continue
                row = section_ref.table.rows[row_index]
                if col_index >= len(row.cells):
                    continue
                _replace_cell_text_in_place(
                    row.cells[col_index],
                    cell_diff.get("value"),
                )
                applied += 1

    if applied > 0:
        document.save(str(path))
    return applied


__all__ = [
    "apply_word_section_diffs",
    "parse_word_document_to_json",
]
