from __future__ import annotations

from importlib import import_module
from pathlib import Path

from openpyxl import Workbook, load_workbook


template_service = import_module("app.modules.module3.services.template_service")


def _build_template(template_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "城市名"
    sheet["B1"] = "GDP总量（亿元）"
    sheet["C1"] = "一般公共预算收入（亿元）"
    sheet["D1"] = "编码"
    workbook.save(template_path)
    workbook.close()


def test_write_filled_excel_coerces_numeric_strings_but_preserves_ids(tmp_path: Path) -> None:
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "filled.xlsx"
    _build_template(template_path)

    template_service.write_filled_excel(
        template_path=template_path,
        output_path=output_path,
        sheet_name="Sheet1",
        header_columns=[
            (1, "城市名"),
            (2, "GDP总量（亿元）"),
            (3, "一般公共预算收入（亿元）"),
            (4, "编码"),
        ],
        rows_values=[
            {
                "城市名": "北京",
                "GDP总量（亿元）": "52073.40",
                "一般公共预算收入（亿元）": "6680.60",
                "编码": "00123",
            }
        ],
        write_row=2,
    )

    workbook = load_workbook(output_path, data_only=False)
    sheet = workbook["Sheet1"]

    assert sheet["A2"].value == "北京"
    assert isinstance(sheet["B2"].value, float)
    assert sheet["B2"].value == 52073.4
    assert isinstance(sheet["C2"].value, float)
    assert sheet["C2"].value == 6680.6
    assert sheet["D2"].value == "00123"

    workbook.close()
