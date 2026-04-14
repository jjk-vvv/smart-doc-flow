from __future__ import annotations

import io
import os
import time
import zipfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from xml.etree import ElementTree as ET

from ..core.format.xlsx_artifact_io import load_xlsx_sheet_rows_artifact
from ..core.format.xlsx_runtime import XlsxRuntime, XlsxTableRuntime
from ..core.table.table_header_runtime import TableHeaderRuntime
from ..core.table.table_metadata import (
    build_header_meta,
    extract_number_meta,
    normalize_text,
    normalize_token_for_compare,
)

_WORKBOOK_ARTIFACT_VERSION = "module2-xlsx-workbook-v1"
_SHEET_ROWS_ARTIFACT_VERSION = "module2-xlsx-sheet-rows-v2"


class XlsxWorkbookArtifactRuntime:
    """Persist full XLSX workbook artifacts while keeping assistant payloads slim."""

    def __init__(
        self,
        *,
        artifact_support_runtime,
        logger: Any,
        getenv=os.getenv,
    ) -> None:
        self._artifact_support_runtime = artifact_support_runtime
        self._logger = logger
        sheet_limit_raw = int(str(getenv("MODULE2_XLSX_WORKBOOK_SHEET_LIMIT", "128") or "128"))
        self._sheet_limit = max(1, sheet_limit_raw) if sheet_limit_raw > 0 else 128
        self._projection_row_limit = max(
            1,
            int(
                str(
                    getenv(
                        "MODULE2_XLSX_ARTIFACT_PROJECTION_ROWS",
                        getenv("MODULE2_TABLE_MAX_ROWS", "80"),
                    )
                    or getenv("MODULE2_TABLE_MAX_ROWS", "80")
                )
            ),
        )
        self._projection_col_limit = max(
            1,
            int(
                str(
                    getenv(
                        "MODULE2_XLSX_ARTIFACT_PROJECTION_COLS",
                        getenv("MODULE2_TABLE_MAX_COLS", "24"),
                    )
                    or getenv("MODULE2_TABLE_MAX_COLS", "24")
                )
            ),
        )
        self._xlsx_runtime = XlsxRuntime(
            normalize_text=normalize_text,
            sheet_limit=self._sheet_limit,
        )
        self._table_header_runtime = TableHeaderRuntime(
            normalize_text=normalize_text,
            normalize_token_for_compare=normalize_token_for_compare,
            contains_figure_marker=lambda _: False,
            contains_catalog_marker=lambda _: False,
            runtime_page_count_provider=lambda: 1,
        )

    def build_artifacts_for_pipeline(
        self,
        *,
        file_name: str,
        file_hash: str,
        file_bytes: bytes,
        table_blocks: List[Dict[str, Any]],
        parse_stats: Dict[str, Any],
    ) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
        if not str(file_name or "").lower().endswith(".xlsx"):
            return list(table_blocks or []), dict(parse_stats or {})

        source_tag = self._resolve_source_tag(table_blocks)
        extract_t0 = time.perf_counter()
        workbook_payload, workbook_mode = self._extract_full_workbook_payload(
            file_bytes=file_bytes,
            source_tag=source_tag,
        )
        extract_ms = int((time.perf_counter() - extract_t0) * 1000)
        sheets = list(workbook_payload.get("sheets") or [])
        if not sheets:
            stats = dict(parse_stats or {})
            stats.update(
                {
                    "xlsx_workbook_build_mode": workbook_mode,
                    "xlsx_workbook_extract_ms": extract_ms,
                    "xlsx_workbook_persist_ms": 0,
                    "xlsx_workbook_total_ms": extract_ms,
                }
            )
            return list(table_blocks or []), stats

        persist_t0 = time.perf_counter()
        manifest_fragment, manifest_url, sheet_manifests = self._persist_workbook_artifacts(
            file_name=file_name,
            file_hash=file_hash,
            workbook_payload=workbook_payload,
            workbook_mode=workbook_mode,
        )
        persist_ms = int((time.perf_counter() - persist_t0) * 1000)
        if not manifest_fragment:
            stats = dict(parse_stats or {})
            stats.update(
                {
                    "xlsx_workbook_build_mode": workbook_mode,
                    "xlsx_workbook_extract_ms": extract_ms,
                    "xlsx_workbook_persist_ms": persist_ms,
                    "xlsx_workbook_total_ms": extract_ms + persist_ms,
                }
            )
            return list(table_blocks or []), stats

        enriched_tables = self._merge_table_blocks_with_sheet_manifests(
            table_blocks=table_blocks,
            sheet_manifests=sheet_manifests,
            source_tag=source_tag,
            workbook_manifest_path=manifest_fragment,
            workbook_manifest_url=manifest_url,
        )
        stats = dict(parse_stats or {})
        stats.update(
            {
                "xlsx_workbook_artifact_version": _WORKBOOK_ARTIFACT_VERSION,
                "xlsx_sheet_rows_artifact_version": _SHEET_ROWS_ARTIFACT_VERSION,
                "xlsx_workbook_manifest_path": manifest_fragment,
                "xlsx_workbook_manifest_url": manifest_url,
                "xlsx_workbook_sheet_count": len(sheet_manifests),
                "xlsx_source_sheet_count": int(workbook_payload.get("sheet_count", 0) or len(sheet_manifests)),
                "xlsx_source_total_rows": int(workbook_payload.get("source_total_rows", 0) or 0),
                "xlsx_source_total_cols": int(workbook_payload.get("source_total_cols", 0) or 0),
                "xlsx_source_total_nonempty_rows": int(workbook_payload.get("source_total_nonempty_rows", 0) or 0),
                "xlsx_source_total_nonempty_cells": int(workbook_payload.get("source_total_nonempty_cells", 0) or 0),
                "xlsx_workbook_build_mode": workbook_mode,
                "xlsx_workbook_extract_ms": extract_ms,
                "xlsx_workbook_persist_ms": persist_ms,
                "xlsx_workbook_total_ms": extract_ms + persist_ms,
            }
        )
        return enriched_tables, stats

    def _extract_full_workbook_payload(
        self,
        *,
        file_bytes: bytes,
        source_tag: str,
    ) -> Tuple[Dict[str, Any], str]:
        payload = self._extract_openpyxl_workbook(file_bytes=file_bytes, source_tag=source_tag, data_only=True)
        if self._has_sheet_content(payload):
            return payload, "openpyxl_data_only"

        payload = self._extract_zip_workbook(file_bytes=file_bytes, source_tag=source_tag)
        if self._has_sheet_content(payload):
            return payload, "zip_xml"

        payload = self._extract_openpyxl_workbook(file_bytes=file_bytes, source_tag=source_tag, data_only=False)
        if self._has_sheet_content(payload):
            return payload, "openpyxl_formula"
        return {"sheets": []}, "unavailable"

    @staticmethod
    def _has_sheet_content(payload: Dict[str, Any]) -> bool:
        sheets = payload.get("sheets") if isinstance(payload, dict) else []
        return bool(sheets)

    def _extract_openpyxl_workbook(
        self,
        *,
        file_bytes: bytes,
        source_tag: str,
        data_only: bool,
    ) -> Dict[str, Any]:
        try:
            from openpyxl import load_workbook
        except Exception:
            return {"sheets": []}

        try:
            workbook = load_workbook(io.BytesIO(file_bytes), data_only=data_only, read_only=True)
        except Exception:
            return {"sheets": []}

        sheets: List[Dict[str, Any]] = []
        total_nonempty_rows = 0
        total_nonempty_cells = 0
        total_source_rows = 0
        total_source_cols = 0
        try:
            for sheet_idx, worksheet in enumerate(getattr(workbook, "worksheets", [])[: self._sheet_limit], start=1):
                rows_payload: List[Dict[str, Any]] = []
                used_cols: set[int] = set()
                nonempty_cells = 0
                source_max_row = max(int(getattr(worksheet, "max_row", 0) or 0), 1)
                source_max_col = max(int(getattr(worksheet, "max_column", 0) or 0), 1)
                for row_idx, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                    cells_out: List[Dict[str, Any]] = []
                    for col_idx, cell_value in enumerate(row, start=1):
                        text = normalize_text("" if cell_value is None else str(cell_value))
                        if not text:
                            continue
                        used_cols.add(col_idx)
                        nonempty_cells += 1
                        bbox = XlsxTableRuntime.build_bbox(
                            sheet_idx=sheet_idx,
                            row_idx=row_idx,
                            col_idx=col_idx,
                            max_rows=max(1, int(source_max_row or row_idx)),
                            max_cols=max(1, int(source_max_col or col_idx)),
                        )
                        cells_out.append(
                            {
                                "col_index": col_idx,
                                "text": text,
                                "ref_id": f"{source_tag}_t{sheet_idx}_r{row_idx}_c{col_idx}",
                                "bbox": bbox,
                            }
                        )
                    if cells_out:
                        rows_payload.append({"row_index": row_idx, "cells": cells_out})
                sheet_payload = self._finalize_sheet_payload(
                    sheet_index=sheet_idx,
                    sheet_name=str(getattr(worksheet, "title", "") or f"Sheet{sheet_idx}"),
                    rows_payload=rows_payload,
                    source_max_row=source_max_row,
                    source_max_col=source_max_col,
                    used_cols=used_cols,
                    nonempty_cells=nonempty_cells,
                )
                sheets.append(sheet_payload)
                total_nonempty_rows += int(sheet_payload.get("source_nonempty_row_count", 0) or 0)
                total_nonempty_cells += int(sheet_payload.get("source_nonempty_cell_count", 0) or 0)
                total_source_rows += int(sheet_payload.get("source_max_row", 0) or 0)
                total_source_cols += int(sheet_payload.get("source_max_col", 0) or 0)
            return {
                "sheet_count": len(sheets),
                "source_total_rows": total_source_rows,
                "source_total_cols": total_source_cols,
                "source_total_nonempty_rows": total_nonempty_rows,
                "source_total_nonempty_cells": total_nonempty_cells,
                "sheets": sheets,
            }
        finally:
            try:
                workbook.close()
            except Exception:
                pass

    def _extract_zip_workbook(
        self,
        *,
        file_bytes: bytes,
        source_tag: str,
    ) -> Dict[str, Any]:
        try:
            zf = zipfile.ZipFile(io.BytesIO(file_bytes))
        except Exception:
            return {"sheets": []}

        sheets: List[Dict[str, Any]] = []
        total_nonempty_rows = 0
        total_nonempty_cells = 0
        total_source_rows = 0
        total_source_cols = 0
        try:
            shared_strings = self._xlsx_runtime.shared_strings_from_zip(zf)
            for sheet_idx, sheet_name, sheet_path in self._xlsx_runtime.sheet_entries_from_zip(zf):
                try:
                    root = ET.fromstring(zf.read(sheet_path))
                except Exception:
                    continue
                dimension_node = root.find(".//{*}dimension")
                source_max_row, source_max_col = self._xlsx_runtime.dimension_hints(
                    str(dimension_node.get("ref") or "") if dimension_node is not None else ""
                )
                sheet_data = root.find(".//{*}sheetData")
                if sheet_data is None:
                    continue

                rows_payload: List[Dict[str, Any]] = []
                used_cols: set[int] = set()
                nonempty_cells = 0
                for row_pos, row_node in enumerate(sheet_data.findall("./{*}row"), start=1):
                    try:
                        row_idx = int(row_node.get("r") or row_pos)
                    except Exception:
                        row_idx = row_pos
                    source_max_row = max(source_max_row, row_idx)
                    cells_out: List[Dict[str, Any]] = []
                    for cell_pos, cell_node in enumerate(row_node.findall("./{*}c"), start=1):
                        _, col_idx = self._xlsx_runtime.ref_to_row_col(str(cell_node.get("r") or "").strip())
                        if col_idx <= 0:
                            col_idx = cell_pos
                        source_max_col = max(source_max_col, col_idx)
                        text = normalize_text(self._xlsx_runtime.cell_text_from_xml(cell_node, shared_strings))
                        if not text:
                            continue
                        used_cols.add(col_idx)
                        nonempty_cells += 1
                        bbox = XlsxTableRuntime.build_bbox(
                            sheet_idx=sheet_idx,
                            row_idx=row_idx,
                            col_idx=col_idx,
                            max_rows=max(1, int(source_max_row or row_idx)),
                            max_cols=max(1, int(source_max_col or col_idx)),
                        )
                        cells_out.append(
                            {
                                "col_index": col_idx,
                                "text": text,
                                "ref_id": f"{source_tag}_t{sheet_idx}_r{row_idx}_c{col_idx}",
                                "bbox": bbox,
                            }
                        )
                    if cells_out:
                        rows_payload.append({"row_index": row_idx, "cells": cells_out})
                sheet_payload = self._finalize_sheet_payload(
                    sheet_index=sheet_idx,
                    sheet_name=sheet_name or f"Sheet{sheet_idx}",
                    rows_payload=rows_payload,
                    source_max_row=source_max_row,
                    source_max_col=source_max_col,
                    used_cols=used_cols,
                    nonempty_cells=nonempty_cells,
                )
                sheets.append(sheet_payload)
                total_nonempty_rows += int(sheet_payload.get("source_nonempty_row_count", 0) or 0)
                total_nonempty_cells += int(sheet_payload.get("source_nonempty_cell_count", 0) or 0)
                total_source_rows += int(sheet_payload.get("source_max_row", 0) or 0)
                total_source_cols += int(sheet_payload.get("source_max_col", 0) or 0)
        finally:
            try:
                zf.close()
            except Exception:
                pass

        return {
            "sheet_count": len(sheets),
            "source_total_rows": total_source_rows,
            "source_total_cols": total_source_cols,
            "source_total_nonempty_rows": total_nonempty_rows,
            "source_total_nonempty_cells": total_nonempty_cells,
            "sheets": sheets,
        }

    def _finalize_sheet_payload(
        self,
        *,
        sheet_index: int,
        sheet_name: str,
        rows_payload: List[Dict[str, Any]],
        source_max_row: int,
        source_max_col: int,
        used_cols: set[int],
        nonempty_cells: int,
    ) -> Dict[str, Any]:
        if not rows_payload:
            return {
                "sheet_index": sheet_index,
                "sheet_name": normalize_text(sheet_name or f"Sheet{sheet_index}") or f"Sheet{sheet_index}",
                "headers": [],
                "header_meta": [],
                "rows": [],
                "source_max_row": max(1, int(source_max_row or 1)),
                "source_max_col": max(1, int(source_max_col or 1)),
                "source_used_col_count": 0,
                "source_nonempty_row_count": 0,
                "source_nonempty_cell_count": 0,
            }

        headers, data_rows = self._table_header_runtime.infer_headers_and_data_rows(rows_payload)
        data_rows = self._table_header_runtime.drop_redundant_header_like_rows(data_rows, headers)
        compact_rows = self._build_compact_rows_for_artifact(data_rows, headers)
        if not compact_rows:
            fallback_col_count = max(
                len(headers),
                max((len(row.get("cells", [])) for row in rows_payload), default=0),
            )
            fallback_headers = [header if header else f"col_{idx}" for idx, header in enumerate(headers, start=1)]
            if len(fallback_headers) < fallback_col_count:
                fallback_headers.extend(
                    [f"col_{idx}" for idx in range(len(fallback_headers) + 1, fallback_col_count + 1)]
                )
            fallback_rows = rows_payload[1:] if len(rows_payload) > 1 else rows_payload
            headers = fallback_headers[:fallback_col_count]
            compact_rows = self._build_compact_rows_for_artifact(fallback_rows, headers)

        headers = [normalize_text(header) or f"col_{idx}" for idx, header in enumerate(headers, start=1)]
        return {
            "sheet_index": sheet_index,
            "sheet_name": normalize_text(sheet_name or f"Sheet{sheet_index}") or f"Sheet{sheet_index}",
            "headers": headers,
            "header_meta": build_header_meta(headers),
            "rows": compact_rows,
            "source_max_row": max(1, int(source_max_row or 1)),
            "source_max_col": max(1, int(source_max_col or 1)),
            "source_used_col_count": len(used_cols),
            "source_nonempty_row_count": len(rows_payload),
            "source_nonempty_cell_count": nonempty_cells,
        }

    @staticmethod
    def _build_compact_rows_for_artifact(rows_payload: List[Dict[str, Any]], headers: List[str]) -> List[List[Any]]:
        header_meta = build_header_meta(headers)
        header_map = {int(item.get("col_index", 0) or 0): item for item in header_meta}
        compact_rows: List[List[Any]] = []
        row_idx = 1
        for row in rows_payload:
            cells = list(row.get("cells") or [])
            out_cells: List[List[Any]] = []
            for cell in cells:
                col_idx = int(cell.get("col_index", 0) or 0)
                if col_idx <= 0:
                    continue
                text = normalize_text(str(cell.get("text", "") or ""))
                if not text:
                    continue
                header = header_map.get(col_idx, {"name": f"col_{col_idx}", "unit": None})
                number_meta = extract_number_meta(text, header.get("unit"))
                bbox = cell.get("bbox")
                out: List[Any] = [
                    col_idx,
                    text,
                    str(cell.get("ref_id", "") or ""),
                    list(bbox[:5]) if isinstance(bbox, list) else [],
                ]
                numeric_value = number_meta.get("numeric_value")
                if numeric_value is not None:
                    out.append(numeric_value)
                out_cells.append(out)
            if out_cells:
                source_row_index = int(row.get("row_index", row_idx) or row_idx)
                compact_rows.append([row_idx, source_row_index, out_cells])
                row_idx += 1
        return compact_rows

    def _persist_workbook_artifacts(
        self,
        *,
        file_name: str,
        file_hash: str,
        workbook_payload: Dict[str, Any],
        workbook_mode: str,
    ) -> Tuple[str, str, List[Dict[str, Any]]]:
        sheet_manifests: List[Dict[str, Any]] = []
        for sheet in workbook_payload.get("sheets", []) or []:
            sheet_index = int(sheet.get("sheet_index", 0) or 0)
            sheet_name = str(sheet.get("sheet_name", "") or f"Sheet{sheet_index}")
            rows_payload = {
                "version": _SHEET_ROWS_ARTIFACT_VERSION,
                "file_name": file_name,
                "file_hash": file_hash,
                "sheet_index": sheet_index,
                "sheet_name": sheet_name,
                "headers": list(sheet.get("headers") or []),
                "rows": list(sheet.get("rows") or []),
                "source_max_row": int(sheet.get("source_max_row", 0) or 0),
                "source_max_col": int(sheet.get("source_max_col", 0) or 0),
                "source_used_col_count": int(sheet.get("source_used_col_count", 0) or 0),
                "source_nonempty_row_count": int(sheet.get("source_nonempty_row_count", 0) or 0),
                "source_nonempty_cell_count": int(sheet.get("source_nonempty_cell_count", 0) or 0),
            }
            rows_fragment, rows_url = self._artifact_support_runtime.write_json_artifact(
                payload=rows_payload,
                file_name=f"{Path(file_name).stem}_sheet_{sheet_index}_{sheet_name}_rows",
                sub_dir=f"xlsx_workbook/{file_hash[:12]}/sheets",
            )
            sheet_manifest = {
                "sheet_index": sheet_index,
                "sheet_name": sheet_name,
                "headers": list(sheet.get("headers") or []),
                "header_meta": list(sheet.get("header_meta") or []),
                "rows_artifact_path": rows_fragment,
                "rows_artifact_url": rows_url,
                "source_max_row": int(sheet.get("source_max_row", 0) or 0),
                "source_max_col": int(sheet.get("source_max_col", 0) or 0),
                "source_used_col_count": int(sheet.get("source_used_col_count", 0) or 0),
                "source_nonempty_row_count": int(sheet.get("source_nonempty_row_count", 0) or 0),
                "source_nonempty_cell_count": int(sheet.get("source_nonempty_cell_count", 0) or 0),
            }
            sheet_manifests.append(sheet_manifest)

        workbook_manifest = {
            "version": _WORKBOOK_ARTIFACT_VERSION,
            "file_name": file_name,
            "file_hash": file_hash,
            "build_mode": workbook_mode,
            "sheet_count": int(workbook_payload.get("sheet_count", 0) or len(sheet_manifests)),
            "source_total_rows": int(workbook_payload.get("source_total_rows", 0) or 0),
            "source_total_cols": int(workbook_payload.get("source_total_cols", 0) or 0),
            "source_total_nonempty_rows": int(workbook_payload.get("source_total_nonempty_rows", 0) or 0),
            "source_total_nonempty_cells": int(workbook_payload.get("source_total_nonempty_cells", 0) or 0),
            "sheets": sheet_manifests,
        }
        manifest_fragment, manifest_url = self._artifact_support_runtime.write_json_artifact(
            payload=workbook_manifest,
            file_name=f"{Path(file_name).stem}_xlsx_workbook_manifest",
            sub_dir=f"xlsx_workbook/{file_hash[:12]}",
        )
        return manifest_fragment, manifest_url, sheet_manifests

    def _merge_table_blocks_with_sheet_manifests(
        self,
        *,
        table_blocks: List[Dict[str, Any]],
        sheet_manifests: List[Dict[str, Any]],
        source_tag: str,
        workbook_manifest_path: str,
        workbook_manifest_url: str,
    ) -> List[Dict[str, Any]]:
        existing_by_sheet: Dict[int, Dict[str, Any]] = {}
        for table in table_blocks or []:
            sheet_index = int((table or {}).get("sheet_index", 0) or 0)
            if sheet_index <= 0:
                table_id = str((table or {}).get("table_id") or "").strip()
                if table_id.startswith(f"{source_tag}_t"):
                    try:
                        sheet_index = int(table_id.split(f"{source_tag}_t", 1)[1].split("_", 1)[0])
                    except Exception:
                        sheet_index = 0
            if sheet_index > 0 and sheet_index not in existing_by_sheet:
                existing_by_sheet[sheet_index] = dict(table or {})

        merged_tables: List[Dict[str, Any]] = []
        for sheet in sheet_manifests:
            sheet_index = int(sheet.get("sheet_index", 0) or 0)
            existing = existing_by_sheet.get(sheet_index)
            if existing is None:
                merged_tables.append(
                    self._build_projection_table_block(
                        sheet_manifest=sheet,
                        source_tag=source_tag,
                        workbook_manifest_path=workbook_manifest_path,
                        workbook_manifest_url=workbook_manifest_url,
                    )
                )
                continue
            merged_tables.append(
                self._enrich_existing_table_block(
                    table=existing,
                    sheet_manifest=sheet,
                    workbook_manifest_path=workbook_manifest_path,
                    workbook_manifest_url=workbook_manifest_url,
                )
            )

        known_sheets = {int(item.get("sheet_index", 0) or 0) for item in sheet_manifests}
        for table in table_blocks or []:
            table_sheet_index = int((table or {}).get("sheet_index", 0) or 0)
            if table_sheet_index in known_sheets:
                continue
            merged_tables.append(dict(table or {}))
        return merged_tables

    def _enrich_existing_table_block(
        self,
        *,
        table: Dict[str, Any],
        sheet_manifest: Dict[str, Any],
        workbook_manifest_path: str,
        workbook_manifest_url: str,
    ) -> Dict[str, Any]:
        enriched = dict(table or {})
        projected_rows = list(enriched.get("rows") or [])
        projected_row_count = int(enriched.get("row_count", 0) or len(projected_rows))
        projected_col_count = int(
            enriched.get("col_count", 0)
            or max((len((row or {}).get("cells", [])) for row in projected_rows), default=0)
        )
        enriched.setdefault("sheet_index", int(sheet_manifest.get("sheet_index", 0) or 0))
        enriched.setdefault("sheet_name", str(sheet_manifest.get("sheet_name", "") or ""))
        enriched.setdefault("headers", list(sheet_manifest.get("headers") or []))
        enriched.setdefault("header_meta", list(sheet_manifest.get("header_meta") or []))
        enriched["projected_row_count"] = projected_row_count
        enriched["projected_col_count"] = projected_col_count
        enriched["source_headers"] = list(sheet_manifest.get("headers") or [])
        enriched["source_header_meta"] = list(sheet_manifest.get("header_meta") or [])
        enriched["source_max_row"] = int(sheet_manifest.get("source_max_row", 0) or 0)
        enriched["source_max_col"] = int(sheet_manifest.get("source_max_col", 0) or 0)
        enriched["source_used_col_count"] = int(sheet_manifest.get("source_used_col_count", 0) or 0)
        enriched["source_nonempty_row_count"] = int(sheet_manifest.get("source_nonempty_row_count", 0) or 0)
        enriched["source_nonempty_cell_count"] = int(sheet_manifest.get("source_nonempty_cell_count", 0) or 0)
        enriched["xlsx_workbook_manifest_path"] = workbook_manifest_path
        enriched["xlsx_workbook_manifest_url"] = workbook_manifest_url
        enriched["xlsx_sheet_rows_artifact_path"] = str(sheet_manifest.get("rows_artifact_path") or "")
        enriched["xlsx_sheet_rows_artifact_url"] = str(sheet_manifest.get("rows_artifact_url") or "")
        enriched["xlsx_artifact_version"] = _SHEET_ROWS_ARTIFACT_VERSION
        enriched["xlsx_projection_truncated"] = bool(
            int(sheet_manifest.get("source_nonempty_row_count", 0) or 0) > projected_row_count
            or int(sheet_manifest.get("source_used_col_count", 0) or 0) > projected_col_count
        )
        return enriched

    def _build_projection_table_block(
        self,
        *,
        sheet_manifest: Dict[str, Any],
        source_tag: str,
        workbook_manifest_path: str,
        workbook_manifest_url: str,
    ) -> Dict[str, Any]:
        rows_payload = load_xlsx_sheet_rows_artifact(sheet_manifest.get("rows_artifact_path") or "")
        full_rows = list(rows_payload.get("rows") or [])
        projected_rows = self._project_rows(full_rows)
        projected_col_count = max((len((row or {}).get("cells", [])) for row in projected_rows), default=0)
        return {
            "table_id": f"{source_tag}_t{int(sheet_manifest.get('sheet_index', 0) or 0)}",
            "source": source_tag,
            "page": int(sheet_manifest.get("sheet_index", 0) or 0),
            "sheet_index": int(sheet_manifest.get("sheet_index", 0) or 0),
            "sheet_name": str(sheet_manifest.get("sheet_name", "") or ""),
            "headers": list(sheet_manifest.get("headers") or []),
            "header_meta": list(sheet_manifest.get("header_meta") or []),
            "rows": projected_rows,
            "row_count": len(projected_rows),
            "col_count": projected_col_count,
            "projected_row_count": len(projected_rows),
            "projected_col_count": projected_col_count,
            "source_headers": list(sheet_manifest.get("headers") or []),
            "source_header_meta": list(sheet_manifest.get("header_meta") or []),
            "source_max_row": int(sheet_manifest.get("source_max_row", 0) or 0),
            "source_max_col": int(sheet_manifest.get("source_max_col", 0) or 0),
            "source_used_col_count": int(sheet_manifest.get("source_used_col_count", 0) or 0),
            "source_nonempty_row_count": int(sheet_manifest.get("source_nonempty_row_count", 0) or 0),
            "source_nonempty_cell_count": int(sheet_manifest.get("source_nonempty_cell_count", 0) or 0),
            "xlsx_workbook_manifest_path": workbook_manifest_path,
            "xlsx_workbook_manifest_url": workbook_manifest_url,
            "xlsx_sheet_rows_artifact_path": str(sheet_manifest.get("rows_artifact_path") or ""),
            "xlsx_sheet_rows_artifact_url": str(sheet_manifest.get("rows_artifact_url") or ""),
            "xlsx_artifact_version": _SHEET_ROWS_ARTIFACT_VERSION,
            "xlsx_projection_truncated": bool(
                int(sheet_manifest.get("source_nonempty_row_count", 0) or 0) > len(projected_rows)
                or int(sheet_manifest.get("source_used_col_count", 0) or 0) > projected_col_count
            ),
        }

    def _project_rows(self, rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        projected: List[Dict[str, Any]] = []
        for row in rows[: self._projection_row_limit]:
            if not isinstance(row, dict):
                continue
            cells = list(row.get("cells") or [])
            if self._projection_col_limit > 0:
                cells = [
                    cell
                    for cell in cells
                    if int((cell or {}).get("col_index", 0) or 0) <= self._projection_col_limit
                ]
            next_row = dict(row)
            next_row["cells"] = cells
            projected.append(next_row)
        return projected

    @staticmethod
    def _resolve_source_tag(table_blocks: Iterable[Dict[str, Any]]) -> str:
        for table in table_blocks or []:
            source = str((table or {}).get("source") or "").strip().lower()
            if source in {"xlsx", "xlsx_formula"}:
                return source
        return "xlsx"
