from __future__ import annotations

import io
import re
import zipfile
from typing import Any, Callable, Dict, List, Optional, Tuple
from xml.etree import ElementTree as ET


NormalizeText = Callable[[str], str]
RowsPayload = List[Dict[str, Any]]
TableBlocks = List[Dict[str, Any]]
LoadWorkbook = Callable[..., Any]
BuildIndexedRowsPayload = Callable[..., Tuple[RowsPayload, int, bool]]
BuildTableBlock = Callable[..., Optional[Dict[str, Any]]]
ConsumeOpenpyxlWorkbook = Callable[..., Tuple[TableBlocks, int]]
ConsumeZipWorkbook = Callable[..., Tuple[TableBlocks, int, bool]]
AppendXlsxSheetPlaceholders = Callable[..., int]
AppendTextRef = Callable[..., bool]


class XlsxRuntime:
    """Parse XLSX workbook metadata helpers while keeping parser glue thin."""

    def __init__(
        self,
        *,
        normalize_text: NormalizeText,
        sheet_limit: int,
    ) -> None:
        self._normalize_text = normalize_text
        self.sheet_limit = max(1, int(sheet_limit or 1))

    @staticmethod
    def col_to_index(col_ref: str) -> int:
        idx = 0
        for ch in str(col_ref or "").upper():
            if "A" <= ch <= "Z":
                idx = idx * 26 + (ord(ch) - 64)
        return idx

    @classmethod
    def ref_to_row_col(cls, cell_ref: str) -> Tuple[int, int]:
        match = re.match(r"^([A-Za-z]+)(\d+)$", str(cell_ref or "").strip())
        if not match:
            return 0, 0
        return int(match.group(2)), cls.col_to_index(match.group(1))

    @classmethod
    def dimension_hints(cls, ref_value: str) -> Tuple[int, int]:
        raw = str(ref_value or "").strip()
        if not raw:
            return 1, 1
        tail = raw.split(":")[-1]
        row_idx, col_idx = cls.ref_to_row_col(tail)
        return max(1, row_idx), max(1, col_idx)

    def shared_strings_from_zip(self, zf: zipfile.ZipFile) -> List[str]:
        try:
            root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
        except Exception:
            return []
        shared: List[str] = []
        for node in root.findall(".//{*}si"):
            fragments = [str(text_node.text or "") for text_node in node.iterfind(".//{*}t")]
            if fragments:
                shared.append(self._normalize_text("".join(fragments)))
            else:
                shared.append("")
        return shared

    def sheet_entries_from_zip(self, zf: zipfile.ZipFile) -> List[Tuple[int, str, str]]:
        rel_map: Dict[str, str] = {}
        try:
            rel_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
            for rel in rel_root.findall(".//{*}Relationship"):
                rel_id = str(rel.get("Id") or "").strip()
                rel_type = str(rel.get("Type") or "").strip().lower()
                target = str(rel.get("Target") or "").strip().replace("\\", "/")
                if not rel_id or not target or not rel_type.endswith("/worksheet"):
                    continue
                if target.startswith("/"):
                    normalized_target = target.lstrip("/")
                elif target.startswith("xl/"):
                    normalized_target = target
                else:
                    normalized_target = f"xl/{target.lstrip('./')}"
                rel_map[rel_id] = normalized_target
        except Exception:
            rel_map = {}

        entries: List[Tuple[int, str, str]] = []
        try:
            root = ET.fromstring(zf.read("xl/workbook.xml"))
            rid_keys = (
                "{http://purl.oclc.org/ooxml/officeDocument/relationships}id",
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
            )
            for sheet_idx, sheet in enumerate(root.findall(".//{*}sheet"), start=1):
                sheet_name = self._normalize_text(str(sheet.get("name") or f"Sheet{sheet_idx}"))
                rel_id = ""
                for rid_key in rid_keys:
                    rel_id = str(sheet.get(rid_key) or "").strip()
                    if rel_id:
                        break
                target = rel_map.get(rel_id)
                if target:
                    entries.append((sheet_idx, sheet_name or f"Sheet{sheet_idx}", target))
        except Exception:
            entries = []

        if entries:
            return entries[: self.sheet_limit]

        sheet_paths = sorted(
            name
            for name in zf.namelist()
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        )
        return [
            (sheet_idx, f"Sheet{sheet_idx}", sheet_path)
            for sheet_idx, sheet_path in enumerate(sheet_paths[: self.sheet_limit], start=1)
        ]

    def cell_text_from_xml(self, cell_node: ET.Element, shared_strings: List[str]) -> str:
        cell_type = str(cell_node.get("t") or "").strip().lower()
        if cell_type == "inlinestr":
            text_parts = [str(node.text or "") for node in cell_node.findall(".//{*}t")]
            return "".join(text_parts).strip()

        value_node = cell_node.find("./{*}v")
        formula_node = cell_node.find("./{*}f")
        raw_value = "".join(value_node.itertext()).strip() if value_node is not None else ""
        raw_formula = "".join(formula_node.itertext()).strip() if formula_node is not None else ""

        if cell_type == "s":
            try:
                shared_idx = int(float(raw_value))
                if 0 <= shared_idx < len(shared_strings):
                    return shared_strings[shared_idx]
            except Exception:
                return raw_value
            return raw_value
        if cell_type == "b":
            if raw_value.lower() in {"1", "true"}:
                return "TRUE"
            if raw_value.lower() in {"0", "false"}:
                return "FALSE"
            return raw_value
        if cell_type in {"str", "d", "e"}:
            return raw_value
        if raw_value:
            return raw_value
        if raw_formula:
            return f"={raw_formula}"
        return ""


class XlsxTableRuntime:
    """Consume XLSX workbook structures into parser-compatible table blocks."""

    def __init__(
        self,
        *,
        sheet_limit: int,
        table_max_rows: int,
        table_max_cols: int,
        xlsx_cell_cap: int,
        xlsx_runtime: Any,
        build_indexed_rows_payload: BuildIndexedRowsPayload,
        build_table_block: BuildTableBlock,
    ) -> None:
        self.sheet_limit = max(1, int(sheet_limit or 1))
        self.table_max_rows = max(1, int(table_max_rows or 1))
        self.table_max_cols = max(1, int(table_max_cols or 1))
        self.xlsx_cell_cap = max(1, int(xlsx_cell_cap or 1))
        self._xlsx_runtime = xlsx_runtime
        self._build_indexed_rows_payload = build_indexed_rows_payload
        self._build_table_block = build_table_block

    @staticmethod
    def build_bbox(
        *,
        sheet_idx: int,
        row_idx: int,
        col_idx: int,
        max_rows: int,
        max_cols: int,
    ) -> List[float]:
        rows = max(1, max_rows)
        cols = max(1, max_cols)
        x = min(1.0, max(0.0, (col_idx - 1) / float(cols)))
        y = min(1.0, max(0.0, (row_idx - 1) / float(rows)))
        w = min(1.0, max(0.001, 1.0 / float(cols)))
        h = min(1.0, max(0.001, 1.0 / float(rows)))
        return [float(sheet_idx), x, y, w, h]

    def consume_openpyxl_workbook(
        self,
        *,
        workbook: Any,
        source_tag: str,
        emitted_cells: int,
        parts: List[str],
        bbox_map: Dict[str, List[float]],
    ) -> Tuple[TableBlocks, int]:
        table_blocks: TableBlocks = []
        for sheet_idx, worksheet in enumerate(getattr(workbook, "worksheets", [])[: self.sheet_limit], start=1):
            max_rows_hint = min(self.table_max_rows, max(int(getattr(worksheet, "max_row", 0) or 0), 1))
            max_cols_hint = min(self.table_max_cols, max(int(getattr(worksheet, "max_column", 0) or 0), 1))
            rows_payload, emitted_cells, hit_cap = self._build_indexed_rows_payload(
                rows=worksheet.iter_rows(values_only=True),
                get_cells=lambda row: list(row),
                get_cell_text=lambda cell: "" if cell is None else str(cell),
                get_row_index=lambda row, row_pos: row_pos,
                get_col_index=lambda cell, cell_pos: cell_pos,
                build_cell_bbox=lambda row_idx, col_idx, cell: self.build_bbox(
                    sheet_idx=sheet_idx,
                    row_idx=row_idx,
                    col_idx=col_idx,
                    max_rows=max_rows_hint,
                    max_cols=max_cols_hint,
                ),
                build_cell_ref=lambda row_idx, col_idx, cell: f"{source_tag}_t{sheet_idx}_r{row_idx}_c{col_idx}",
                max_rows=self.table_max_rows,
                max_cols=self.table_max_cols,
                emitted_cells=emitted_cells,
                max_emitted_cells=self.xlsx_cell_cap,
                parts=parts,
                bbox_map=bbox_map,
                break_on_col_overflow=True,
            )
            table_block = self._build_table_block(
                table_id=f"{source_tag}_t{sheet_idx}",
                source=source_tag,
                page=sheet_idx,
                rows_payload=rows_payload,
                extra_fields={
                    "sheet_index": sheet_idx,
                    "sheet_name": str(getattr(worksheet, "title", "") or f"Sheet{sheet_idx}"),
                },
            )
            if table_block:
                table_blocks.append(table_block)
            if hit_cap:
                break
        return table_blocks, emitted_cells

    def consume_zip_workbook(
        self,
        *,
        file_bytes: bytes,
        source_tag: str,
        emitted_cells: int,
        parts: List[str],
        bbox_map: Dict[str, List[float]],
    ) -> Tuple[TableBlocks, int, bool]:
        table_blocks: TableBlocks = []
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            shared_strings = self._xlsx_runtime.shared_strings_from_zip(zf)
            sheet_entries = self._xlsx_runtime.sheet_entries_from_zip(zf)
            if not sheet_entries:
                return table_blocks, emitted_cells, False

            for sheet_idx, sheet_name, sheet_path in sheet_entries:
                if emitted_cells >= self.xlsx_cell_cap:
                    break
                try:
                    root = ET.fromstring(zf.read(sheet_path))
                except Exception:
                    continue

                dimension_node = root.find(".//{*}dimension")
                max_rows_hint, max_cols_hint = self._xlsx_runtime.dimension_hints(
                    str(dimension_node.get("ref") or "") if dimension_node is not None else ""
                )
                sheet_data = root.find(".//{*}sheetData")
                if sheet_data is None:
                    continue

                def _row_index(row_node: ET.Element, row_pos: int) -> int:
                    nonlocal max_rows_hint
                    try:
                        row_idx = int(row_node.get("r") or row_pos)
                    except Exception:
                        row_idx = row_pos
                    max_rows_hint = max(max_rows_hint, row_idx)
                    return row_idx

                def _col_index(cell_node: ET.Element, cell_pos: int) -> int:
                    return (self._xlsx_runtime.ref_to_row_col(str(cell_node.get("r") or "").strip())[1]) or cell_pos

                def _cell_bbox(row_idx: int, col_idx: int, cell_node: ET.Element) -> List[float]:
                    return self.build_bbox(
                        sheet_idx=sheet_idx,
                        row_idx=row_idx,
                        col_idx=col_idx,
                        max_rows=max_rows_hint,
                        max_cols=max(max_cols_hint, col_idx),
                    )

                def _after_cell(row_idx: int, col_idx: int, cell_node: ET.Element, text: str) -> None:
                    nonlocal max_cols_hint
                    max_cols_hint = max(max_cols_hint, col_idx)

                rows_payload, emitted_cells, hit_cap = self._build_indexed_rows_payload(
                    rows=sheet_data.findall("./{*}row"),
                    get_cells=lambda row_node: row_node.findall("./{*}c"),
                    get_cell_text=lambda cell_node: self._xlsx_runtime.cell_text_from_xml(cell_node, shared_strings),
                    get_row_index=_row_index,
                    get_col_index=_col_index,
                    build_cell_bbox=_cell_bbox,
                    build_cell_ref=lambda row_idx, col_idx, cell_node: f"{source_tag}_t{sheet_idx}_r{row_idx}_c{col_idx}",
                    max_rows=self.table_max_rows,
                    max_cols=self.table_max_cols,
                    emitted_cells=emitted_cells,
                    max_emitted_cells=self.xlsx_cell_cap,
                    parts=parts,
                    bbox_map=bbox_map,
                    after_cell=_after_cell,
                )
                table_block = self._build_table_block(
                    table_id=f"{source_tag}_t{sheet_idx}",
                    source=source_tag,
                    page=sheet_idx,
                    rows_payload=rows_payload,
                    extra_fields={
                        "sheet_index": sheet_idx,
                        "sheet_name": sheet_name or f"Sheet{sheet_idx}",
                    },
                )
                if table_block:
                    table_blocks.append(table_block)
                if hit_cap:
                    break

        return table_blocks, emitted_cells, True


class XlsxParseRuntime:
    """Coordinate XLSX workbook fallbacks while preserving parser-compatible output."""

    def __init__(
        self,
        *,
        sheet_limit: int,
        consume_openpyxl_workbook: ConsumeOpenpyxlWorkbook,
        consume_zip_workbook: ConsumeZipWorkbook,
        append_xlsx_sheet_placeholders: AppendXlsxSheetPlaceholders,
        append_text_ref: AppendTextRef,
    ) -> None:
        self.sheet_limit = max(1, int(sheet_limit or 1))
        self._consume_openpyxl_workbook = consume_openpyxl_workbook
        self._consume_zip_workbook = consume_zip_workbook
        self._append_xlsx_sheet_placeholders = append_xlsx_sheet_placeholders
        self._append_text_ref = append_text_ref

    def parse_xlsx(
        self,
        *,
        file_bytes: bytes,
        load_workbook: Optional[LoadWorkbook] = None,
    ) -> Tuple[str, Dict[str, List[float]], TableBlocks, bool]:
        workbook_loader = load_workbook or self._default_load_workbook
        parts: List[str] = []
        bbox_map: Dict[str, List[float]] = {}
        table_blocks: TableBlocks = []
        emitted_cells = 0
        used_zip_fallback = False

        try:
            wb_data = workbook_loader(io.BytesIO(file_bytes), data_only=True, read_only=True)
            if getattr(wb_data, "worksheets", None):
                workbook_tables, emitted_cells = self._consume_openpyxl_workbook(
                    workbook=wb_data,
                    source_tag="xlsx",
                    emitted_cells=emitted_cells,
                    parts=parts,
                    bbox_map=bbox_map,
                )
                if workbook_tables:
                    table_blocks.extend(workbook_tables)
        except Exception:
            pass

        if not parts:
            try:
                zip_tables, emitted_cells, used_zip_fallback = self._consume_zip_workbook(
                    file_bytes=file_bytes,
                    source_tag="xlsx",
                    emitted_cells=emitted_cells,
                    parts=parts,
                    bbox_map=bbox_map,
                )
                if zip_tables:
                    table_blocks.extend(zip_tables)
            except Exception:
                used_zip_fallback = False

        if not parts:
            try:
                wb_formula = workbook_loader(io.BytesIO(file_bytes), data_only=False, read_only=True)
                if getattr(wb_formula, "worksheets", None):
                    formula_tables, emitted_cells = self._consume_openpyxl_workbook(
                        workbook=wb_formula,
                        source_tag="xlsx_formula",
                        emitted_cells=emitted_cells,
                        parts=parts,
                        bbox_map=bbox_map,
                    )
                    if formula_tables:
                        table_blocks.extend(formula_tables)
            except Exception:
                pass

        if not parts:
            sheet_names: List[str] = []
            try:
                wb_meta = workbook_loader(io.BytesIO(file_bytes), data_only=False, read_only=True)
                sheet_names = list(getattr(wb_meta, "sheetnames", []) or [])
            except Exception:
                sheet_names = []

            emitted_sheet_placeholders = self._append_xlsx_sheet_placeholders(
                sheet_names=sheet_names,
                sheet_limit=self.sheet_limit,
                parts=parts,
                bbox_map=bbox_map,
            )
            if emitted_sheet_placeholders == 0:
                self._append_text_ref(
                    ref_id="xlsx_sheet_1",
                    text="\u5de5\u4f5c\u7c3f\u4e2d\u672a\u68c0\u6d4b\u5230\u53ef\u63d0\u53d6\u5de5\u4f5c\u8868\uff0c\u5df2\u6309\u7a7a\u7ed3\u6784\u6587\u4ef6\u5904\u7406\u3002",
                    parts=parts,
                    bbox_map=bbox_map,
                )

        if not parts:
            raise ValueError("XLSX contains no extractable rows.")
        return "\n".join(parts), bbox_map, table_blocks, used_zip_fallback

    @staticmethod
    def _default_load_workbook(stream: Any, **kwargs: Any) -> Any:
        from openpyxl import load_workbook

        return load_workbook(stream, **kwargs)
