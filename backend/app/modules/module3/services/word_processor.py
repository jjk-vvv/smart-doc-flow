from __future__ import annotations

import base64
import mimetypes
import re
import zipfile
from pathlib import Path

from docx import Document


CAPTION_PATTERN = re.compile(r"^(?:图|表)\s*[0-9０-９一二三四五六七八九十百千]+\s*[：:、.．\-—]?\s*.+")


def squish(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def looks_like_figure_caption(text: str) -> bool:
    normalized = squish(text.replace("\u3000", " "))
    return bool(CAPTION_PATTERN.match(normalized))


def find_nearby_caption(paragraph_texts: list[str], image_para_idx: int) -> str:
    for offset in range(1, 4):
        above_idx = image_para_idx - offset
        if above_idx >= 0:
            above_text = squish(paragraph_texts[above_idx])
            if above_text and looks_like_figure_caption(above_text):
                return above_text

    for offset in range(1, 4):
        below_idx = image_para_idx + offset
        if below_idx < len(paragraph_texts):
            below_text = squish(paragraph_texts[below_idx])
            if below_text and looks_like_figure_caption(below_text):
                return below_text

    for offset in range(1, 3):
        above_idx = image_para_idx - offset
        if above_idx >= 0:
            above_text = squish(paragraph_texts[above_idx])
            if above_text:
                return above_text
    return ""


def read_docx_content(word_path: Path) -> str:
    if word_path.suffix.lower() != ".docx":
        raise ValueError("Only .docx is currently supported. Please convert .doc to .docx first.")

    doc = Document(str(word_path))
    parts: list[str] = []

    paragraphs = [squish(p.text) for p in doc.paragraphs if squish(p.text)]
    if paragraphs:
        parts.append("[Body Text]\n" + "\n".join(paragraphs))

    for idx, table in enumerate(doc.tables, start=1):
        lines: list[str] = []
        for row in table.rows:
            cells = [squish(cell.text) for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
        if lines:
            parts.append(f"[Table {idx}]\n" + "\n".join(lines))

    content = "\n\n".join(parts).strip()
    if not content:
        raise ValueError("No extractable text content was found in the Word file.")
    return content


def get_word_data_url(word_path: Path) -> tuple[str, str]:
    mime = mimetypes.guess_type(word_path.name)[0]
    if not mime:
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    b64_data = base64.b64encode(word_path.read_bytes()).decode("utf-8")
    return mime, b64_data


def extract_docx_images(word_path: Path, max_images: int, max_image_bytes: int) -> list[dict[str, str]]:
    if word_path.suffix.lower() != ".docx":
        return []

    images: list[dict[str, str]] = []

    try:
        doc = Document(str(word_path))
        paragraph_texts = [p.text for p in doc.paragraphs]

        for para_idx, paragraph in enumerate(doc.paragraphs):
            caption = find_nearby_caption(paragraph_texts, para_idx)
            for run in paragraph.runs:
                rel_ids = run._element.xpath('.//*[local-name()="blip"]/@*[local-name()="embed"]')
                for rel_id in rel_ids:
                    part = doc.part.related_parts.get(rel_id)
                    if part is None:
                        continue

                    data = getattr(part, "blob", b"")
                    if not data or len(data) > max_image_bytes:
                        continue

                    mime = getattr(part, "content_type", None)
                    part_name = str(getattr(part, "partname", ""))
                    if not mime:
                        mime = mimetypes.guess_type(part_name)[0] or "application/octet-stream"
                    if not mime.startswith("image/"):
                        continue

                    name = Path(part_name).name if part_name else f"image_{len(images) + 1}.bin"
                    data_url = f"data:{mime};base64,{base64.b64encode(data).decode('utf-8')}"
                    images.append(
                        {
                            "name": name,
                            "mime": mime,
                            "data_url": data_url,
                            "caption": caption,
                        }
                    )
                    if len(images) >= max_images:
                        return images
    except Exception:
        pass

    if images:
        return images

    with zipfile.ZipFile(word_path, "r") as zf:
        media_paths = sorted(
            name
            for name in zf.namelist()
            if name.startswith("word/media/") and not name.endswith("/")
        )

        for media_path in media_paths:
            try:
                data = zf.read(media_path)
            except KeyError:
                continue

            if not data or len(data) > max_image_bytes:
                continue

            mime = mimetypes.guess_type(media_path)[0] or "application/octet-stream"
            if not mime.startswith("image/"):
                continue

            data_url = f"data:{mime};base64,{base64.b64encode(data).decode('utf-8')}"
            images.append(
                {
                    "name": Path(media_path).name,
                    "mime": mime,
                    "data_url": data_url,
                    "caption": "",
                }
            )
            if len(images) >= max_images:
                break

    return images


__all__ = [
    "CAPTION_PATTERN",
    "extract_docx_images",
    "find_nearby_caption",
    "get_word_data_url",
    "looks_like_figure_caption",
    "read_docx_content",
    "squish",
]
