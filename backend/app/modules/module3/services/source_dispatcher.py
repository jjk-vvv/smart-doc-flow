from __future__ import annotations

import json
import re
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Callable

import pandas as pd
from openpyxl import load_workbook

from ..core.config import DEFAULT_PROMPT, read_structured_metric_profile
from .llm_service import (
    build_messages,
    build_repair_messages,
    normalize_rows_output,
    parse_model_json,
    request_chat_completions,
)
from .markdown_processor import extract_rows_from_markdown_source
from .text_source_common import log_event
from .text_processor import extract_rows_from_txt_source
from .word_processor import extract_docx_images, read_docx_content, squish

_STRUCTURED_METRIC_PROFILE = read_structured_metric_profile()
_METRIC_HEADER_ALIASES = {
    key: [str(item) for item in value]
    for key, value in dict(_STRUCTURED_METRIC_PROFILE.get("header_aliases") or {}).items()
}
_ROOT_REGION_LABELS = tuple(str(item) for item in (_STRUCTURED_METRIC_PROFILE.get("root_regions") or []))
_PRIMARY_ROOT_REGION = _ROOT_REGION_LABELS[0] if _ROOT_REGION_LABELS else ""
_ROOT_REGION_SET = {item for item in _ROOT_REGION_LABELS if item}
_REGION_SUFFIXES = tuple(str(item) for item in (_STRUCTURED_METRIC_PROFILE.get("region_suffixes") or []))
_REGION_SUFFIX_PATTERN = (
    re.compile(rf"({'|'.join(re.escape(item) for item in _REGION_SUFFIXES if item)})$")
    if any(_REGION_SUFFIXES)
    else None
)
_CN_NUMERIC_MULTIPLIERS = tuple(
    (str(unit), float(multiplier))
    for unit, multiplier in (_STRUCTURED_METRIC_PROFILE.get("cn_numeric_multipliers") or [])
)
_CASE_EXCLUDE_KEYWORDS = tuple(str(item) for item in (_STRUCTURED_METRIC_PROFILE.get("case_exclude_keywords") or []))
_CASE_STRONG_KEYWORDS = tuple(str(item) for item in (_STRUCTURED_METRIC_PROFILE.get("case_strong_keywords") or []))
_CASE_WEAK_KEYWORDS = tuple(str(item) for item in (_STRUCTURED_METRIC_PROFILE.get("case_weak_keywords") or []))


def prefixed_logger(logger: Callable[[str], None] | None, prefix: str) -> Callable[[str], None] | None:
    if logger is None:
        return None

    def _wrapped(message: str) -> None:
        logger(f"{prefix}{message}")

    return _wrapped


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    text = squish(str(value))
    return "" if text.lower() in {"nan", "none", "null", "nat"} else text


def _excel_engine(path: Path) -> str:
    return "xlrd" if path.suffix.lower() == ".xls" else "openpyxl"


def _normalize_docx_excel_header(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", str(text or ""))
    normalized = normalized.replace("（", "(").replace("）", ")")
    return " ".join(normalized.split()).strip()


def _normalize_header(text: str) -> str:
    return re.sub(r"[\s_\-—–\\()（）\[\]【】:：,.，、/]+", "", str(text or "").strip()).lower()


def _tokenize_header(text: str) -> list[str]:
    raw = str(text or "").strip().lower()
    ascii_tokens = [token for token in re.split(r"[^a-z0-9]+", raw) if token]
    if ascii_tokens:
        return ascii_tokens
    chinese_tokens = re.findall(r"[\u4e00-\u9fff]+", raw)
    if chinese_tokens:
        return chinese_tokens
    normalized = _normalize_header(raw)
    return [normalized] if normalized else []


def _non_empty_columns(values: list[str]) -> list[tuple[int, str]]:
    return [(idx, value) for idx, value in enumerate(values, start=1) if value]


def _normalize_match_text(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", str(text or "")).lower()
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", normalized)


def _row_filled_count(row: dict[str, str], headers: list[str]) -> int:
    return sum(1 for header in headers if _cell_text(row.get(header, "")))


def _rows_quality(rows: list[dict[str, str]], headers: list[str]) -> tuple[int, int]:
    return len(rows), sum(_row_filled_count(row, headers) for row in rows)


def _match_header(headers: list[str], candidates: list[str]) -> str:
    candidate_keys = {_normalize_header(item) for item in candidates if _normalize_header(item)}
    for header in headers:
        header_key = _normalize_header(header)
        if header_key and header_key in candidate_keys:
            return header
    return ""


def _resolve_structured_metric_headers(headers: list[str]) -> dict[str, str]:
    aliases = _METRIC_HEADER_ALIASES
    return {
        "region": _match_header(headers, list(aliases.get("region") or [])),
        "gdp_per_capita": _match_header(headers, list(aliases.get("gdp_per_capita") or [])),
        "population": _match_header(headers, list(aliases.get("population") or [])),
        "daily_tests": _match_header(headers, list(aliases.get("daily_tests") or [])),
        "cases": _match_header(headers, list(aliases.get("cases") or [])),
    }


def _looks_like_nested_region(value: str) -> bool:
    text = _cell_text(value)
    if not text:
        return False
    if text in _ROOT_REGION_SET:
        return False
    if _REGION_SUFFIX_PATTERN is None:
        return False
    return bool(_REGION_SUFFIX_PATTERN.search(text))


def _format_number_text(value: float) -> str:
    if abs(value - round(value)) < 1e-9:
        return str(int(round(value)))
    return f"{value:.6f}".rstrip("0").rstrip(".")


def _extract_cn_numeric_text(value: str, *, sum_all_numbers: bool = False) -> str:
    text = _cell_text(value)
    if not text:
        return ""
    normalized = unicodedata.normalize("NFKC", text).replace(",", "").replace("，", "").strip()
    if not normalized:
        return ""

    number_tokens = re.findall(r"[-+]?\d+(?:\.\d+)?", normalized)
    if not number_tokens:
        return ""

    numbers: list[float] = []
    for token in number_tokens:
        try:
            numbers.append(float(token))
        except ValueError:
            continue
    if not numbers:
        return ""

    base_value = sum(numbers) if sum_all_numbers else numbers[0]
    multiplier = 1.0
    for unit, unit_multiplier in _CN_NUMERIC_MULTIPLIERS:
        if unit and unit in normalized:
            multiplier = unit_multiplier
            break
    numeric = base_value * multiplier
    return _format_number_text(numeric)


def _extract_case_count_from_doc_text(doc_text: str, anchor_text: str) -> str:
    context_text = str(doc_text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    anchor = _cell_text(anchor_text)
    if not context_text or not anchor:
        return ""

    anchor_candidates: list[str] = []

    def _add_anchor_candidate(value: str) -> None:
        text = _cell_text(value)
        if text and text not in anchor_candidates:
            anchor_candidates.append(text)

    _add_anchor_candidate(anchor)
    cleaned_anchor = re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff]+", "", anchor)
    _add_anchor_candidate(cleaned_anchor)
    for suffix in _REGION_SUFFIXES:
        if not suffix:
            continue
        if cleaned_anchor.endswith(suffix):
            _add_anchor_candidate(cleaned_anchor[: -len(suffix)])

    position = -1
    for candidate in anchor_candidates:
        if not candidate:
            continue
        hit = context_text.find(candidate)
        if hit >= 0:
            position = hit
            break
    if position < 0:
        return ""
    window = context_text[position : position + 260]

    newline_index = window.find("\n")
    if newline_index > 0:
        window = window[:newline_index]

    case_matches = re.findall(r"(\d+(?:\.\d+)?)\s*例([^，。；;\n]{0,24})", window)
    if not case_matches:
        return ""

    qualified_matches = [
        match for match in case_matches
        if any(keyword and keyword in match[1] for keyword in _CASE_STRONG_KEYWORDS)
    ]
    if qualified_matches:
        case_matches = qualified_matches

    total = 0.0
    matched = False
    for token, qualifier in case_matches:
        try:
            number_value = float(token)
        except ValueError:
            continue
        qualifier_text = _cell_text(qualifier)
        if _CASE_EXCLUDE_KEYWORDS and any(keyword and keyword in qualifier_text for keyword in _CASE_EXCLUDE_KEYWORDS):
            continue
        if _CASE_WEAK_KEYWORDS and any(keyword and keyword in qualifier_text for keyword in _CASE_WEAK_KEYWORDS):
            if _CASE_STRONG_KEYWORDS and not any(
                keyword and keyword in qualifier_text
                for keyword in _CASE_STRONG_KEYWORDS
            ):
                continue
        total += number_value
        matched = True
    if not matched:
        return ""
    return _format_number_text(total) if total >= 0 else ""




def _normalize_word_rows_for_structured_metrics(
    rows: list[dict[str, str]],
    headers: list[str],
    *,
    doc_text: str = "",
) -> list[dict[str, str]]:
    if not rows:
        return rows

    resolved_headers = _resolve_structured_metric_headers(headers)
    country_header = resolved_headers.get("region", "")
    gdp_header = resolved_headers.get("gdp_per_capita", "")
    population_header = resolved_headers.get("population", "")
    daily_tests_header = resolved_headers.get("daily_tests", "")
    cases_header = resolved_headers.get("cases", "")

    available_metric_headers = [item for item in (gdp_header, population_header, daily_tests_header, cases_header) if item]
    if not country_header or len(available_metric_headers) < 2:
        return rows

    normalized_rows: list[dict[str, str]] = []
    for row in rows:
        normalized = dict(row)
        source_region = str(normalized.get(country_header, "")) if country_header else ""
        if country_header and _PRIMARY_ROOT_REGION and _looks_like_nested_region(source_region):
            normalized[country_header] = _PRIMARY_ROOT_REGION

        if gdp_header:
            parsed = _extract_cn_numeric_text(str(normalized.get(gdp_header, "")))
            if parsed:
                normalized[gdp_header] = parsed
        if population_header:
            parsed = _extract_cn_numeric_text(str(normalized.get(population_header, "")))
            if parsed:
                normalized[population_header] = parsed
        if daily_tests_header:
            parsed = _extract_cn_numeric_text(str(normalized.get(daily_tests_header, "")))
            if parsed:
                normalized[daily_tests_header] = parsed
        if cases_header:
            parsed = _extract_cn_numeric_text(str(normalized.get(cases_header, "")), sum_all_numbers=True)
            if (not parsed or parsed == "0") and source_region:
                derived = _extract_case_count_from_doc_text(doc_text, source_region)
                if derived:
                    parsed = derived
            if parsed:
                normalized[cases_header] = parsed

        normalized_rows.append(normalized)

    return normalized_rows


def _anchor_value(row: dict[str, str], anchor_header: str) -> str:
    return _normalize_match_text(row.get(anchor_header, ""))


def _contains_digit(text: str) -> bool:
    return any(char.isdigit() for char in str(text or ""))


def _resolve_word_row_roles(headers: list[str]) -> dict[str, str]:
    anchor_header = headers[0] if headers else ""
    unit_header = next((header for header in headers if "单位" in header), headers[1] if len(headers) > 1 else "")
    value_header = next(
        (
            header
            for header in headers
            if any(token in header for token in ("产量", "数量", "数值", "总量", "值"))
            and "增长" not in header
            and "%" not in header
            and "同比" not in header
        ),
        headers[2] if len(headers) > 2 else (headers[1] if len(headers) > 1 else ""),
    )
    yoy_header = next(
        (header for header in headers if any(token in header for token in ("增长", "同比", "%", "增速", "变化"))),
        headers[-1] if headers else "",
    )
    return {
        "anchor": anchor_header,
        "unit": unit_header,
        "value": value_header,
        "yoy": yoy_header,
    }


def _find_number_token(text: str, start: int = 0) -> tuple[str, int, int]:
    index = max(start, 0)
    length = len(text)
    while index < length:
        char = text[index]
        if char.isdigit() or char in {"-", "−"}:
            break
        index += 1
    if index >= length:
        return "", -1, -1

    end = index + 1
    while end < length and (text[end].isdigit() or text[end] == "."):
        end += 1
    return text[index:end], index, end


def _extract_anchor_metrics_from_text(
    doc_text: str,
    anchor_text: str,
) -> tuple[str, str, str]:
    anchor = str(anchor_text or "").strip()
    if not anchor:
        return "", "", ""

    position = doc_text.find(anchor)
    if position < 0:
        return "", "", ""

    window = doc_text[position : position + 180]
    pivot = window.find("产量")
    parse_start = pivot + 2 if pivot >= 0 else len(anchor)

    value, _, value_end = _find_number_token(window, parse_start)
    if not value:
        return "", "", ""

    unit_start = value_end
    while unit_start < len(window) and window[unit_start] in {" ", "\t"}:
        unit_start += 1
    unit_end = unit_start
    while unit_end < len(window) and window[unit_end] not in {",", "，", ";", "；", "。", "\n", "\r"}:
        if window[unit_end] == "%":
            break
        unit_end += 1
    unit = window[unit_start:unit_end].strip()

    trend_keyword = ""
    trend_pos = len(window)
    for keyword in ("增长", "增⻓", "下降", "减少"):
        hit = window.find(keyword)
        if 0 <= hit < trend_pos:
            trend_keyword = keyword
            trend_pos = hit

    rate = ""
    if trend_keyword:
        rate, _, _ = _find_number_token(window, trend_pos + len(trend_keyword))
        if rate and trend_keyword in {"下降", "减少"} and not rate.startswith("-"):
            rate = f"-{rate}"

    return value.strip(), unit.strip(), rate.strip()


def _repair_recall_row_from_source(
    row: dict[str, str],
    *,
    doc_text: str,
    roles: dict[str, str],
) -> dict[str, str]:
    anchor_header = roles.get("anchor", "")
    value_header = roles.get("value", "")
    unit_header = roles.get("unit", "")
    yoy_header = roles.get("yoy", "")

    anchor_text = str(row.get(anchor_header, "")).strip()
    if not anchor_text:
        return row

    value = str(row.get(value_header, "")).strip() if value_header else ""
    unit = str(row.get(unit_header, "")).strip() if unit_header else ""
    yoy = str(row.get(yoy_header, "")).strip() if yoy_header else ""

    needs_repair = False
    if value_header and not _contains_digit(value):
        needs_repair = True
    if unit_header and not unit:
        needs_repair = True
    if yoy_header and yoy and not _contains_digit(yoy):
        needs_repair = True

    if not needs_repair:
        return row

    parsed_value, parsed_unit, parsed_rate = _extract_anchor_metrics_from_text(doc_text, anchor_text)
    if not any((parsed_value, parsed_unit, parsed_rate)):
        return row

    repaired = dict(row)
    if value_header and parsed_value and not _contains_digit(value):
        repaired[value_header] = parsed_value
    if unit_header and parsed_unit and not unit:
        repaired[unit_header] = parsed_unit
    if yoy_header and parsed_rate and (not yoy or not _contains_digit(yoy)):
        repaired[yoy_header] = parsed_rate
    return repaired


def _build_word_recall_messages(
    headers: list[str],
    custom_prompt: str,
    doc_text: str,
    extracted_rows: list[dict[str, str]],
) -> tuple[str, str]:
    system_prompt = (
        "You are a strict extraction completeness auditor. "
        "Your job is to output a COMPLETE row list for the target headers based on the source text. "
        "Return strict JSON only and do not include explanations."
    )
    user_prompt = (
        f"[Target Headers]\n{json.dumps(headers, ensure_ascii=False)}\n\n"
        f"[User Requirements]\n{custom_prompt}\n\n"
        f"[Current Extracted Rows]\n{json.dumps(extracted_rows, ensure_ascii=False)}\n\n"
        f"[Parsed Word Text]\n{doc_text}\n\n"
        "Return JSON in this format only:\n"
        '{\n  "rows": [\n    {"target_header_1": "value"}\n  ]\n}\n\n'
        "Rules:\n"
        "1) Keep every row key exactly equal to target headers.\n"
        "2) Keep existing correct rows and补全遗漏行.\n"
        "3) For semicolon-separated items (e.g. A; B; C), extract each item as a separate row.\n"
        "4) If a value is unavailable, return an empty string.\n"
        "5) Do not hallucinate facts not present in source text.\n"
    )
    return system_prompt, user_prompt


def _extract_rows_with_word_recall_pass(
    *,
    headers: list[str],
    prompt: str,
    doc_text: str,
    extracted_rows: list[dict[str, str]],
    base_url: str,
    api_key: str,
    model: str,
    api_mode: str,
    logger: Callable[[str], None] | None = None,
) -> list[dict[str, str]]:
    if not extracted_rows or not doc_text.strip():
        return []

    if headers and len(extracted_rows) >= 50 and all(
        _cell_text(row.get(header)) for row in extracted_rows for header in headers
    ):
        log_event(logger, "[recall skipped] high-confidence primary extraction")
        return []

    system_prompt, user_prompt = _build_word_recall_messages(headers, prompt, doc_text, extracted_rows)
    try:
        if api_mode != "chat":
            raise ValueError("Module 3 only supports chat API mode.")
        log_event(logger, "[recall] running Word completeness pass via /v1/chat/completions")
        model_text = request_chat_completions(
            base_url=base_url,
            api_key=api_key,
            model=model,
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            image_inputs=None,
        )

        parsed = parse_model_json(model_text)
        rows = normalize_rows_output(headers, parsed)
        log_event(logger, f"[recall] completeness pass returned rows={len(rows)}")
        return rows
    except Exception as exc:
        log_event(logger, f"[recall skipped] completeness pass failed: {exc}")
        return []


def _merge_rows_with_word_recall(
    *,
    headers: list[str],
    primary_rows: list[dict[str, str]],
    recall_rows: list[dict[str, str]],
    doc_text: str,
) -> list[dict[str, str]]:
    if not recall_rows or not primary_rows or not headers:
        return primary_rows

    roles = _resolve_word_row_roles(headers)
    anchor_header = roles.get("anchor", "")
    if not anchor_header:
        return primary_rows

    doc_norm = _normalize_match_text(doc_text)
    result_rows = [dict(row) for row in primary_rows]
    unit_header = roles.get("unit", "")
    allowed_unit_tokens = {
        _normalize_match_text(str(row.get(unit_header, "")).strip())
        for row in primary_rows
        if unit_header and str(row.get(unit_header, "")).strip()
    }

    def _rebuild_anchor_index() -> dict[str, int]:
        index: dict[str, int] = {}
        for row_idx, row in enumerate(result_rows):
            anchor = _anchor_value(row, anchor_header)
            if anchor and anchor not in index:
                index[anchor] = row_idx
        return index

    anchor_index = _rebuild_anchor_index()
    recall_anchor_order: list[str] = []
    recall_rows_by_anchor: dict[str, dict[str, str]] = {}

    for row in recall_rows:
        anchor = _anchor_value(row, anchor_header)
        if not anchor:
            continue
        if doc_norm and anchor not in doc_norm:
            continue
        repaired_row = _repair_recall_row_from_source(row, doc_text=doc_text, roles=roles)
        if allowed_unit_tokens and unit_header:
            unit_token = _normalize_match_text(str(repaired_row.get(unit_header, "")).strip())
            if unit_token and unit_token not in allowed_unit_tokens and anchor not in anchor_index:
                continue
        if anchor not in recall_rows_by_anchor:
            recall_anchor_order.append(anchor)
            recall_rows_by_anchor[anchor] = repaired_row
            continue
        if _row_filled_count(repaired_row, headers) > _row_filled_count(recall_rows_by_anchor[anchor], headers):
            recall_rows_by_anchor[anchor] = repaired_row

    if not recall_rows_by_anchor:
        return primary_rows

    # First, improve existing rows without changing row order.
    for anchor, row in recall_rows_by_anchor.items():
        if anchor not in anchor_index:
            continue
        idx = anchor_index[anchor]
        if _row_filled_count(row, headers) > _row_filled_count(result_rows[idx], headers):
            result_rows[idx] = row

    # Then, insert anchors that are truly missing, following recall relative order.
    for anchor in recall_anchor_order:
        if anchor in anchor_index:
            continue
        row = recall_rows_by_anchor[anchor]
        anchor_pos = recall_anchor_order.index(anchor)
        insert_at = len(result_rows)

        for prev_anchor in reversed(recall_anchor_order[:anchor_pos]):
            if prev_anchor in anchor_index:
                insert_at = anchor_index[prev_anchor] + 1
                break
        else:
            for next_anchor in recall_anchor_order[anchor_pos + 1 :]:
                if next_anchor in anchor_index:
                    insert_at = anchor_index[next_anchor]
                    break

        result_rows.insert(insert_at, row)
        anchor_index = _rebuild_anchor_index()

    return result_rows


def read_excel_source_sheets(
    source_path: Path,
    *,
    scan_rows: int = 12,
    max_cols: int = 80,
) -> list[dict[str, Any]]:
    wb = load_workbook(filename=str(source_path), read_only=True, data_only=True)
    sheets: list[dict[str, Any]] = []

    try:
        for ws in wb.worksheets:
            header_row_index = 0
            header_columns: list[tuple[int, str]] = []

            for row_index, row in enumerate(ws.iter_rows(values_only=True, max_row=scan_rows, max_col=max_cols), start=1):
                values = [_cell_text(value) for value in row]
                columns = _non_empty_columns(values)
                if not columns:
                    continue

                normalized = [_normalize_header(value) for _, value in columns]
                distinct = len({item for item in normalized if item})
                if len(columns) >= 2 and distinct >= 2:
                    header_row_index = row_index
                    header_columns = columns
                    break

                if not header_columns:
                    header_row_index = row_index
                    header_columns = columns

            if not header_columns:
                continue

            header_by_col = {col_idx: header for col_idx, header in header_columns}
            rows: list[dict[str, str]] = []
            for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if row_index <= header_row_index:
                    continue
                row_values = {col_idx: _cell_text(value) for col_idx, value in enumerate(row, start=1)}
                record = {header: row_values.get(col_idx, "") for col_idx, header in header_by_col.items()}
                if any(value for value in record.values()):
                    rows.append(record)

            sheets.append(
                {
                    "sheet_name": ws.title,
                    "header_row": header_row_index,
                    "header_columns": header_columns,
                    "headers": [header for _, header in header_columns],
                    "rows": rows,
                }
            )
    finally:
        wb.close()

    if not sheets:
        raise ValueError("Excel source does not contain any usable worksheets")

    return sheets


def _score_header_match(target: str, source: str) -> tuple[int, int, int, int, int]:
    target_norm = _normalize_header(target)
    source_norm = _normalize_header(source)
    if not target_norm or not source_norm:
        return (0, 0, 0, 0, 0)

    if target_norm == source_norm:
        return (4, 0, 0, 0, -len(source))

    if target_norm in source_norm or source_norm in target_norm:
        return (3, min(len(target_norm), len(source_norm)), 0, 0, -len(source))

    target_tokens = set(_tokenize_header(target))
    source_tokens = set(_tokenize_header(source))
    overlap = len(target_tokens & source_tokens)
    if overlap:
        coverage = int((overlap / max(1, len(target_tokens))) * 100)
        precision = int((overlap / max(1, len(source_tokens))) * 100)
        return (2, coverage, precision, -abs(len(target_tokens) - len(source_tokens)), -len(source))

    return (0, 0, 0, 0, 0)


def simple_match_excel_headers(source_headers: list[str], target_headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used_sources: set[str] = set()

    for target in target_headers:
        best_source = ""
        best_score = (0, 0, 0, 0, 0)

        for source in source_headers:
            if source in used_sources:
                continue
            score = _score_header_match(target, source)
            if score > best_score:
                best_source = source
                best_score = score

        if best_source and best_score[0] > 0:
            mapping[target] = best_source
            used_sources.add(best_source)

    return mapping


def _build_excel_header_mapping_messages(
    source_headers: list[str],
    target_headers: list[str],
    custom_prompt: str,
) -> tuple[str, str]:
    system_prompt = (
        "You align target fields to source Excel headers. "
        "Return strict JSON only. "
        "Do not extract rows. "
        "Only map each target field to one source header or an empty string."
    )
    user_prompt = (
        f"[Source Headers]\n{json.dumps(source_headers, ensure_ascii=False)}\n\n"
        f"[Target Headers]\n{json.dumps(target_headers, ensure_ascii=False)}\n\n"
        f"[User Rules]\n{custom_prompt}\n\n"
        "Return JSON in this format only:\n"
        '{\n  "mapping": {\n    "target_header_1": "source_header_or_empty"\n  }\n}\n\n'
        "Rules:\n"
        "1) Prefer exact and semantic header alignment.\n"
        "2) Only use headers from [Source Headers].\n"
        "3) If there is no confident match, return an empty string.\n"
        "4) Do not invent new headers.\n"
    )
    return system_prompt, user_prompt


def ai_match_excel_headers(
    source_headers: list[str],
    target_headers: list[str],
    *,
    prompt: str,
    settings: dict[str, Any],
    logger: Callable[[str], None] | None = None,
) -> dict[str, str]:
    if not target_headers:
        return {}

    system_prompt, user_prompt = _build_excel_header_mapping_messages(source_headers, target_headers, prompt)
    if settings["api_mode"] != "chat":
        raise ValueError("Module 3 only supports chat API mode.")
    log_event(logger, "[model] header mapping via /v1/chat/completions")
    model_text = request_chat_completions(
        base_url=settings["base_url"],
        api_key=settings["api_key"],
        model=settings["model"],
        system_prompt=system_prompt,
        user_prompt=user_prompt,
        image_inputs=None,
    )

    parsed = parse_model_json(model_text)
    if isinstance(parsed, dict) and isinstance(parsed.get("mapping"), dict):
        parsed_mapping = parsed["mapping"]
    elif isinstance(parsed, dict):
        parsed_mapping = parsed
    else:
        raise ValueError("Header mapping output is not a JSON object")

    allowed = set(source_headers)
    mapping: dict[str, str] = {}
    for target in target_headers:
        value = parsed_mapping.get(target, "")
        if value is None:
            continue
        selected = str(value).strip()
        if selected in allowed:
            mapping[target] = selected

    return mapping


def _ordered_source_headers(sheets: list[dict[str, Any]]) -> list[str]:
    ordered: list[str] = []
    seen: set[str] = set()
    for sheet in sheets:
        for header in sheet.get("headers", []):
            if header and header not in seen:
                seen.add(header)
                ordered.append(header)
    return ordered


def transform_excel_rows(
    sheets: list[dict[str, Any]],
    mapping: dict[str, str],
    target_headers: list[str],
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []

    for sheet in sheets:
        for source_row in sheet.get("rows", []):
            target_row: dict[str, str] = {}
            for target_header in target_headers:
                source_header = mapping.get(target_header, "")
                value = source_row.get(source_header, "") if source_header else ""
                target_row[target_header] = _cell_text(value)
            for source_header, value in source_row.items():
                source_key = _cell_text(source_header)
                if not source_key or source_key in target_row:
                    continue
                target_row[source_key] = _cell_text(value)
            if any(value for value in target_row.values()):
                rows.append(target_row)

    return rows


def _match_docx_headers_to_excel(word_headers: list[str], excel_columns: pd.Index) -> tuple[dict[str, Any], list[str]]:
    labels = list(excel_columns)
    if not labels:
        raise ValueError("Excel 数据源没有可用表头")

    normalized_to_label: dict[str, Any] = {}
    stripped_to_label: dict[str, Any] = {}
    for label in labels:
        stripped = str(label).strip()
        normalized = _normalize_docx_excel_header(label)
        stripped_to_label.setdefault(stripped, label)
        normalized_to_label.setdefault(normalized, label)

    matched: dict[str, Any] = {}
    unresolved: list[str] = []
    used_labels: set[Any] = set()
    for index, header in enumerate(word_headers, start=1):
        stripped = str(header or "").strip()
        if not stripped:
            raise ValueError(
                f"DOCX 模板表格第 {index} 列表头为空，请保留非空表头行。"
            )
        if stripped in stripped_to_label and stripped_to_label[stripped] not in used_labels:
            label = stripped_to_label[stripped]
            matched[stripped] = label
            used_labels.add(label)
            continue
        normalized = _normalize_docx_excel_header(stripped)
        if normalized in normalized_to_label and normalized_to_label[normalized] not in used_labels:
            label = normalized_to_label[normalized]
            matched[stripped] = label
            used_labels.add(label)
            continue

        best_label: Any | None = None
        best_score = (0, 0, 0, 0, 0)
        for label in labels:
            if label in used_labels:
                continue
            score = _score_header_match(stripped, str(label))
            if score > best_score:
                best_score = score
                best_label = label

        if best_label is not None and best_score[0] >= 2:
            matched[stripped] = best_label
            used_labels.add(best_label)
            continue
        unresolved.append(stripped)
    return matched, unresolved


def _is_safe_docx_ai_header_match(target_header: str, source_header: str) -> bool:
    target_text = re.sub(r"[_\-\s]+", "", _normalize_docx_excel_header(target_header))
    source_text = re.sub(r"[_\-\s]+", "", _normalize_docx_excel_header(source_header))
    if not target_text or not source_text:
        return False
    if target_text == source_text:
        return True
    if target_text in source_text or source_text in target_text:
        return True
    return SequenceMatcher(None, target_text, source_text).ratio() >= 0.84


def _safe_int(value: Any, fallback: int = 0) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return fallback


def _table_context_value(
    table_context: dict[str, Any] | None,
    key: str,
    default: Any = "",
) -> Any:
    if not isinstance(table_context, dict):
        return default
    value = table_context.get(key)
    if value is not None and value != "":
        return value
    table_block = table_context.get("table_block")
    if isinstance(table_block, dict):
        nested_value = table_block.get(key)
        if nested_value is not None and nested_value != "":
            return nested_value
    return default


def _extract_context_tokens_from_text(text: str) -> list[str]:
    merged_text = _cell_text(text)
    if not merged_text:
        return []

    raw_location_tokens = re.findall(r"[\u4e00-\u9fff]{2,24}(?:省|市|区|县)", merged_text)
    raw_location_tokens.extend(
        re.findall(r"(?:城市|地区|省份)\s*[：:]\s*([\u4e00-\u9fff]{2,12}(?:省|市|区|县)?)", merged_text)
    )
    date_tokens = re.findall(
        r"\d{4}[年/-]\d{1,2}[月/-]\d{1,2}(?:\s+\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?",
        merged_text,
    )
    time_tokens = re.findall(r"\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?", merged_text)

    city_tokens: list[str] = []
    location_stopwords = {
        "城市",
        "行政区",
        "市区",
        "城区",
        "地区",
        "省份",
    }
    for token in raw_location_tokens:
        normalized = _cell_text(token)
        if not normalized:
            continue
        # Keep only the trailing place-name segment to avoid phrases like "时刻德州市".
        tail_match = re.search(r"([\u4e00-\u9fff]{2,9}(?:省|市|区|县))$", normalized)
        location = tail_match.group(1) if tail_match else normalized
        if location in location_stopwords:
            continue
        if location not in city_tokens:
            city_tokens.append(location)

    tokens: list[str] = []
    for token in city_tokens + date_tokens + time_tokens:
        normalized = _cell_text(token)
        if normalized and normalized not in tokens:
            tokens.append(normalized)
    return tokens[:16]


def _extract_docx_context_tokens(prompt: str, table_context: dict[str, Any] | None) -> list[str]:
    context_text = _cell_text(_table_context_value(table_context, "context_text", ""))
    table_hint = _cell_text(_table_context_value(table_context, "title", ""))
    contextual_tokens = _extract_context_tokens_from_text("\n".join([context_text, table_hint]))
    if contextual_tokens:
        return contextual_tokens
    return _extract_context_tokens_from_text(prompt)


def _row_match_score(row: pd.Series, tokens: list[str]) -> int:
    if not tokens:
        return 0
    row_text = " ".join(filter(None, (_cell_text(value) for value in row.values.tolist())))
    if not row_text:
        return 0
    return sum(1 for token in tokens if token and token in row_text)


def _filter_docx_dataframe_by_context(
    dataframe: pd.DataFrame,
    *,
    prompt: str,
    table_context: dict[str, Any] | None,
    logger: Callable[[str], None] | None = None,
) -> pd.DataFrame:
    context_text = _cell_text(_table_context_value(table_context, "context_text", ""))
    table_hint = _cell_text(_table_context_value(table_context, "title", ""))
    context_tokens = _extract_context_tokens_from_text("\n".join([context_text, table_hint]))
    prompt_tokens = _extract_context_tokens_from_text(prompt)

    token_candidates: list[tuple[str, list[str]]] = []
    if context_tokens:
        token_candidates.append(("context", context_tokens))
    if prompt_tokens and prompt_tokens != context_tokens:
        token_candidates.append(("prompt", prompt_tokens))
    if not token_candidates:
        return dataframe

    city_stopwords = {"城市", "行政区", "市区", "城区", "地区", "省份"}
    city_column_candidates = [
        column_name
        for column_name in dataframe.columns.tolist()
        if "城市" in _cell_text(column_name)
    ]
    known_city_values: set[str] = set()
    for column_name in city_column_candidates[:3]:
        for value in dataframe[column_name].head(5000).tolist():
            city_value = _cell_text(value)
            if city_value:
                known_city_values.add(city_value)

    def _normalize_city_tokens(raw_tokens: list[str]) -> list[str]:
        normalized_tokens: list[str] = []
        for raw_token in raw_tokens:
            token_text = _cell_text(raw_token)
            if not token_text:
                continue
            matched_city_values = [
                city_name
                for city_name in known_city_values
                if city_name and city_name in token_text
            ]
            if matched_city_values:
                for city_name in sorted(matched_city_values, key=len, reverse=True):
                    if city_name not in normalized_tokens:
                        normalized_tokens.append(city_name)
                continue
            if known_city_values:
                # When we already know concrete city values from the source data,
                # skip unmatched phrase-like tokens to avoid over-broad filtering.
                continue
            if token_text in city_stopwords:
                continue
            if token_text not in normalized_tokens:
                normalized_tokens.append(token_text)
        return normalized_tokens

    for source_name, tokens in token_candidates:
        city_tokens = [token for token in tokens if re.search(r"(省|市|区|县)$", token)]
        city_tokens = _normalize_city_tokens(city_tokens)
        if city_tokens:
            specific_city_tokens = [token for token in city_tokens if token.endswith(("市", "区", "县"))]
            if specific_city_tokens:
                city_tokens = specific_city_tokens
        if source_name == "context" and len(city_tokens) > 1:
            city_tokens = [city_tokens[0]]

        matched_indices: list[int] = []
        match_scores: list[int] = []
        for idx, (_, row) in enumerate(dataframe.iterrows()):
            row_text = " ".join(filter(None, (_cell_text(value) for value in row.values.tolist())))
            if not row_text:
                continue
            if city_tokens and not any(token in row_text for token in city_tokens):
                continue
            score = _row_match_score(row, tokens)
            if city_tokens and score <= 0:
                score = 1
            if score > 0:
                matched_indices.append(idx)
                match_scores.append(score)

        if not matched_indices:
            log_event(
                logger,
                f"[context filter] source={source_name} produced zero row matches, "
                f"trying next strategy (token_count={len(tokens)}, city_token_count={len(city_tokens)})",
            )
            continue

        filtered = dataframe.iloc[matched_indices].copy()
        score_max = max(match_scores) if match_scores else 0
        score_min = min(match_scores) if match_scores else 0
        log_event(
            logger,
            "[context filter] source="
            + source_name
            + ", tokens="
            + json.dumps(tokens, ensure_ascii=False)
            + f", city_tokens={json.dumps(city_tokens, ensure_ascii=False)}, "
            + f"matched_rows={len(filtered)}, score_range={score_min}-{score_max}",
        )
        return filtered

    log_event(
        logger,
        "[context filter] no rows matched contextual or prompt tokens; fallback to full source rows",
    )
    return dataframe


def _docx_candidate_row_limit(table_context: dict[str, Any] | None) -> int:
    available_row_slots = _safe_int(_table_context_value(table_context, "available_row_slots", 0), 0)
    if available_row_slots <= 0:
        return 2000
    # Keep a generous buffer over writable slots while preventing huge row explosions.
    return min(max(available_row_slots * 8, available_row_slots + 40), 1500)


def _extract_rows_from_excel_for_docx_template(
    source_path: Path,
    headers: list[str],
    prompt: str,
    *,
    settings: dict[str, Any],
    table_context: dict[str, Any] | None,
    logger: Callable[[str], None] | None = None,
) -> list[dict[str, str]]:
    read_kwargs: dict[str, Any] = {
        "sheet_name": 0,
        "engine": _excel_engine(source_path),
        "header": 0,
    }

    log_event(logger, f"[step1] loaded Excel source: {source_path.name}")
    dataframe = pd.read_excel(source_path, **read_kwargs)
    dataframe = _filter_docx_dataframe_by_context(
        dataframe,
        prompt=prompt,
        table_context=table_context,
        logger=logger,
    )
    matched_mapping, unresolved = _match_docx_headers_to_excel(headers, dataframe.columns)
    if unresolved:
        ai_prompt = prompt
        context_text = _cell_text((table_context or {}).get("context_text", ""))
        if context_text:
            ai_prompt = f"{prompt}\n\n[Current DOCX Table Context]\n{context_text}" if prompt else f"[Current DOCX Table Context]\n{context_text}"
        ai_mapping = ai_match_excel_headers(
            [str(label) for label in dataframe.columns],
            unresolved,
            prompt=ai_prompt or DEFAULT_PROMPT,
            settings=settings,
            logger=logger,
        )
        direct_ai_mapping = {
            header: label
            for header, label in ai_mapping.items()
            if _is_safe_docx_ai_header_match(header, str(label))
        }
        inferred_ai_mapping = {
            header: label
            for header, label in ai_mapping.items()
            if header not in direct_ai_mapping
        }
        matched_mapping.update(ai_mapping)
        if inferred_ai_mapping:
            log_event(
                logger,
                "[header ai detail] 以下 DOCX 表头由 AI 联想匹配："
                + json.dumps(inferred_ai_mapping, ensure_ascii=False),
            )
        unresolved = [header for header in headers if str(header).strip() not in matched_mapping]

    if not matched_mapping:
        log_event(
            logger,
            "[table skipped] 当前源文件与当前 DOCX 表格没有可匹配表头，已跳过。"
            + json.dumps(
                {
                    "target_headers": headers,
                    "excel_headers": [str(label) for label in dataframe.columns],
                },
                ensure_ascii=False,
            ),
        )
        return []

    matched_headers = [header for header in headers if str(header).strip() in matched_mapping]
    rows: list[dict[str, str]] = []
    invalid_row_count = 0
    for _, row in dataframe.iterrows():
        mapped_row = {header: "" for header in headers}
        for header in matched_headers:
            label = matched_mapping[str(header).strip()]
            mapped_row[header] = _cell_text(row[label])
        for source_header in dataframe.columns:
            source_key = _cell_text(source_header)
            if not source_key or source_key in mapped_row:
                continue
            mapped_row[source_key] = _cell_text(row[source_header])
        filled_matched_headers = [header for header in matched_headers if mapped_row.get(header)]
        if not filled_matched_headers:
            continue
        if len(matched_headers) >= 3 and len(filled_matched_headers) == 1:
            invalid_row_count += 1
            continue
        if any(mapped_row.values()):
            rows.append(mapped_row)

    if not rows:
        log_event(logger, "[table skipped] 当前源文件在已匹配列下没有可写入当前 DOCX 表格的有效行，已跳过。")
        return []

    log_event(logger, "[step2] DOCX table headers matched to Excel columns")
    for header in matched_headers:
        label = matched_mapping[str(header).strip()]
        log_event(logger, f"[header match] {header} -> {label}")
    non_exact_mapping = {
        header: matched_mapping[str(header).strip()]
        for header in matched_headers
        if str(header).strip() != str(matched_mapping[str(header).strip()]).strip()
    }
    if non_exact_mapping:
        log_event(
            logger,
            "[header inferred] 以下 DOCX 表头不是完全同名匹配，请用户确认："
            + json.dumps(non_exact_mapping, ensure_ascii=False),
        )
    if unresolved:
        log_event(
            logger,
            "[header unresolved] 以下 DOCX 表头未匹配到 Excel 列，将暂时按空值处理："
            + json.dumps(unresolved, ensure_ascii=False),
        )
    if invalid_row_count:
        log_event(logger, f"[step2.1] filtered invalid Excel rows={invalid_row_count}")
    row_limit = _docx_candidate_row_limit(table_context)
    if len(rows) > row_limit:
        log_event(
            logger,
            f"[step2.2] candidate rows capped from {len(rows)} to {row_limit} "
            f"(available_row_slots={_safe_int(_table_context_value(table_context, 'available_row_slots', 0), 0)})",
        )
        rows = rows[:row_limit]
    log_event(logger, f"[step3] DOCX-oriented Excel transform complete, rows={len(rows)}")
    return rows


def extract_rows_from_excel_with_prompt(
    source_path: Path,
    headers: list[str],
    prompt: str,
    *,
    settings: dict[str, Any],
    table_context: dict[str, Any] | None = None,
    logger: Callable[[str], None] | None = None,
) -> list[dict[str, str]]:
    if table_context and str(table_context.get("template_type") or "").lower() == "docx":
        return _extract_rows_from_excel_for_docx_template(
            source_path=source_path,
            headers=headers,
            prompt=prompt,
            settings=settings,
            table_context=table_context,
            logger=logger,
        )

    log_event(logger, f"[step1] loaded Excel source: {source_path.name}")
    log_event(logger, f"[step1.1] extraction guidance received: {prompt}")

    sheets = read_excel_source_sheets(source_path)
    total_rows = sum(len(sheet.get("rows", [])) for sheet in sheets)
    source_headers = _ordered_source_headers(sheets)
    log_event(
        logger,
        f"[step1.2] worksheets={len(sheets)} source_headers={len(source_headers)} source_rows={total_rows}",
    )

    mapping = simple_match_excel_headers(source_headers, headers)
    unresolved = [header for header in headers if header not in mapping]
    log_event(logger, f"[step2] direct header match resolved {len(mapping)}/{len(headers)} fields")

    if unresolved:
        ai_mapping = ai_match_excel_headers(
            source_headers,
            unresolved,
            prompt=prompt,
            settings=settings,
            logger=logger,
        )
        mapping.update(ai_mapping)
        unresolved = [header for header in headers if header not in mapping]
        log_event(logger, f"[step2.1] AI header match resolved {len(ai_mapping)} additional fields")

    if unresolved:
        log_event(logger, f"[step2.2] unresolved target fields: {json.dumps(unresolved, ensure_ascii=False)}")

    rows = transform_excel_rows(sheets, mapping, headers)
    if not rows:
        raise ValueError("No writable row data extracted from Excel source")

    log_event(logger, f"[step3] structured Excel transform complete, rows={len(rows)}")
    log_event(logger, "[mapped rows]\n" + json.dumps({"rows": rows[:10]}, ensure_ascii=False, indent=2))
    return rows


def extract_rows_from_word_source(
    word_path: Path,
    headers: list[str],
    prompt: str,
    *,
    base_url: str,
    api_key: str,
    model: str,
    api_mode: str,
    max_images: int,
    max_image_bytes: int,
    vision_model: str | None = None,
    logger: Callable[[str], None] | None = None,
) -> list[dict[str, str]]:
    if word_path.suffix.lower() != ".docx":
        raise ValueError("Word source files must use the .docx extension")

    image_inputs = extract_docx_images(word_path, max_images=max_images, max_image_bytes=max_image_bytes)
    effective_vision_model = str(vision_model or model or "").strip() or model
    try:
        doc_text = read_docx_content(word_path)
    except ValueError as exc:
        if image_inputs:
            doc_text = ""
            log_event(logger, f"[step1] Word 正文为空，改为仅基于图片继续抽取。原因: {exc}")
        else:
            raise
    system_prompt, user_prompt = build_messages(headers, prompt, doc_text, image_count=len(image_inputs))

    log_event(logger, f"[step1] loaded Word source: {word_path.name}")
    log_event(logger, f"[step1.1] target headers={len(headers)}")
    log_event(
        logger,
        f"[step1.2] extracted {len(image_inputs)} images (max_images={max_images}, max_image_bytes={max_image_bytes})",
    )
    log_event(logger, "[step1.3] sending Word content to model")

    if api_mode != "chat":
        raise ValueError("Module 3 only supports chat API mode.")
    try:
        log_event(logger, "[model] /v1/chat/completions with images")
        model_text = request_chat_completions(
            base_url=base_url,
            api_key=api_key,
            model=effective_vision_model if image_inputs else model,
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            image_inputs=image_inputs,
        )
    except Exception as exc:
        if not image_inputs:
            raise
        log_event(logger, f"[model fallback] chat image input failed, retrying text-only. reason: {exc}")
        log_event(logger, "[model] /v1/chat/completions text-only fallback")
        model_text = request_chat_completions(
            base_url=base_url,
            api_key=api_key,
            model=model,
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            image_inputs=None,
        )

    log_event(logger, "[model output]\n" + model_text)

    try:
        parsed_data = parse_model_json(model_text)
    except Exception as exc:
        log_event(logger, f"[json repair] initial parse failed, retrying with repair prompt. reason: {exc}")
        repair_system_prompt, repair_user_prompt = build_repair_messages(headers, model_text)
        if api_mode != "chat":
            raise ValueError("Module 3 only supports chat API mode.")
        log_event(logger, "[model] JSON repair via /v1/chat/completions")
        repair_text = request_chat_completions(
            base_url=base_url,
            api_key=api_key,
            model=model,
            system_prompt=repair_system_prompt,
            user_prompt=repair_user_prompt,
            image_inputs=None,
        )
        log_event(logger, "[repair output]\n" + repair_text)
        parsed_data = parse_model_json(repair_text)

    rows_values = normalize_rows_output(headers, parsed_data)
    recall_rows = _extract_rows_with_word_recall_pass(
        headers=headers,
        prompt=prompt,
        doc_text=doc_text,
        extracted_rows=rows_values,
        base_url=base_url,
        api_key=api_key,
        model=model,
        api_mode=api_mode,
        logger=logger,
    )
    merged_rows = _merge_rows_with_word_recall(
        headers=headers,
        primary_rows=rows_values,
        recall_rows=recall_rows,
        doc_text=doc_text,
    )
    if _rows_quality(merged_rows, headers) > _rows_quality(rows_values, headers):
        rows_values = merged_rows
        log_event(logger, f"[recall] merged rows improved to {len(rows_values)}")
    rows_values = _normalize_word_rows_for_structured_metrics(rows_values, headers, doc_text=doc_text)
    log_event(logger, f"[step2] normalized Word rows={len(rows_values)}")
    log_event(logger, "[normalized rows]\n" + json.dumps({"rows": rows_values}, ensure_ascii=False, indent=2))
    return rows_values


def extract_rows_from_source(
    source_path: Path,
    headers: list[str],
    prompt: str,
    *,
    settings: dict[str, Any],
    table_context: dict[str, Any] | None = None,
    logger: Callable[[str], None] | None = None,
) -> list[dict[str, str]]:
    suffix = source_path.suffix.lower()
    if suffix == ".docx":
        return extract_rows_from_word_source(
            word_path=source_path,
            headers=headers,
            prompt=prompt,
            base_url=settings["base_url"],
            api_key=settings["api_key"],
            model=settings["model"],
            vision_model=settings.get("vision_model", settings["model"]),
            api_mode=settings["api_mode"],
            max_images=settings["max_images"],
            max_image_bytes=settings["max_image_bytes"],
            logger=logger,
        )
    if suffix == ".xlsx":
        normalized_prompt = prompt.strip() or DEFAULT_PROMPT
        return extract_rows_from_excel_with_prompt(
            source_path=source_path,
            headers=headers,
            prompt=normalized_prompt,
            settings=settings,
            table_context=table_context,
            logger=logger,
        )
    if suffix == ".md":
        return extract_rows_from_markdown_source(
            markdown_path=source_path,
            headers=headers,
            prompt=prompt,
            base_url=settings["base_url"],
            api_key=settings["api_key"],
            model=settings["model"],
            api_mode=settings["api_mode"],
            logger=logger,
        )
    if suffix == ".txt":
        return extract_rows_from_txt_source(
            txt_path=source_path,
            headers=headers,
            prompt=prompt,
            base_url=settings["base_url"],
            api_key=settings["api_key"],
            model=settings["model"],
            api_mode=settings["api_mode"],
            logger=logger,
        )
    raise ValueError(f"Unsupported source file type: {source_path.suffix}")


__all__ = [
    "ai_match_excel_headers",
    "extract_rows_from_excel_with_prompt",
    "extract_rows_from_source",
    "extract_rows_from_word_source",
    "log_event",
    "prefixed_logger",
    "read_excel_source_sheets",
    "simple_match_excel_headers",
    "transform_excel_rows",
]
