from __future__ import annotations

from collections import deque
from copy import deepcopy
from pathlib import Path
import re
from typing import Any, Callable

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..core.config import DEFAULT_HEADER_ROW
from .storage_service import normalize_owner_id, utcnow_iso


def _squish(text: str) -> str:
    return " ".join(text.split()).strip()


def _header_columns_to_pairs(header_columns: list[Any]) -> list[tuple[int, str]]:
    pairs: list[tuple[int, str]] = []
    for item in header_columns:
        if isinstance(item, dict):
            col_idx = int(item.get("col_idx", 0) or 0)
            header = _squish(str(item.get("header", "")))
        elif isinstance(item, (list, tuple)) and len(item) >= 2:
            col_idx = int(item[0] or 0)
            header = _squish(str(item[1]))
        else:
            continue
        if col_idx > 0 and header:
            pairs.append((col_idx, header))
    return pairs


def _pairs_to_header_columns(header_columns: list[tuple[int, str]]) -> list[dict[str, Any]]:
    return [{"col_idx": col_idx, "header": header} for col_idx, header in header_columns if header]


def _excel_row_values(ws: Any, row_index: int) -> list[str]:
    values: list[str] = []
    merged_ranges = list(ws.merged_cells.ranges)
    for col_index in range(1, ws.max_column + 1):
        value = ws.cell(row=row_index, column=col_index).value
        if value is None:
            for merged_range in merged_ranges:
                if (
                    merged_range.min_row <= row_index <= merged_range.max_row
                    and merged_range.min_col <= col_index <= merged_range.max_col
                ):
                    value = ws.cell(row=merged_range.min_row, column=merged_range.min_col).value
                    break
        values.append(_squish(str(value)) if value is not None else "")
    return values


def _looks_like_multirow_excel_header(ws: Any, header_row: int) -> bool:
    if header_row + 1 > ws.max_row:
        return False

    first_row_values = _excel_row_values(ws, header_row)
    second_row_values = _excel_row_values(ws, header_row + 1)
    first_non_empty = [value for value in first_row_values if value]
    second_non_empty = [value for value in second_row_values if value]

    if len(second_non_empty) < 2:
        return False

    merged_first_row = any(
        merged_range.min_row <= header_row <= merged_range.max_row and merged_range.max_col > merged_range.min_col
        for merged_range in ws.merged_cells.ranges
    )
    sparse_first_row = 0 < len(first_non_empty) < len(second_non_empty)

    if not merged_first_row and not sparse_first_row:
        return False

    return any(first_row_values[col_index - 1] for col_index, value in enumerate(second_row_values, start=1) if value)


def _combine_excel_header_rows(ws: Any, header_row: int) -> list[tuple[int, str]]:
    first_row_values = _excel_row_values(ws, header_row)
    second_row_values = _excel_row_values(ws, header_row + 1)

    headers: list[tuple[int, str]] = []
    for col_index, second_value in enumerate(second_row_values, start=1):
        if not second_value:
            continue
        first_value = first_row_values[col_index - 1]
        if first_value and first_value != second_value:
            header = _squish(f"{first_value} {second_value}")
        else:
            header = second_value or first_value
        if header:
            headers.append((col_index, header))
    return headers


def _iter_docx_blocks(doc: DocxDocument):
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield "paragraph", Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield "table", Table(child, doc)


def _row_values(row: Any) -> list[str]:
    return [_squish(cell.text) for cell in row.cells]


def _non_empty_columns(values: list[str]) -> list[tuple[int, str]]:
    return [(idx, value) for idx, value in enumerate(values, start=1) if value]


def _looks_like_title_row(values: list[str]) -> bool:
    non_empty = [value for value in values if value]
    if len(non_empty) <= 1:
        return True
    return len(set(non_empty)) == 1


def _find_header_row(table: Table) -> tuple[int, list[tuple[int, str]]]:
    fallback: tuple[int, list[tuple[int, str]]] | None = None
    scan_limit = min(len(table.rows), 8)

    for row_index, row in enumerate(table.rows[:scan_limit], start=1):
        values = _row_values(row)
        columns = _non_empty_columns(values)
        if not columns:
            continue
        if fallback is None:
            fallback = (row_index, columns)
        if len(columns) >= 2 and not _looks_like_title_row(values):
            return row_index, columns

    if fallback is not None:
        return fallback

    raise ValueError("No usable table header row found in Word template")


def _detect_key_value_targets(table: Table) -> list[dict[str, Any]] | None:
    if not table.rows:
        return None

    first_row_values = _row_values(table.rows[0])
    first_row_columns = _non_empty_columns(first_row_values)
    if len(first_row_columns) >= 2 and not _looks_like_title_row(first_row_values):
        return None

    targets: list[dict[str, Any]] = []
    for row_index, row in enumerate(table.rows, start=1):
        values = _row_values(row)
        if len(values) < 2:
            continue
        label = values[0]
        if not label:
            continue
        if sum(1 for value in values[1:] if value) > 1:
            return None
        targets.append(
            {
                "header": label,
                "row_index": row_index,
                "label_col": 1,
                "value_col": 2,
            }
        )

    if len(targets) >= 2:
        return targets
    return None


def _build_docx_table_block(table_index: int, table: Table, context_text: str) -> dict[str, Any]:
    key_value_targets = _detect_key_value_targets(table)
    block_id = f"table_{table_index + 1}"
    template_row_count = len(table.rows)

    if key_value_targets:
        header_columns = [{"col_idx": target["value_col"], "header": target["header"]} for target in key_value_targets]
        return {
            "block_id": block_id,
            "table_index": table_index,
            "context_text": context_text,
            "fill_mode": "key_value",
            "template_row_count": template_row_count,
            "available_row_slots": len(key_value_targets),
            "header_row": None,
            "header_rows": [],
            "write_row": None,
            "data_start_row": None,
            "header_columns": header_columns,
            "cell_targets": key_value_targets,
        }

    header_row, header_pairs = _find_header_row(table)
    return {
        "block_id": block_id,
        "table_index": table_index,
        "context_text": context_text,
        "fill_mode": "detail_rows",
        "template_row_count": template_row_count,
        "available_row_slots": max(template_row_count - header_row, 0),
        "header_row": header_row,
        "header_rows": [header_row],
        "write_row": header_row + 1,
        "data_start_row": header_row + 1,
        "header_columns": _pairs_to_header_columns(header_pairs),
        "cell_targets": [],
    }


def read_excel_headers(template_path: Path, sheet_name: str | None, header_row: int) -> tuple[str, list[tuple[int, str]], int]:
    wb = load_workbook(filename=str(template_path))
    ws = wb[sheet_name] if sheet_name else wb.active

    if _looks_like_multirow_excel_header(ws, header_row):
        headers = _combine_excel_header_rows(ws, header_row)
        write_row = header_row + 2
    else:
        headers = []
        for cell in ws[header_row]:
            if cell.value is None:
                continue
            header = str(cell.value).strip()
            if header:
                headers.append((cell.column, header))
        write_row = header_row + 1

    wb.close()

    if not headers:
        raise ValueError(f"No headers found on row {header_row} in Excel template")

    return ws.title, headers, write_row


def read_docx_tables(template_path: Path) -> list[dict[str, Any]]:
    doc = Document(str(template_path))
    blocks: list[dict[str, Any]] = []
    context_window: deque[str] = deque(maxlen=3)
    table_position = 0

    for block_type, block in _iter_docx_blocks(doc):
        if block_type == "paragraph":
            text = _squish(block.text)
            if text:
                context_window.append(text)
            continue

        context_text = "\n".join(context_window).strip()
        blocks.append(_build_docx_table_block(table_position, block, context_text))
        table_position += 1
        context_window.clear()

    if not blocks:
        raise ValueError("No tables found in Word template")

    return blocks


def read_docx_headers(template_path: Path) -> tuple[int, int, list[tuple[int, str]]]:
    first_table = read_docx_tables(template_path)[0]
    header_pairs = _header_columns_to_pairs(first_table.get("header_columns", []))
    if not header_pairs:
        raise ValueError("No usable fields found in first Word template table")
    header_row = int(first_table.get("header_row") or 1)
    return int(first_table["table_index"]), header_row, header_pairs


def read_template_headers(template_path: Path, sheet_name: str | None, header_row: int) -> dict[str, Any]:
    suffix = template_path.suffix.lower()
    if suffix == ".xlsx":
        resolved_sheet_name, header_columns, write_row = read_excel_headers(template_path, sheet_name, header_row)
        return {
            "template_type": "xlsx",
            "sheet_name": resolved_sheet_name,
            "header_row": header_row,
            "write_row": write_row,
            "table_index": None,
            "header_columns": header_columns,
            "tables": [],
        }
    if suffix == ".docx":
        tables = read_docx_tables(template_path)
        first_table = tables[0]
        header_columns = _header_columns_to_pairs(first_table.get("header_columns", []))
        return {
            "template_type": "docx",
            "sheet_name": f"table_{int(first_table['table_index']) + 1}",
            "header_row": first_table.get("header_row"),
            "write_row": first_table.get("write_row"),
            "table_index": int(first_table["table_index"]),
            "header_columns": header_columns,
            "tables": tables,
        }
    raise ValueError(f"Unsupported template format: {template_path.suffix}")


def build_template_fields(
    sheet_name: str,
    header_columns: list[tuple[int, str]],
    *,
    header_row: int,
    write_row: int,
    template_type: str = "xlsx",
    table_index: int | None = None,
) -> list[dict[str, Any]]:
    fields: list[dict[str, Any]] = []
    for col_idx, header in header_columns:
        if template_type == "xlsx":
            cell = f"{get_column_letter(col_idx)}{header_row}"
            write_cell = f"{get_column_letter(col_idx)}{write_row}"
        else:
            cell = f"table[{(table_index or 0) + 1}].row[{header_row}].col[{col_idx}]"
            write_cell = f"table[{(table_index or 0) + 1}].row[{write_row}].col[{col_idx}]"
        fields.append(
            {
                "key": header,
                "header": header,
                "sheet": sheet_name if template_type == "xlsx" else None,
                "cell": cell,
                "write_cell": write_cell,
                "unit": None,
                "template_type": template_type,
                "table_index": table_index,
                "col_idx": col_idx,
                "write_row": write_row,
                "aliases": [header],
            }
        )
    return fields


def build_docx_template_fields(tables: list[dict[str, Any]]) -> list[dict[str, Any]]:
    fields: list[dict[str, Any]] = []
    for table in tables:
        table_index = int(table["table_index"])
        fill_mode = str(table.get("fill_mode") or "detail_rows")
        if fill_mode == "key_value":
            for target in table.get("cell_targets", []):
                header = _squish(str(target.get("header", "")))
                if not header:
                    continue
                fields.append(
                    {
                        "key": header,
                        "header": header,
                        "sheet": None,
                        "cell": f"table[{table_index + 1}].row[{target['row_index']}].col[{target['label_col']}]",
                        "write_cell": f"table[{table_index + 1}].row[{target['row_index']}].col[{target['value_col']}]",
                        "unit": None,
                        "template_type": "docx",
                        "table_index": table_index,
                        "col_idx": int(target["value_col"]),
                        "write_row": int(target["row_index"]),
                        "aliases": [header],
                        "fill_mode": fill_mode,
                        "block_id": table.get("block_id"),
                    }
                )
            continue

        header_pairs = _header_columns_to_pairs(table.get("header_columns", []))
        header_row = int(table.get("header_row") or 1)
        write_row = int(table.get("write_row") or header_row + 1)
        block_fields = build_template_fields(
            f"table_{table_index + 1}",
            header_pairs,
            header_row=header_row,
            write_row=write_row,
            template_type="docx",
            table_index=table_index,
        )
        for item in block_fields:
            item["fill_mode"] = fill_mode
            item["block_id"] = table.get("block_id")
        fields.extend(block_fields)
    return fields


def build_default_output_path(template_path: Path) -> Path:
    candidate = template_path.with_name(f"{template_path.stem}_filled{template_path.suffix}")
    if not candidate.exists():
        return candidate

    index = 2
    while True:
        candidate = template_path.with_name(f"{template_path.stem}_filled_{index}{template_path.suffix}")
        if not candidate.exists():
            return candidate
        index += 1


_INTEGER_TEXT_RE = re.compile(r"^[+-]?\d+$")
_DECIMAL_TEXT_RE = re.compile(r"^[+-]?(?:\d+\.\d*|\.\d+)(?:[eE][+-]?\d+)?$")


def _coerce_excel_cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value

    text = str(value).strip()
    if not text:
        return ""

    normalized = text.replace(",", "")
    if "%" in normalized:
        return text

    signless = normalized.lstrip("+-")
    if _INTEGER_TEXT_RE.fullmatch(normalized):
        # Preserve identifiers such as 00123 rather than forcing Excel to trim leading zeros.
        if len(signless) > 1 and signless.startswith("0"):
            return text
        if len(signless) > 15:
            return text
        return int(normalized)

    if _DECIMAL_TEXT_RE.fullmatch(normalized):
        try:
            return float(normalized)
        except ValueError:
            return text

    return text


def write_filled_excel(
    template_path: Path,
    output_path: Path,
    sheet_name: str,
    header_columns: list[tuple[int, str]],
    rows_values: list[dict[str, str]],
    write_row: int,
    cancel_callback: Callable[[], None] | None = None,
) -> None:
    wb = load_workbook(filename=str(template_path))
    ws = wb[sheet_name]

    for row_offset, values in enumerate(rows_values):
        if cancel_callback is not None:
            cancel_callback()
        row_idx = write_row + row_offset
        for col_idx, header in header_columns:
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=_coerce_excel_cell_value(values.get(header, "")),
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    wb.close()


def _write_docx_detail_rows(
    table: Table,
    block: dict[str, Any],
    rows_values: list[dict[str, str]],
    cancel_callback: Callable[[], None] | None = None,
) -> None:
    if not rows_values:
        return

    header_columns = _header_columns_to_pairs(block.get("header_columns", []))
    write_row = int(block.get("data_start_row") or block.get("write_row") or 1)
    target_row_index = max(0, write_row - 1)
    if target_row_index >= len(table.rows):
        raise ValueError("DOCX template table requires a sample data row beneath the header row")

    def _row_matches_header(row_values: list[str]) -> bool:
        if not header_columns:
            return False
        matched = 0
        for col_idx, header in header_columns:
            cell_idx = col_idx - 1
            if cell_idx >= len(row_values):
                continue
            if _squish(row_values[cell_idx]) == _squish(header):
                matched += 1
        return matched >= max(1, len(header_columns) // 2)

    def _is_writable_detail_row(row_index: int) -> bool:
        values = _row_values(table.rows[row_index])
        if _row_matches_header(values):
            return False
        if any(values) and _looks_like_title_row(values):
            return False
        return True

    writable_row_indexes: list[int] = []
    for row_index in range(target_row_index, len(table.rows)):
        if _is_writable_detail_row(row_index):
            writable_row_indexes.append(row_index)

    sample_row_index = target_row_index
    while len(writable_row_indexes) < len(rows_values):
        if cancel_callback is not None:
            cancel_callback()
        new_tr = deepcopy(table.rows[sample_row_index]._tr)
        table.rows[-1]._tr.addnext(new_tr)
        writable_row_indexes.append(len(table.rows) - 1)

    for row_offset, values in enumerate(rows_values):
        if cancel_callback is not None:
            cancel_callback()
        row_cells = table.rows[writable_row_indexes[row_offset]].cells
        for col_idx, header in header_columns:
            if col_idx - 1 >= len(row_cells):
                continue
            row_cells[col_idx - 1].text = values.get(header, "")


def _write_docx_key_value(
    table: Table,
    block: dict[str, Any],
    rows_values: list[dict[str, str]],
    cancel_callback: Callable[[], None] | None = None,
) -> None:
    merged_values: dict[str, str] = {}
    for row in rows_values:
        if cancel_callback is not None:
            cancel_callback()
        for header, value in row.items():
            if header not in merged_values and _squish(str(value)):
                merged_values[header] = str(value).strip()

    for target in block.get("cell_targets", []):
        if cancel_callback is not None:
            cancel_callback()
        header = str(target.get("header", ""))
        row_index = int(target.get("row_index", 0))
        value_col = int(target.get("value_col", 0))
        if row_index <= 0 or value_col <= 0:
            continue
        while len(table.rows) < row_index:
            if cancel_callback is not None:
                cancel_callback()
            table.add_row()
        row_cells = table.rows[row_index - 1].cells
        if value_col - 1 >= len(row_cells):
            continue
        row_cells[value_col - 1].text = merged_values.get(header, "")


def write_filled_docx_blocks(
    template_path: Path,
    output_path: Path,
    tables: list[dict[str, Any]],
    rows_by_block: dict[str, list[dict[str, str]]],
    cancel_callback: Callable[[], None] | None = None,
) -> None:
    doc = Document(str(template_path))

    for block in tables:
        if cancel_callback is not None:
            cancel_callback()
        table_index = int(block["table_index"])
        if table_index < 0 or table_index >= len(doc.tables):
            raise ValueError(f"Word template table index out of range: {table_index}")
        table = doc.tables[table_index]
        block_rows = rows_by_block.get(str(block.get("block_id")), [])
        fill_mode = str(block.get("fill_mode") or "detail_rows")

        if fill_mode == "key_value":
            _write_docx_key_value(table, block, block_rows, cancel_callback=cancel_callback)
            continue

        _write_docx_detail_rows(table, block, block_rows, cancel_callback=cancel_callback)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if cancel_callback is not None:
        cancel_callback()
    doc.save(str(output_path))


def write_filled_docx(
    template_path: Path,
    output_path: Path,
    table_index: int,
    header_columns: list[tuple[int, str]],
    rows_values: list[dict[str, str]],
    write_row: int,
    cancel_callback: Callable[[], None] | None = None,
) -> None:
    block = {
        "block_id": f"table_{table_index + 1}",
        "table_index": table_index,
        "fill_mode": "detail_rows",
        "data_start_row": write_row,
        "write_row": write_row,
        "header_columns": _pairs_to_header_columns(header_columns),
    }
    write_filled_docx_blocks(
        template_path,
        output_path,
        [block],
        {block["block_id"]: rows_values},
        cancel_callback=cancel_callback,
    )


def write_filled_template(
    template_path: Path,
    output_path: Path,
    *,
    template_type: str,
    sheet_name: str | None,
    table_index: int | None,
    header_columns: list[tuple[int, str]],
    rows_values: list[dict[str, str]],
    write_row: int,
    docx_tables: list[dict[str, Any]] | None = None,
    docx_rows_by_block: dict[str, list[dict[str, str]]] | None = None,
    cancel_callback: Callable[[], None] | None = None,
) -> None:
    if template_type == "xlsx":
        if not sheet_name:
            raise ValueError("Excel template is missing sheet information")
        write_filled_excel(
            template_path=template_path,
            output_path=output_path,
            sheet_name=sheet_name,
            header_columns=header_columns,
            rows_values=rows_values,
            write_row=write_row,
            cancel_callback=cancel_callback,
        )
        return

    if template_type == "docx":
        if docx_tables is not None and docx_rows_by_block is not None:
            write_filled_docx_blocks(
                template_path=template_path,
                output_path=output_path,
                tables=docx_tables,
                rows_by_block=docx_rows_by_block,
                cancel_callback=cancel_callback,
            )
            return
        if table_index is None:
            raise ValueError("Word template is missing table information")
        write_filled_docx(
            template_path=template_path,
            output_path=output_path,
            table_index=table_index,
            header_columns=header_columns,
            rows_values=rows_values,
            write_row=write_row,
            cancel_callback=cancel_callback,
        )
        return

    raise ValueError(f"Unsupported template format: {template_type}")


def create_template_record(
    template_id: str,
    original_name: str,
    stored_path: Path,
    *,
    owner_id: str = "default",
) -> dict[str, Any]:
    resolved_owner_id = normalize_owner_id(owner_id)
    template_info = read_template_headers(stored_path, None, DEFAULT_HEADER_ROW)
    template_type = template_info["template_type"]

    if template_type == "docx":
        tables = template_info.get("tables") or []
        fields = build_docx_template_fields(tables)
        first_table = tables[0]
        return {
            "template_id": template_id,
            "owner_id": resolved_owner_id,
            "filename": original_name,
            "path": str(stored_path),
            "sheet_name": template_info["sheet_name"],
            "header_row": first_table.get("header_row"),
            "write_row": first_table.get("write_row"),
            "table_index": int(first_table["table_index"]),
            "template_type": "docx",
            "fields": fields,
            "tables": tables,
            "created_at": utcnow_iso(),
        }

    fields = build_template_fields(
        template_info["sheet_name"],
        template_info["header_columns"],
        header_row=template_info["header_row"],
        write_row=template_info["write_row"],
        template_type=template_type,
        table_index=template_info["table_index"],
    )
    return {
        "template_id": template_id,
        "owner_id": resolved_owner_id,
        "filename": original_name,
        "path": str(stored_path),
        "sheet_name": template_info["sheet_name"],
        "header_row": template_info["header_row"],
        "write_row": template_info["write_row"],
        "table_index": template_info["table_index"],
        "template_type": template_type,
        "fields": fields,
        "tables": [],
        "created_at": utcnow_iso(),
    }


__all__ = [
    "build_default_output_path",
    "build_docx_template_fields",
    "build_template_fields",
    "create_template_record",
    "read_docx_headers",
    "read_docx_tables",
    "read_excel_headers",
    "read_template_headers",
    "write_filled_docx",
    "write_filled_docx_blocks",
    "write_filled_excel",
    "write_filled_template",
]
