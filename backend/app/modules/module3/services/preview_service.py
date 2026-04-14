from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from ..core.config import PREVIEW_CHAR_LIMIT, PREVIEW_LINE_LIMIT
from .word_processor import squish


def read_text_with_fallback(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"鏃犳硶璇嗗埆鏂囨湰鏂囦欢缂栫爜: {path.name}")


def render_excel_preview(path: Path, max_rows: int | None = PREVIEW_LINE_LIMIT) -> str:
    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for index, ws in enumerate(wb.worksheets, start=1):
            remaining = max_rows
            if remaining is not None and remaining <= 0:
                break

            lines: list[str] = [f"[工作表 {index}] {ws.title}"]
            for row in ws.iter_rows(values_only=True):
                cells = ["" if value is None else squish(str(value)) for value in row]
                if not any(cells):
                    continue
                lines.append("\t".join(cells))
                if remaining is not None:
                    remaining -= 1
                if remaining is not None and remaining <= 0:
                    break

            if len(lines) > 1:
                parts.append("\n".join(lines))

        return "\n\n".join(parts).strip() or "[工作表 1] 空白工作表"
    finally:
        wb.close()


def render_docx_preview(path: Path, max_lines: int | None = PREVIEW_LINE_LIMIT) -> str:
    doc = Document(str(path))
    lines: list[str] = [f"[Word 鏂囨。] {path.name}"]
    remaining = max_lines

    for paragraph in doc.paragraphs:
        text = squish(paragraph.text)
        if not text:
            continue
        lines.append(text)
        if remaining is not None:
            remaining -= 1
        if remaining is not None and remaining <= 0:
            return "\n".join(lines)

    for index, table in enumerate(doc.tables, start=1):
        lines.append(f"[琛ㄦ牸 {index}]")
        for row in table.rows:
            cells = [squish(cell.text) for cell in row.cells]
            if not any(cells):
                continue
            lines.append("| " + " | ".join(cells) + " |")
            if remaining is not None:
                remaining -= 1
            if remaining is not None and remaining <= 0:
                return "\n".join(lines)

    return "\n".join(lines).strip()


def render_text_preview(path: Path, max_lines: int | None = PREVIEW_LINE_LIMIT) -> str:
    text = read_text_with_fallback(path).replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    source_lines = lines if max_lines is None else lines[:max_lines]
    kept = [line.rstrip() for line in source_lines]
    return "\n".join(kept).strip()


def build_preview_text(path: Path) -> str:
    line_limit = PREVIEW_LINE_LIMIT if PREVIEW_LINE_LIMIT > 0 else None
    char_limit = PREVIEW_CHAR_LIMIT if PREVIEW_CHAR_LIMIT > 0 else None

    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        content = render_excel_preview(path, max_rows=line_limit)
    elif suffix == ".docx":
        content = render_docx_preview(path, max_lines=line_limit)
    elif suffix in {".txt", ".md", ".json", ".log"}:
        content = render_text_preview(path, max_lines=line_limit)
    else:
        raise ValueError(f"鏆備笉鏀寔棰勮璇ユ枃浠剁被鍨? {path.suffix}")

    if char_limit is not None and len(content) > char_limit:
        return content[:char_limit].rstrip() + "\n..."
    return content


__all__ = [
    "build_preview_text",
    "read_text_with_fallback",
    "render_docx_preview",
    "render_excel_preview",
    "render_text_preview",
]
