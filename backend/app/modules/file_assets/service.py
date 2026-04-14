from __future__ import annotations

import json
import mimetypes
import os
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, TypedDict

from docx import Document
from openpyxl import load_workbook


class FileAssetRecord(TypedDict):
    file_id: str
    owner_id: str
    path: str
    filename: str
    source_module: str
    purpose: str
    suffix: str
    mime_type: str
    size_bytes: int
    created_at: str
    updated_at: str


_INDEX_LOCK = threading.Lock()
_REPO_ROOT = Path(__file__).resolve().parents[4]
_PREVIEWABLE_SUFFIXES = {".txt", ".md", ".json", ".log", ".csv", ".docx", ".xlsx"}


def _resolve_runtime_root() -> Path:
    configured = str(os.getenv("SMART_DOC_FLOW_RUNTIME_ROOT", "") or "").strip()
    if configured:
        return Path(configured).resolve()
    return _REPO_ROOT.resolve()


def _resolve_data_root(runtime_root: Path) -> Path:
    explicit = str(os.getenv("FILE_ASSETS_DATA_ROOT", "") or "").strip()
    if explicit:
        return Path(explicit).resolve()

    managed = str(os.getenv("SMART_DOC_FLOW_DATA_ROOT", "") or "").strip()
    if managed:
        managed_path = Path(managed).resolve()
        if managed_path == runtime_root.resolve():
            return (runtime_root / "runtime-data").resolve()
        return managed_path
    return (runtime_root / "runtime-data").resolve()


def _resolve_assets_root() -> Path:
    explicit = str(os.getenv("FILE_ASSETS_ROOT_DIR", "") or "").strip()
    if explicit:
        return Path(explicit).resolve()
    runtime_root = _resolve_runtime_root()
    data_root = _resolve_data_root(runtime_root)
    return (data_root / "file-assets").resolve()


ROOT_DIR = _resolve_assets_root()
INDEX_PATH = (ROOT_DIR / "index.json").resolve()


def _utcnow_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _normalize_owner_id(owner_id: str | None) -> str:
    normalized = str(owner_id or "").strip()
    return normalized or "default"


def _read_text_with_fallback(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"cannot decode file: {path.name}")


def _write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    temp_path = path.with_name(f"{path.name}.tmp")
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path.write_text(text, encoding="utf-8")
    os.replace(temp_path, path)


def ensure_file_assets_storage() -> None:
    ROOT_DIR.mkdir(parents=True, exist_ok=True)
    if INDEX_PATH.exists():
        return
    _write_json_atomic(INDEX_PATH, {"version": 1, "items": {}})


def _load_index() -> dict[str, Any]:
    ensure_file_assets_storage()
    try:
        raw = json.loads(INDEX_PATH.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        raw = {"version": 1, "items": {}}
    if not isinstance(raw, dict):
        raw = {"version": 1, "items": {}}
    items = raw.get("items")
    if not isinstance(items, dict):
        raw["items"] = {}
    raw.setdefault("version", 1)
    return raw


def _save_index(index_payload: dict[str, Any]) -> None:
    _write_json_atomic(INDEX_PATH, index_payload)


def _coerce_record(raw: dict[str, Any]) -> FileAssetRecord:
    return FileAssetRecord(
        file_id=str(raw.get("file_id") or "").strip(),
        owner_id=_normalize_owner_id(str(raw.get("owner_id") or "")),
        path=str(raw.get("path") or "").strip(),
        filename=str(raw.get("filename") or "").strip(),
        source_module=str(raw.get("source_module") or "").strip() or "unknown",
        purpose=str(raw.get("purpose") or "").strip() or "general",
        suffix=str(raw.get("suffix") or "").strip().lower(),
        mime_type=str(raw.get("mime_type") or "").strip() or "application/octet-stream",
        size_bytes=int(raw.get("size_bytes") or 0),
        created_at=str(raw.get("created_at") or "").strip(),
        updated_at=str(raw.get("updated_at") or "").strip(),
    )


def preview_support_for_suffix(suffix: str | None) -> tuple[bool, str | None]:
    normalized = str(suffix or "").strip().lower()
    if not normalized:
        return False, "preview is not supported for files without suffix"
    if not normalized.startswith("."):
        normalized = f".{normalized}"
    if normalized in _PREVIEWABLE_SUFFIXES:
        return True, None
    return False, f"preview is not supported for suffix: {normalized}"


def build_file_asset_urls(file_id: str) -> dict[str, str]:
    normalized = str(file_id or "").strip()
    if not normalized:
        return {"download_url": "", "preview_url": ""}
    return {
        "download_url": f"/download/{normalized}",
        "preview_url": f"/file-assets/{normalized}/preview",
    }


def register_file(
    *,
    path: str | Path,
    owner_id: str,
    source_module: str,
    purpose: str,
    filename: str | None = None,
) -> FileAssetRecord:
    resolved_path = Path(path).expanduser().resolve()
    if not resolved_path.exists() or not resolved_path.is_file():
        raise FileNotFoundError(f"file not found: {resolved_path}")

    normalized_owner = _normalize_owner_id(owner_id)
    now = _utcnow_iso()
    guessed_mime = mimetypes.guess_type(str(resolved_path))[0] or "application/octet-stream"
    suffix = resolved_path.suffix.lower()
    effective_name = str(filename or resolved_path.name).strip() or resolved_path.name

    with _INDEX_LOCK:
        index_payload = _load_index()
        items = index_payload["items"]

        for raw in items.values():
            if not isinstance(raw, dict):
                continue
            record = _coerce_record(raw)
            if record["owner_id"] != normalized_owner:
                continue
            if Path(record["path"]).resolve() != resolved_path:
                continue
            if record["size_bytes"] != resolved_path.stat().st_size:
                continue

            record["filename"] = effective_name
            record["source_module"] = str(source_module or "unknown").strip() or "unknown"
            record["purpose"] = str(purpose or "general").strip() or "general"
            record["suffix"] = suffix
            record["mime_type"] = guessed_mime
            record["size_bytes"] = int(resolved_path.stat().st_size)
            record["updated_at"] = now
            items[record["file_id"]] = dict(record)
            _save_index(index_payload)
            return record

        file_id = f"fa_{uuid.uuid4().hex[:16]}"
        record = FileAssetRecord(
            file_id=file_id,
            owner_id=normalized_owner,
            path=str(resolved_path),
            filename=effective_name,
            source_module=str(source_module or "unknown").strip() or "unknown",
            purpose=str(purpose or "general").strip() or "general",
            suffix=suffix,
            mime_type=guessed_mime,
            size_bytes=int(resolved_path.stat().st_size),
            created_at=now,
            updated_at=now,
        )
        items[file_id] = dict(record)
        _save_index(index_payload)
        return record


def get_file_record(file_id: str) -> FileAssetRecord | None:
    normalized_id = str(file_id or "").strip()
    if not normalized_id:
        return None
    with _INDEX_LOCK:
        index_payload = _load_index()
        raw = index_payload.get("items", {}).get(normalized_id)
        if not isinstance(raw, dict):
            return None
        record = _coerce_record(raw)
        if not record["file_id"]:
            return None
        return record


def resolve_file_for_owner(file_id: str, owner_id: str) -> tuple[FileAssetRecord, Path]:
    record = get_file_record(file_id)
    if record is None:
        raise FileNotFoundError(f"file asset not found: {file_id}")

    normalized_owner = _normalize_owner_id(owner_id)
    if record["owner_id"] != normalized_owner:
        raise PermissionError("no permission to access this file asset")

    file_path = Path(record["path"]).expanduser().resolve()
    if not file_path.exists() or not file_path.is_file():
        raise FileNotFoundError(f"file missing on disk: {file_path}")
    return record, file_path


def _preview_plain_text(path: Path, *, max_lines: int, max_chars: int) -> str:
    text = _read_text_with_fallback(path).replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    if max_lines > 0:
        lines = lines[:max_lines]
    preview = "\n".join(line.rstrip() for line in lines).strip()
    if max_chars > 0 and len(preview) > max_chars:
        return preview[:max_chars].rstrip() + "\n..."
    return preview


def _preview_docx(path: Path, *, max_lines: int, max_chars: int) -> str:
    doc = Document(str(path))
    lines: list[str] = [f"[DOCX] {path.name}"]
    remaining = max_lines if max_lines > 0 else 999999

    for paragraph in doc.paragraphs:
        text = " ".join(str(paragraph.text or "").split()).strip()
        if not text:
            continue
        lines.append(text)
        remaining -= 1
        if remaining <= 0:
            break

    if remaining > 0:
        for table_index, table in enumerate(doc.tables, start=1):
            lines.append(f"[TABLE {table_index}]")
            remaining -= 1
            if remaining <= 0:
                break
            for row in table.rows:
                cells = [" ".join(str(cell.text or "").split()).strip() for cell in row.cells]
                if not any(cells):
                    continue
                lines.append(" | ".join(cells))
                remaining -= 1
                if remaining <= 0:
                    break
            if remaining <= 0:
                break

    preview = "\n".join(lines).strip()
    if max_chars > 0 and len(preview) > max_chars:
        return preview[:max_chars].rstrip() + "\n..."
    return preview


def _preview_xlsx(path: Path, *, max_lines: int, max_chars: int) -> str:
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        lines: list[str] = []
        remaining = max_lines if max_lines > 0 else 999999
        for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
            lines.append(f"[SHEET {sheet_index}] {worksheet.title}")
            remaining -= 1
            if remaining <= 0:
                break
            for row in worksheet.iter_rows(values_only=True):
                cells = ["" if value is None else " ".join(str(value).split()).strip() for value in row]
                if not any(cells):
                    continue
                lines.append("\t".join(cells))
                remaining -= 1
                if remaining <= 0:
                    break
            if remaining <= 0:
                break

        preview = "\n".join(lines).strip()
        if max_chars > 0 and len(preview) > max_chars:
            return preview[:max_chars].rstrip() + "\n..."
        return preview
    finally:
        workbook.close()


def build_preview_text_for_owner(
    file_id: str,
    owner_id: str,
    *,
    max_lines: int = 240,
    max_chars: int = 12000,
) -> str:
    _, file_path = resolve_file_for_owner(file_id, owner_id)
    suffix = file_path.suffix.lower()
    supported, reason = preview_support_for_suffix(suffix)
    if not supported:
        raise ValueError(reason or f"preview is not supported for suffix: {suffix}")
    if suffix in {".txt", ".md", ".json", ".log", ".csv"}:
        return _preview_plain_text(file_path, max_lines=max_lines, max_chars=max_chars)
    if suffix == ".docx":
        return _preview_docx(file_path, max_lines=max_lines, max_chars=max_chars)
    if suffix == ".xlsx":
        return _preview_xlsx(file_path, max_lines=max_lines, max_chars=max_chars)
    raise ValueError(reason or f"preview is not supported for suffix: {suffix}")


__all__ = [
    "FileAssetRecord",
    "INDEX_PATH",
    "ROOT_DIR",
    "build_file_asset_urls",
    "build_preview_text_for_owner",
    "ensure_file_assets_storage",
    "get_file_record",
    "preview_support_for_suffix",
    "register_file",
    "resolve_file_for_owner",
]
