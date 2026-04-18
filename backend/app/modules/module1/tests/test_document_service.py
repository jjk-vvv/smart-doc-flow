from __future__ import annotations

import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import docx

BACKEND_ROOT = Path(__file__).resolve().parents[4]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.modules.module1.services.document_mutation_support import (
    _apply_modifications_to_docx,
    _resolve_docx_paragraph_indexes,
)
from app.modules.module1.services.document_io_support import _matches_file_signature
from app.modules.module1.services.modification_support import (
    _infer_format_update_policy,
    _normalize_ai_modification,
    _sanitize_modifications_for_instruction,
)


class DocumentServiceTestCase(unittest.TestCase):
    def test_normalize_ai_modification_handles_string_inputs(self) -> None:
        normalized = _normalize_ai_modification(
            {
                "type": "text_run_format",
                "lineIndex": "2",
                "bold": "true",
                "italic": "false",
                "color": "红色",
                "alignment": "居中",
                "runIndices": ["0", "bad", 2, -1, "2"],
            }
        )

        self.assertIsNotNone(normalized)
        self.assertEqual(normalized["type"], "text_run_format")
        self.assertEqual(normalized["line_index"], 2)
        self.assertTrue(normalized["bold"])
        self.assertFalse(normalized["italic"])
        self.assertEqual(normalized["color"], "#FF0000")
        self.assertEqual(normalized["alignment"], "center")
        self.assertEqual(normalized["run_indices"], [0, 2])

    def test_resolve_docx_paragraph_indexes_supports_normalized_reference(self) -> None:
        document = docx.Document()
        document.add_paragraph("封面标题")
        document.add_paragraph("项目总金额：1,250 元")

        paragraph_indexes = _resolve_docx_paragraph_indexes(
            document,
            {
                "type": "text_run_format",
                "reference_text": "项目总金额1250元",
            },
        )

        self.assertEqual(paragraph_indexes, [1])

    def test_apply_replace_text_to_docx_preserves_basic_style(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "sample.docx"
            document = docx.Document()
            paragraph = document.add_paragraph()
            run = paragraph.add_run("原始文本")
            run.bold = True
            document.save(input_path)

            output_path = _apply_modifications_to_docx(
                str(input_path),
                [
                    {
                        "type": "replace_text",
                        "old_text": "原始文本",
                        "new_text": "更新文本",
                    }
                ],
            )

            result = docx.Document(output_path)
            self.assertEqual(result.paragraphs[0].text, "更新文本")
            self.assertTrue(result.paragraphs[0].runs[0].bold)

    def test_sanitize_modifications_keeps_only_explicit_color_change(self) -> None:
        modifications = [
            {
                "type": "text_run_format",
                "reference_text": "项目概述",
                "color": "#FF0000",
                "bold": False,
                "heading_level": 0,
            }
        ]

        sanitized = _sanitize_modifications_for_instruction(modifications, "把标题“项目概述”改成红色")

        self.assertEqual(
            sanitized,
            [
                {
                    "type": "text_run_format",
                    "reference_text": "项目概述",
                    "color": "#FF0000",
                }
            ],
        )

    def test_sanitize_modifications_allows_explicit_false_for_bold(self) -> None:
        modifications = [
            {
                "type": "text_run_format",
                "reference_text": "项目概述",
                "color": "#FF0000",
                "bold": False,
                "heading_level": 0,
            }
        ]

        sanitized = _sanitize_modifications_for_instruction(modifications, "把标题“项目概述”取消加粗并改成红色")

        self.assertEqual(
            sanitized,
            [
                {
                    "type": "text_run_format",
                    "reference_text": "项目概述",
                    "color": "#FF0000",
                    "bold": False,
                }
            ],
        )

    def test_sanitize_modifications_drops_unmentioned_style_fields(self) -> None:
        modifications = [
            {
                "type": "text_run_format",
                "reference_text": "项目概述",
                "alignment": "center",
                "font_size": 16,
                "bold": False,
                "color": "#FF0000",
            }
        ]

        sanitized = _sanitize_modifications_for_instruction(modifications, "把标题“项目概述”居中并改成16磅")

        self.assertEqual(
            sanitized,
            [
                {
                    "type": "text_run_format",
                    "reference_text": "项目概述",
                    "alignment": "center",
                    "font_size": 16.0,
                }
            ],
        )

    def test_sanitize_modifications_drops_empty_text_run_format(self) -> None:
        modifications = [
            {
                "type": "text_run_format",
                "reference_text": "项目概述",
                "bold": False,
            }
        ]

        sanitized = _sanitize_modifications_for_instruction(modifications, "把标题“项目概述”改一下格式")

        self.assertEqual(sanitized, [])

    def test_apply_color_change_preserves_existing_bold_style(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "sample.docx"
            document = docx.Document()
            paragraph = document.add_paragraph()
            run = paragraph.add_run("项目概述")
            run.bold = True
            document.save(input_path)

            raw_modifications = [
                {
                    "type": "text_run_format",
                    "reference_text": "项目概述",
                    "color": "#FF0000",
                    "bold": False,
                }
            ]
            sanitized = _sanitize_modifications_for_instruction(raw_modifications, "把标题“项目概述”改成红色")

            output_path = _apply_modifications_to_docx(str(input_path), sanitized)
            result = docx.Document(output_path)

            self.assertTrue(result.paragraphs[0].runs[0].bold)
            self.assertEqual(str(result.paragraphs[0].runs[0].font.color.rgb), "FF0000")

    def test_infer_format_update_policy_marks_explicit_false_fields(self) -> None:
        policy = _infer_format_update_policy("把标题取消加粗并改成红色")

        self.assertEqual(set(policy["allowed_fields"]), {"bold", "color"})
        self.assertEqual(set(policy["explicit_false_fields"]), {"bold"})

    def test_matches_file_signature_validates_docx_zip_payload(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("[Content_Types].xml", "")
            archive.writestr("word/document.xml", "<w:document/>")

        self.assertTrue(_matches_file_signature(".docx", buffer.getvalue()))
        self.assertFalse(_matches_file_signature(".docx", b"not-a-zip"))


if __name__ == "__main__":
    unittest.main()
