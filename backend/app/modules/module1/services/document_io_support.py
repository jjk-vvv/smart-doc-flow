from __future__ import annotations

import io
import json
import logging
import os
import re
import shutil
import urllib.parse
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from ..domain import SavedUpload
from .conversion_service import convert_doc_bytes_to_docx


logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {".txt", ".md", ".csv"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}
SPREADSHEET_EXTENSIONS = {".xlsx"}
WORD_EXTENSIONS = {".doc", ".docx"}
SUPPORTED_UPLOAD_ORDER = (
    ".pdf",
    ".doc",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".xlsx",
    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".bmp",
    ".webp",
)
SUPPORTED_UPLOAD_EXTENSIONS = set(SUPPORTED_UPLOAD_ORDER)
TEXT_FILE_ENCODINGS = ("utf-8", "utf-8-sig", "gbk")

try:
    import PyPDF2
except ImportError:  # pragma: no cover - optional dependency
    PyPDF2 = None

try:
    import docx
except ImportError:  # pragma: no cover - optional dependency
    docx = None

try:
    import pdfplumber
except ImportError:  # pragma: no cover - optional dependency
    pdfplumber = None

try:
    from PIL import Image
except ImportError:  # pragma: no cover - optional dependency
    Image = None


def _libreoffice_available() -> bool:
    env_candidates = [os.getenv("MODULE1_SOFFICE_PATH", "").strip(), os.getenv("MODULE2_SOFFICE_PATH", "").strip()]
    for item in env_candidates:
        if item and Path(item).exists():
            return True
    return shutil.which("soffice") is not None


def _media_type_for_name(file_name: str) -> str:
    ext = Path(file_name or "").suffix.lower()
    if ext == ".pdf":
        return "application/pdf"
    if ext == ".docx":
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if ext == ".doc":
        return "application/msword"
    if ext == ".md":
        return "text/markdown"
    if ext in {".txt", ".csv"}:
        return "text/plain"
    return "application/octet-stream"


def _build_download_headers(file_name: str, modification_count: int, *, applied: bool) -> dict[str, str]:
    encoded_name = urllib.parse.quote(file_name)
    return {
        "Content-Disposition": f"attachment; filename={encoded_name}; filename*=UTF-8''{encoded_name}",
        "X-Modification-Applied": "true" if applied else "false",
        "X-Modification-Count": str(modification_count),
    }


def _supported_upload_types_text() -> str:
    return "PDF、DOC、DOCX、TXT、MD、CSV、XLSX、JPG、JPEG、PNG、GIF、BMP、WEBP"


def _validate_saved_upload(saved_upload: SavedUpload) -> str:
    source_ext = Path(saved_upload.filename).suffix.lower()
    if source_ext not in SUPPORTED_UPLOAD_EXTENSIONS:
        raise ValueError(f"不支持的文件类型: {source_ext or saved_upload.filename}。当前支持：{_supported_upload_types_text()}")
    if not _matches_file_signature(source_ext, saved_upload.file_bytes):
        raise ValueError("文件内容与扩展名不匹配，请确认上传了正确的文件类型")
    return source_ext


def _matches_file_signature(file_ext: str, file_bytes: bytes) -> bool:
    if not file_bytes:
        return False

    sample = file_bytes[:32]
    normalized_ext = str(file_ext or "").lower()

    if normalized_ext in TEXT_EXTENSIONS:
        return _looks_like_text_payload(file_bytes)
    if normalized_ext == ".pdf":
        return sample.startswith(b"%PDF-")
    if normalized_ext == ".doc":
        return sample.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    if normalized_ext in {".docx", ".xlsx"}:
        if not zipfile.is_zipfile(io.BytesIO(file_bytes)):
            return False
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes), "r") as archive:
                names = set(archive.namelist())
        except zipfile.BadZipFile:
            return False
        if normalized_ext == ".docx":
            return "word/document.xml" in names
        return "xl/workbook.xml" in names
    if normalized_ext == ".png":
        return sample.startswith(b"\x89PNG\r\n\x1a\n")
    if normalized_ext in {".jpg", ".jpeg"}:
        return sample.startswith(b"\xff\xd8\xff")
    if normalized_ext == ".gif":
        return sample.startswith((b"GIF87a", b"GIF89a"))
    if normalized_ext == ".bmp":
        return sample.startswith(b"BM")
    if normalized_ext == ".webp":
        return sample.startswith(b"RIFF") and sample[8:12] == b"WEBP"
    return False


def _looks_like_text_payload(file_bytes: bytes) -> bool:
    sample = file_bytes[:4096]
    if not sample:
        return False
    if b"\x00" in sample:
        return False
    for encoding in TEXT_FILE_ENCODINGS:
        try:
            sample.decode(encoding)
            return True
        except UnicodeDecodeError:
            continue
    return False


def _prepare_docx_source(saved_upload: SavedUpload) -> tuple[str, str]:
    ext = Path(saved_upload.filename).suffix.lower()
    if ext != ".doc":
        return saved_upload.path, ""

    converted_bytes, _ = convert_doc_bytes_to_docx(
        file_name=saved_upload.filename,
        file_bytes=saved_upload.file_bytes,
        timeout_sec=90.0,
    )
    converted_path = f"{saved_upload.path}.docx"
    with open(converted_path, "wb") as converted_file:
        converted_file.write(converted_bytes)
    return converted_path, converted_path


def _extract_text(file_path: str, file_ext: str) -> str:
    if file_ext == ".pdf":
        return _extract_text_from_pdf(file_path)
    if file_ext == ".docx":
        return _extract_text_from_docx(file_path)
    if file_ext in TEXT_EXTENSIONS:
        return _read_text_file(file_path)
    if file_ext in SPREADSHEET_EXTENSIONS:
        return _extract_text_from_xlsx(file_path)
    if file_ext in IMAGE_EXTENSIONS:
        return "这是一个图片文件。"
    return f"不支持的文件类型: {file_ext}"


def _extract_text_simple(file_path: str, file_ext: str) -> str:
    if file_ext == ".pdf":
        return _extract_text_from_pdf(file_path, limit=5000)
    if file_ext == ".docx":
        return _extract_text_from_docx_plain(file_path, limit=5000)
    if file_ext in TEXT_EXTENSIONS:
        return _read_text_file(file_path)[:5000]
    if file_ext in SPREADSHEET_EXTENSIONS:
        return _extract_text_from_xlsx(file_path, limit=5000)
    return ""


def _extract_text_from_pdf(file_path: str, limit: int | None = None) -> str:
    if PyPDF2 is None and pdfplumber is None:
        return "PDF 处理库未安装，无法提取文本。"

    candidates: list[str] = []

    if pdfplumber is not None:
        try:
            pdfplumber_text = _extract_text_from_pdf_with_pdfplumber(file_path)
            if pdfplumber_text.strip():
                candidates.append(pdfplumber_text)
        except Exception as exc:
            logger.warning("pdfplumber 提取 PDF 文本失败: %s", exc)

    if PyPDF2 is not None:
        try:
            pypdf2_text = _extract_text_from_pdf_with_pypdf2(file_path)
            if pypdf2_text.strip():
                candidates.append(pypdf2_text)
        except Exception as exc:
            logger.warning("PyPDF2 提取 PDF 文本失败: %s", exc)

    text = max(candidates, key=_score_extracted_text, default="").strip()
    if limit is not None:
        return text[:limit]
    return text


def _extract_text_from_pdf_with_pypdf2(file_path: str) -> str:
    if PyPDF2 is None:
        return ""

    text_parts: list[str] = []
    with open(file_path, "rb") as pdf_file:
        reader = PyPDF2.PdfReader(pdf_file)
        for page in reader.pages:
            text_parts.append((page.extract_text() or "").strip())
    return "\n".join(part for part in text_parts if part).strip()


def _extract_text_from_pdf_with_pdfplumber(file_path: str) -> str:
    if pdfplumber is None:
        return ""

    text_parts: list[str] = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = (page.extract_text(layout=True) or page.extract_text() or "").strip()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts).strip()


def _score_extracted_text(text: str) -> int:
    normalized = re.sub(r"\s+", "", str(text or ""))
    if not normalized:
        return 0
    non_empty_lines = [line for line in str(text or "").splitlines() if line.strip()]
    return len(normalized) + len(non_empty_lines) * 24 + len(set(normalized))


def _extract_text_from_docx(file_path: str) -> str:
    if docx is None:
        return "Word 处理库未安装，无法提取文本。"

    try:
        document = docx.Document(file_path)
        paragraphs = []
        for index, paragraph in enumerate(document.paragraphs):
            paragraphs.append(
                {
                    "index": index,
                    "text": paragraph.text,
                    "style": paragraph.style.name if paragraph.style else "Normal",
                    "runs": [
                        {
                            "text": run.text,
                            "bold": run.bold,
                            "italic": run.italic,
                            "underline": run.underline,
                            "font_name": run.font.name,
                            "font_size": run.font.size.pt if run.font.size else None,
                            "color": str(run.font.color.rgb) if run.font.color and run.font.color.rgb else None,
                        }
                        for run in paragraph.runs
                    ],
                }
            )

        tables = []
        for table in document.tables:
            tables.append([[cell.text for cell in row.cells] for row in table.rows])

        return json.dumps(
            {
                "paragraphs": paragraphs,
                "tables": tables,
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
            },
            ensure_ascii=False,
            indent=2,
        )
    except Exception as exc:
        logger.warning("python-docx 解析失败，降级为纯文本提取: %s", exc)
        fallback_text = _extract_text_from_docx_plain(file_path)
        fallback_lines = [line.strip() for line in fallback_text.splitlines() if line.strip()]

        return json.dumps(
            {
                "paragraphs": [
                    {
                        "index": index,
                        "text": line,
                        "style": "RecoveredText",
                        "runs": [{"text": line}],
                    }
                    for index, line in enumerate(fallback_lines)
                ],
                "tables": [],
                "paragraph_count": len(fallback_lines),
                "table_count": 0,
                "recovered": True,
            },
            ensure_ascii=False,
            indent=2,
        )


def _extract_text_from_docx_plain(file_path: str, limit: int | None = None) -> str:
    if docx is None:
        return "Word 处理库未安装，无法提取文本。"

    try:
        document = docx.Document(file_path)
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells if cell.text.strip())
                if row_text:
                    lines.append(row_text)
        text = "\n".join(lines).strip()
    except Exception:
        with zipfile.ZipFile(file_path, "r") as docx_zip:
            text_parts: list[str] = []
            if "word/document.xml" in docx_zip.namelist():
                root = ET.fromstring(docx_zip.read("word/document.xml"))
                for element in root.iter():
                    if element.tag.endswith("t") and element.text and element.text.strip():
                        text_parts.append(element.text)
            text = "\n".join(text_parts).strip()

    if limit is not None:
        return text[:limit]
    return text


def _extract_text_from_xlsx(file_path: str, limit: int | None = None) -> str:
    rows: list[str] = []
    with zipfile.ZipFile(file_path, "r") as workbook:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in workbook.namelist():
            root = ET.fromstring(workbook.read("xl/sharedStrings.xml"))
            for item in root.iter():
                if item.tag.endswith("t"):
                    shared_strings.append(item.text or "")

        worksheet_names = sorted(
            name for name in workbook.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        )
        for worksheet_name in worksheet_names:
            root = ET.fromstring(workbook.read(worksheet_name))
            for row in root.iter():
                if not row.tag.endswith("row"):
                    continue
                values: list[str] = []
                for cell in row:
                    if not cell.tag.endswith("c"):
                        continue
                    cell_type = cell.attrib.get("t")
                    raw_value = ""
                    inline_text_parts: list[str] = []

                    for child in cell:
                        if child.tag.endswith("v") and child.text is not None:
                            raw_value = child.text
                        elif child.tag.endswith("is"):
                            inline_text_parts.extend(
                                element.text or ""
                                for element in child.iter()
                                if element.tag.endswith("t") and element.text is not None
                            )

                    if cell_type == "inlineStr" and inline_text_parts:
                        values.append("".join(inline_text_parts).strip())
                        continue
                    if not raw_value:
                        continue
                    if cell_type == "s":
                        try:
                            values.append(shared_strings[int(raw_value)])
                        except Exception:
                            values.append(raw_value)
                    elif cell_type == "str":
                        values.append(raw_value)
                    else:
                        values.append(raw_value)
                if values:
                    rows.append(" | ".join(values))

    text = "\n".join(rows).strip()
    if limit is not None:
        return text[:limit]
    return text


def _read_text_file(file_path: str) -> str:
    for encoding in TEXT_FILE_ENCODINGS:
        try:
            with open(file_path, "r", encoding=encoding) as text_file:
                return text_file.read()
        except UnicodeDecodeError:
            continue
    raise ValueError("无法解码文本文件")


__all__ = [
    "Image",
    "IMAGE_EXTENSIONS",
    "PyPDF2",
    "SPREADSHEET_EXTENSIONS",
    "SUPPORTED_UPLOAD_EXTENSIONS",
    "SUPPORTED_UPLOAD_ORDER",
    "TEXT_EXTENSIONS",
    "TEXT_FILE_ENCODINGS",
    "WORD_EXTENSIONS",
    "_build_download_headers",
    "_extract_text",
    "_extract_text_from_docx",
    "_extract_text_from_docx_plain",
    "_extract_text_from_pdf",
    "_extract_text_from_pdf_with_pdfplumber",
    "_extract_text_from_pdf_with_pypdf2",
    "_extract_text_from_xlsx",
    "_extract_text_simple",
    "_libreoffice_available",
    "_looks_like_text_payload",
    "_matches_file_signature",
    "_media_type_for_name",
    "_prepare_docx_source",
    "_read_text_file",
    "_score_extracted_text",
    "_supported_upload_types_text",
    "_validate_saved_upload",
    "docx",
    "pdfplumber",
]
