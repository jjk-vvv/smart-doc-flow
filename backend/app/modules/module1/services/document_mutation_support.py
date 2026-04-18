from __future__ import annotations

import re
from typing import Any

from .document_io_support import _read_text_file
from .modification_support import (
    _find_best_matching_paragraph,
    _normalize_alignment,
    _normalize_color,
    _normalize_match_text,
)
from .sandbox_service import derive_output_path

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Pt, RGBColor
    from docx.text.paragraph import Paragraph
except ImportError:  # pragma: no cover - optional dependency
    docx = None
    WD_ALIGN_PARAGRAPH = None
    OxmlElement = None
    Paragraph = None
    Pt = None
    RGBColor = None


def _write_content_to_docx(output_path: str, content: str) -> str:
    if docx is None:
        raise RuntimeError("python-docx 未安装")

    document = docx.Document()
    normalized_content = str(content or "").replace("\r", "").strip()
    blocks = re.split(r"\n\s*\n", normalized_content) if normalized_content else []

    if not blocks:
        document.add_paragraph("")
    else:
        for block in blocks:
            for line in block.split("\n"):
                stripped_line = line.strip()
                if not stripped_line:
                    continue

                heading_match = re.match(r"^(#{1,3})\s+(.+)$", stripped_line)
                if heading_match:
                    paragraph = document.add_paragraph(heading_match.group(2).strip())
                    try:
                        paragraph.style = f"Heading {len(heading_match.group(1))}"
                    except Exception:
                        pass
                    continue

                document.add_paragraph(stripped_line)

    document.save(output_path)
    return output_path


def _apply_modifications_to_docx(input_path: str, modifications: list[dict[str, Any]]) -> str:
    if docx is None:
        raise RuntimeError("python-docx 未安装")

    document = docx.Document(input_path)
    for modification in modifications:
        modification_type = str(modification.get("type") or "").lower()
        if modification_type == "text_run_format":
            _apply_text_run_format(document, modification)
        elif modification_type == "replace_text":
            _apply_replace_text_to_docx(document, modification)
        elif modification_type == "insert_text":
            _apply_insert_text_to_docx(document, modification)

    output_path = derive_output_path(input_path, prefix="modified_")
    document.save(output_path)
    return output_path


def _resolve_docx_paragraph_indexes(document: Any, modification: dict[str, Any]) -> list[int]:
    paragraph_index = modification.get("paragraph_index")
    if isinstance(paragraph_index, int) and 0 <= paragraph_index < len(document.paragraphs):
        return [paragraph_index]

    line_index = modification.get("line_index")
    if isinstance(line_index, int) and 0 <= line_index < len(document.paragraphs):
        return [line_index]

    reference_text = str(modification.get("reference_text") or "").strip()
    if reference_text:
        for index, paragraph in enumerate(document.paragraphs):
            if reference_text and reference_text in paragraph.text:
                return [index]
        normalized_reference = _normalize_match_text(reference_text)
        if normalized_reference:
            for index, paragraph in enumerate(document.paragraphs):
                if normalized_reference in _normalize_match_text(paragraph.text):
                    return [index]
            best_index = _find_best_matching_paragraph(document.paragraphs, reference_text)
            if best_index is not None:
                return [best_index]

    return []


def _insert_docx_paragraph_after(paragraph: Any, text: str) -> None:
    if OxmlElement is None or Paragraph is None:
        paragraph.add_run(f"\n{text}")
        return

    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = Paragraph(new_element, paragraph._parent)
    try:
        new_paragraph.style = paragraph.style
    except Exception:
        pass
    if text:
        new_paragraph.add_run(text)


def _apply_text_run_format(document: Any, modification: dict[str, Any]) -> None:
    paragraph_indexes = _resolve_docx_paragraph_indexes(document, modification)
    if not paragraph_indexes:
        return

    requested_indices = modification.get("run_indices")
    color = _normalize_color(modification.get("color"))
    font_name = str(modification.get("font_name") or "").strip()
    font_size = modification.get("font_size")
    alignment = _normalize_alignment(modification.get("alignment"))
    heading_level = modification.get("heading_level")
    line_spacing = modification.get("line_spacing")
    space_before = modification.get("space_before")
    space_after = modification.get("space_after")

    for paragraph_index in paragraph_indexes:
        paragraph = document.paragraphs[paragraph_index]
        if not paragraph.runs and paragraph.text:
            paragraph.add_run(paragraph.text)

        runs = paragraph.runs
        if runs:
            if isinstance(requested_indices, list) and requested_indices:
                run_indices = [index for index in requested_indices if isinstance(index, int) and 0 <= index < len(runs)]
            else:
                run_indices = list(range(len(runs)))

            for run_index in run_indices:
                run = runs[run_index]
                if "bold" in modification:
                    run.bold = bool(modification["bold"])
                if "italic" in modification:
                    run.italic = bool(modification["italic"])
                if "underline" in modification:
                    run.underline = bool(modification["underline"])
                if font_name:
                    run.font.name = font_name
                if isinstance(font_size, (int, float)) and Pt is not None:
                    run.font.size = Pt(float(font_size))
                if color and RGBColor is not None:
                    run.font.color.rgb = RGBColor(int(color[1:3], 16), int(color[3:5], 16), int(color[5:7], 16))

        if alignment is not None and WD_ALIGN_PARAGRAPH is not None:
            alignment_value_map = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }
            paragraph.alignment = alignment_value_map.get(alignment, paragraph.alignment)

        if isinstance(heading_level, int):
            if heading_level <= 0:
                try:
                    paragraph.style = document.styles["Normal"]
                except Exception:
                    pass
            else:
                try:
                    paragraph.style = document.styles[f"Heading {max(1, min(3, heading_level))}"]
                except Exception:
                    pass

        paragraph_format = paragraph.paragraph_format
        if isinstance(line_spacing, (int, float)):
            paragraph_format.line_spacing = float(line_spacing)

        if isinstance(space_before, (int, float)) and Pt is not None:
            paragraph_format.space_before = Pt(float(space_before))

        if isinstance(space_after, (int, float)) and Pt is not None:
            paragraph_format.space_after = Pt(float(space_after))


def _apply_replace_text_to_docx(document: Any, modification: dict[str, Any]) -> None:
    old_text = str(modification.get("old_text") or "")
    new_text = str(modification.get("new_text") or "")
    if not old_text:
        return

    for paragraph in document.paragraphs:
        _replace_text_in_paragraph(paragraph, old_text, new_text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_text_in_paragraph(paragraph, old_text, new_text)


def _apply_insert_text_to_docx(document: Any, modification: dict[str, Any]) -> None:
    insert_text = str(modification.get("text") or "")
    if not insert_text:
        return

    position = str(modification.get("position") or "end").lower()
    paragraph_indexes = _resolve_docx_paragraph_indexes(document, modification)

    if paragraph_indexes:
        paragraph = document.paragraphs[paragraph_indexes[0]]
        if position == "start":
            _prepend_text_to_paragraph(paragraph, insert_text)
        elif position == "before":
            paragraph.insert_paragraph_before(insert_text)
        elif position == "after":
            _insert_docx_paragraph_after(paragraph, insert_text)
        else:
            _append_text_to_paragraph(paragraph, insert_text)
        return

    if position == "start" and document.paragraphs:
        _prepend_text_to_paragraph(document.paragraphs[0], insert_text)
    else:
        document.add_paragraph(insert_text)


def _set_paragraph_text(paragraph: Any, text: str, *, template_run: Any | None = None) -> None:
    template_style = _capture_run_style(template_run or (paragraph.runs[0] if paragraph.runs else None))
    paragraph_element = paragraph._element
    for child in list(paragraph_element):
        paragraph_element.remove(child)
    new_run = paragraph.add_run(text)
    _apply_run_style(new_run, template_style)


def _replace_text_in_paragraph(paragraph: Any, old_text: str, new_text: str) -> bool:
    if old_text not in paragraph.text:
        return False

    for run in paragraph.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            return True

    template_run = paragraph.runs[0] if paragraph.runs else None
    _set_paragraph_text(paragraph, paragraph.text.replace(old_text, new_text), template_run=template_run)
    return True


def _prepend_text_to_paragraph(paragraph: Any, text: str) -> None:
    if not text:
        return
    if not paragraph.runs:
        paragraph.add_run(text)
        return

    first_run = paragraph.runs[0]
    new_run = paragraph.add_run(text)
    _apply_run_style(new_run, _capture_run_style(first_run))
    first_run._r.addprevious(new_run._r)


def _append_text_to_paragraph(paragraph: Any, text: str) -> None:
    if not text:
        return

    template_run = paragraph.runs[-1] if paragraph.runs else None
    new_run = paragraph.add_run(text)
    _apply_run_style(new_run, _capture_run_style(template_run))


def _capture_run_style(run: Any | None) -> dict[str, Any]:
    if run is None:
        return {}

    style: dict[str, Any] = {
        "bold": run.bold,
        "italic": run.italic,
        "underline": run.underline,
    }

    if getattr(run.font, "name", None):
        style["font_name"] = run.font.name
    if getattr(run.font, "size", None) is not None:
        style["font_size"] = run.font.size
    if getattr(run.font, "color", None) is not None and getattr(run.font.color, "rgb", None) is not None:
        style["color_rgb"] = run.font.color.rgb

    return style


def _apply_run_style(run: Any, style: dict[str, Any]) -> None:
    if not style:
        return

    if "bold" in style:
        run.bold = style["bold"]
    if "italic" in style:
        run.italic = style["italic"]
    if "underline" in style:
        run.underline = style["underline"]
    if style.get("font_name"):
        run.font.name = style["font_name"]
    if style.get("font_size") is not None:
        run.font.size = style["font_size"]
    if style.get("color_rgb") is not None:
        run.font.color.rgb = style["color_rgb"]


def _apply_modifications_to_text(input_path: str, modifications: list[dict[str, Any]]) -> str:
    content = _read_text_file(input_path)
    for modification in modifications:
        modification_type = str(modification.get("type") or "").lower()
        if modification_type == "replace_text":
            old_text = str(modification.get("old_text") or "")
            new_text = str(modification.get("new_text") or "")
            if old_text:
                content = content.replace(old_text, new_text)
        elif modification_type == "insert_text":
            insert_text = str(modification.get("text") or "")
            position = str(modification.get("position") or "end").lower()
            line_index = modification.get("line_index")
            if not isinstance(line_index, int):
                paragraph_index = modification.get("paragraph_index")
                if isinstance(paragraph_index, int):
                    line_index = paragraph_index
            reference_text = str(modification.get("reference_text") or "")
            if not insert_text:
                continue
            if isinstance(line_index, int):
                lines = content.split("\n")
                safe_index = max(0, min(line_index, len(lines) - 1 if lines else 0))

                if position == "before":
                    lines.insert(safe_index, insert_text)
                elif position == "after":
                    lines.insert(min(safe_index + 1, len(lines)), insert_text)
                elif position == "start" and lines:
                    lines[safe_index] = f"{insert_text}{lines[safe_index]}"
                elif position == "end" and lines:
                    lines[safe_index] = f"{lines[safe_index]}{insert_text}"
                else:
                    lines.append(insert_text)

                content = "\n".join(lines)
            elif reference_text and reference_text in content and position in {"before", "after"}:
                replacement = f"{insert_text}{reference_text}" if position == "before" else f"{reference_text}{insert_text}"
                content = content.replace(reference_text, replacement, 1)
            elif position == "start":
                content = f"{insert_text}{content}"
            else:
                content = f"{content}{insert_text}"

    output_path = derive_output_path(input_path, prefix="modified_")
    with open(output_path, "w", encoding="utf-8") as output_file:
        output_file.write(content)
    return output_path


__all__ = [
    "_apply_modifications_to_docx",
    "_apply_modifications_to_text",
    "_resolve_docx_paragraph_indexes",
    "_write_content_to_docx",
    "docx",
]
