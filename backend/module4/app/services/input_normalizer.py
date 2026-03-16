from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence

from app.schemas import UploadSyncFile


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
}
OFFICE_TEXT_EXTENSIONS = {".docx", ".xlsx", ".xlsm", ".pdf"}
EXPLICIT_UNSUPPORTED_EXTENSIONS = {".doc", ".xls"}


@dataclass(frozen=True)
class NormalizedInputFile:
    file_name: str
    file_hash: str
    text_content: str
    source_kind: str = "upload"
    original_path: str | None = None
    artifact_path: str | None = None
    artifact_url: str | None = None
    mime_type: str | None = None
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class FailedInputFile:
    file_name: str
    file_hash: str
    error: str
    source_kind: str = "upload"
    original_path: str | None = None
    artifact_path: str | None = None
    artifact_url: str | None = None
    mime_type: str | None = None
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class InputNormalizationResult:
    normalized: List[NormalizedInputFile]
    failed: List[FailedInputFile]

    @property
    def partial(self) -> bool:
        return bool(self.normalized and self.failed)


def _clean_text(text: str) -> str:
    value = str(text or "").replace("\x00", " ")
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    value = re.sub(r"[ \t]{2,}", " ", value)
    return value.strip()


def _resolve_existing_path(upload: UploadSyncFile) -> Path:
    candidates = [upload.storage_path, upload.artifact_path]
    for raw in candidates:
        if not raw:
            continue
        path = Path(str(raw)).expanduser()
        if path.exists():
            return path.resolve()
    raise FileNotFoundError(f"Artifact path not found for {upload.file_name}")


def _read_text_file(path: Path) -> str:
    last_error: Exception | None = None
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "gbk"):
        try:
            return _clean_text(path.read_text(encoding=encoding))
        except Exception as exc:  # pragma: no cover - exercised by fallback path
            last_error = exc
    raw = path.read_bytes()
    try:
        return _clean_text(raw.decode("utf-8", errors="ignore"))
    except Exception as exc:  # pragma: no cover - defensive
        raise RuntimeError(f"Unable to decode text file {path.name}: {exc}") from (last_error or exc)


def _extract_pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: List[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = _clean_text(page.extract_text() or "")
        if text:
            parts.append(f"[Page {index}]\n{text}")
    return _clean_text("\n\n".join(parts))


def _iter_docx_blocks(document: Any) -> Iterable[Any]:
    from docx.document import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent = document.element.body if isinstance(document, DocxDocument) else document._tc  # type: ignore[attr-defined]
    for child in parent.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _extract_docx_text_xml_fallback(path: Path) -> str:
    import zipfile
    from xml.etree import ElementTree as ET

    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

    def _paragraph_text(node: Any) -> str:
        text = "".join(str(item.text or "") for item in node.findall(".//w:t", namespace))
        return _clean_text(text)

    parts: List[str] = []
    with zipfile.ZipFile(path) as archive:
        document_xml = ET.fromstring(archive.read("word/document.xml"))
        body = document_xml.find("w:body", namespace)
        if body is None:
            return ""

        table_index = 0
        for child in list(body):
            if child.tag.endswith("}p"):
                text = _paragraph_text(child)
                if text:
                    parts.append(text)
                continue
            if child.tag.endswith("}tbl"):
                table_index += 1
                rows: List[str] = []
                for row in child.findall(".//w:tr", namespace):
                    cells: List[str] = []
                    for cell in row.findall("./w:tc", namespace):
                        cell_text = _clean_text(" ".join(filter(None, (_paragraph_text(p) for p in cell.findall(".//w:p", namespace)))))
                        if cell_text:
                            cells.append(cell_text)
                    if cells:
                        rows.append(" | ".join(cells))
                if rows:
                    parts.append(f"[Table {table_index}]")
                    parts.extend(rows)

        for extra_name, label in (("word/footnotes.xml", "Footnotes"), ("word/endnotes.xml", "Endnotes")):
            if extra_name not in archive.namelist():
                continue
            extra_root = ET.fromstring(archive.read(extra_name))
            extra_parts: List[str] = []
            for paragraph in extra_root.findall(".//w:p", namespace):
                text = _paragraph_text(paragraph)
                if text:
                    extra_parts.append(text)
            if extra_parts:
                parts.append(f"[{label}]")
                parts.extend(extra_parts)

    return _clean_text("\n\n".join(parts))


def _extract_docx_text(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    fallback_error: Exception | None = None
    try:
        document = Document(str(path))
        parts: List[str] = []
        table_index = 0
        for block in _iter_docx_blocks(document):
            if isinstance(block, Paragraph):
                text = _clean_text(block.text)
                if text:
                    parts.append(text)
                continue
            if isinstance(block, Table):
                table_index += 1
                rows: List[str] = []
                for row in block.rows:
                    cells = [_clean_text(cell.text) for cell in row.cells]
                    cells = [cell for cell in cells if cell]
                    if cells:
                        rows.append(" | ".join(cells))
                if rows:
                    parts.append(f"[Table {table_index}]")
                    parts.extend(rows)
        text = _clean_text("\n\n".join(parts))
        if text:
            return text
    except Exception as exc:
        fallback_error = exc

    fallback_text = _extract_docx_text_xml_fallback(path)
    if fallback_text:
        return fallback_text
    if fallback_error is not None:
        raise RuntimeError(f"Unable to extract DOCX text from {path.name}: {fallback_error}") from fallback_error
    return ""


def _extract_xlsx_text(path: Path, *, sheet_limit: int = 12, cell_cap: int = 4000) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    emitted_cells = 0
    parts: List[str] = []
    try:
        for sheet in workbook.worksheets[:sheet_limit]:
            rows: List[str] = []
            for raw_row in sheet.iter_rows(values_only=True):
                values = [_clean_text(str(cell)) for cell in raw_row if cell is not None and _clean_text(str(cell))]
                if not values:
                    continue
                emitted_cells += len(values)
                rows.append(" | ".join(values))
                if emitted_cells >= cell_cap:
                    break
            if rows:
                parts.append(f"# Sheet: {sheet.title}")
                parts.extend(rows)
            if emitted_cells >= cell_cap:
                break
    finally:
        workbook.close()
    return _clean_text("\n\n".join(parts))


class UploadInputNormalizer:
    def normalize_uploads(self, uploads: Sequence[UploadSyncFile | Dict[str, Any]]) -> InputNormalizationResult:
        normalized: List[NormalizedInputFile] = []
        failed: List[FailedInputFile] = []

        for raw_item in uploads:
            fallback = raw_item.model_dump(mode="json") if isinstance(raw_item, UploadSyncFile) else dict(raw_item or {})
            try:
                upload = raw_item if isinstance(raw_item, UploadSyncFile) else UploadSyncFile.model_validate(raw_item)
                text = self._normalize_single(upload)
                normalized.append(
                    NormalizedInputFile(
                        file_name=upload.file_name,
                        file_hash=upload.file_hash,
                        text_content=text,
                        source_kind="upload",
                        original_path=upload.storage_path or upload.artifact_path,
                        artifact_path=upload.artifact_path,
                        artifact_url=upload.artifact_url,
                        mime_type=upload.mime_type,
                        metadata=dict(upload.metadata or {}),
                    )
                )
            except Exception as exc:
                failed.append(
                    FailedInputFile(
                        file_name=str(fallback.get("file_name") or "unknown"),
                        file_hash=str(fallback.get("file_hash") or ""),
                        error=f"{type(exc).__name__}: {exc}",
                        source_kind="upload",
                        original_path=str(fallback.get("storage_path") or fallback.get("artifact_path") or "") or None,
                        artifact_path=str(fallback.get("artifact_path") or "") or None,
                        artifact_url=str(fallback.get("artifact_url") or "") or None,
                        mime_type=str(fallback.get("mime_type") or "") or None,
                        metadata=dict(fallback.get("metadata") or {}),
                    )
                )

        return InputNormalizationResult(normalized=normalized, failed=failed)

    def _normalize_single(self, upload: UploadSyncFile) -> str:
        inline_text = _clean_text(upload.text_content or "")
        if inline_text:
            return inline_text

        path = _resolve_existing_path(upload)
        suffix = Path(upload.file_name or path.name).suffix.lower()

        if suffix in EXPLICIT_UNSUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported legacy Office format: {suffix}. Convert it to .docx or .xlsx first.")
        if suffix in TEXT_EXTENSIONS:
            text = _read_text_file(path)
        elif suffix == ".pdf":
            text = _extract_pdf_text(path)
        elif suffix == ".docx":
            text = _extract_docx_text(path)
        elif suffix in {".xlsx", ".xlsm"}:
            text = _extract_xlsx_text(path)
        else:
            raise ValueError(f"Unsupported file type for GraphRAG normalization: {suffix or 'unknown'}")

        text = _clean_text(text)
        if not text:
            raise ValueError(f"No extractable text found in {upload.file_name}")
        return text
